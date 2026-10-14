"""Разбор банка вопросов (bank_parse.py, bank_docx.py).

Корпус: синтетические файлы fixtures/banks с эталоном expected.json (их пересобирает make_corpus.py) — каждый
файл разбирается в эталон: вопросы, варианты, ключи, виды, флаги, тексты дословно, темы, предметы,
предупреждения, пропуски номеров. Синтетические банки fixtures/banks/synthetic (C#, базы данных, сети, ОС,
алгоритмизация, архитектура ЭВМ; эталон truth.json) — каждый вопрос, вариант и ключ дословно. Дальше — код
без изменений, источники ключа, флаги, доли отметок (marks), предметы и темы, открытые вопросы, format и
key_mode, пределы, безопасность разбора файлов, ошибки, найденные на синтетических банках, и скорость (банк на
240 вопросов — меньше 0,5 с, сессия 4 × 240 — меньше 1,5 с). Файлы колледжа сюда не кладутся.
"""
import io
import json
import sys
import time
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import bank_docx  # noqa: E402
import bank_parse  # noqa: E402
from bank_parse import BankError, parse_bank_file, parse_bank_text  # noqa: E402

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "banks"
SPECS = json.loads((FIXTURES / "expected.json").read_text(encoding="utf-8"))
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
SIDE_FLAGS = ("has_image", "has_formula", "joined_lines")  # сопутствующие флаги: в эталоне не перечислены


# ── Корпус ──────────────────────────────────────────────────────────────────────


def corpus_problems(spec: dict, result: dict) -> list[str]:
    """Расхождения разбора с эталоном (как check.py прототипа, плюс тексты, номера, темы, предметы)."""
    got, want = result["items"], spec["items"]
    problems = []
    if len(got) != len(want):
        problems.append(f"items {len(got)} != {len(want)}")
    topics_in_file = spec.get("topics_in_file", True) and bool(spec["topics"])
    for g, w in zip(got, want):
        tag = f"№{w['label'] or w['text'][:20]}"
        expected_flags = spec["flags"].get(w["label"], [])
        if w["kind"] in ("single", "multiple", "order") and w["options"] and g["options"] != w["options"]:
            problems.append(f"{tag} options {g['options']} != {w['options']}")
        if w["kind"] in ("single", "multiple") and sorted(g["correct"]) != sorted(w["correct"]) \
                and not set(expected_flags) & {"key_conflict", "no_key", "several_marked"}:
            problems.append(f"{tag} correct {g['correct']} != {w['correct']}")
        if w["kind"] != g["kind"] and not (w["kind"] == "single" and g["kind"] == "truefalse"):
            problems.append(f"{tag} kind {g['kind']} != {w['kind']}")
        if w["kind"] == "open" and (g["options"] or g["correct"] or "no_key" in g["flags"]):
            problems.append(f"{tag} open item with options, key or no_key: {g}")
        if w["kind"] == "order" and w.get("order") and g["order"] != [w["order"].index(p) for p in range(1, len(w["order"]) + 1)]:
            problems.append(f"{tag} order {g['order']} != places {w['order']}")
        if w["kind"] == "open" and w.get("answer") and g["answer"] != w["answer"]:
            problems.append(f"{tag} answer {g['answer']} != {w['answer']}")
        missing = [f for f in expected_flags if f not in g["flags"]]
        extra = [f for f in g["flags"] if f not in expected_flags and f not in SIDE_FLAGS]
        if missing or extra:
            problems.append(f"{tag} flags missing {missing} extra {extra}")
        # формула OMML остаётся в тексте вопроса (флаг has_formula): эталон — текст до неё
        if g["text"] != w["text"] and not ("has_formula" in g["flags"] and g["text"].startswith(w["text"])):
            problems.append(f"{tag} text {g['text']!r} != {w['text']!r}")
        if g["label"] != w["label"]:
            problems.append(f"{tag} label {g['label']!r} != {w['label']!r}")
        if "subject" in w and g["subject"] != w["subject"]:
            problems.append(f"{tag} subject {g['subject']} != {w['subject']}")
        if topics_in_file and g["topic"] != w.get("topic"):
            problems.append(f"{tag} topic {g['topic']} != {w.get('topic')}")
    if topics_in_file and [t["title"] for t in result["topics"]] != spec["topics"]:
        problems.append(f"topics {[t['title'] for t in result['topics']]} != {spec['topics']}")
    if "topic_subjects" in spec and [t["subject"] for t in result["topics"]] != spec["topic_subjects"]:
        problems.append(f"topic subjects {[t['subject'] for t in result['topics']]} != {spec['topic_subjects']}")
    subjects = [{"title": s["title"], "evidence": s["evidence"]} for s in result["subjects"]]
    if subjects != spec.get("subjects", [{"title": None, "evidence": "none"}]):
        problems.append(f"subjects {subjects}")
    if sorted(w["code"] for w in result["warnings"]) != sorted(spec.get("warnings", [])):
        problems.append(f"warnings {result['warnings']}")
    if result["stats"]["gaps"] != spec.get("gaps", []):
        problems.append(f"gaps {result['stats']['gaps']}")
    return problems


def parse_spec(spec: dict) -> dict:
    key_text = None
    if spec["key"].startswith("key_file:"):
        key_text = (FIXTURES / spec["key"].split(":", 1)[1]).read_text(encoding="utf-8")
    contents = (FIXTURES / spec["file"]).read_bytes()
    return parse_bank_file(contents, spec["file"], key_text=key_text, format_mode=spec.get("format", "auto"))


@pytest.mark.parametrize("spec", SPECS, ids=[f"{s['file']}-{s.get('format', 'auto')}" for s in SPECS])
def test_every_corpus_file_parses_to_its_expected_result(spec):
    assert corpus_problems(spec, parse_spec(spec)) == []


def test_corpus_covers_every_layout_key_source_and_the_session_cases():
    files = {spec["file"] for spec in SPECS}
    assert len(files) == 24 and len(SPECS) == 25  # 17 файлов прототипа (без PDF) и 7 файлов сессии
    layouts = {parse_spec(spec)["layout"] for spec in SPECS}
    assert {"table_option_rows", "table_options_in_cell", "table_option_columns", "numbered_paragraphs", "blocks", "gift"} <= layouts
    sources = set()
    for spec in SPECS:
        sources |= set(parse_spec(spec)["key_source"])
    assert {"bold", "underline", "color", "key_column", "key_list", "key_text", "marker", "answer_line", "gift",
            "order_numbers"} <= sources


def test_generator_is_committed_next_to_its_output():
    source = (FIXTURES / "make_corpus.py").read_text(encoding="utf-8")
    assert "expected.json" in source and "def main()" in source
    for spec in SPECS:
        assert (FIXTURES / spec["file"]).is_file()


# ── Синтетические банки: эталон subject_*.json (truth.json) ─────────────────────

SYNTHETIC = FIXTURES / "synthetic"
TRUTH = json.loads((SYNTHETIC / "truth.json").read_text(encoding="utf-8"))


def synthetic_expected(entry: dict) -> list[tuple[dict, str, str]]:
    """Вопросы файла по эталону — (вопрос с номером в файле, предмет, тема): выбранные номера идут подряд с 1."""
    rows = []
    for part in entry["parts"]:
        subject = TRUTH["subjects"][part["subject"]]
        topics = {topic["index"]: topic["title"] for topic in subject["topics"]}
        by_number = {item["n"]: item for item in subject["items"]}
        for k, number in enumerate(part["pick"] or sorted(by_number), start=1):
            item = dict(by_number[number], n=k)
            rows.append((item, subject["subject"], topics[item["topic"]]))
    return rows


def synthetic_problems(entry: dict, result: dict) -> list[str]:
    """Расхождения с эталоном: номер, текст с кодом (дословно), варианты, ключ, вид, тема, предмет, флаги проверки."""
    want, got = synthetic_expected(entry), result["items"]
    problems = [] if len(got) == len(want) else [f"items {len(got)} != {len(want)}"]
    topics = {topic["index"]: topic["title"] for topic in result["topics"]}
    subjects = {subject["index"]: subject["title"] for subject in result["subjects"]}
    several = len(entry["parts"]) > 1
    for (w, subject, topic), g in zip(want, got):
        tag = f"{subject} №{w['n']}"
        text = w["text"] + ("\n" + w["code"] if w["code"] else "")
        checks = {
            "label": (g["label"], str(w["n"])), "text": (g["text"], text), "options": (g["options"], w["options"]),
            "correct": (sorted(g["correct"]), w["correct"]), "kind": (g["kind"], w["kind"]),
            "topic": (topics.get(g["topic"]), topic), "flags": (sorted(set(g["flags"]) & bank_parse.PROBLEM_FLAGS), []),
        }
        if several:
            checks["subject"] = (subjects.get(g["subject"]), subject)
        problems += [f"{tag} {name}: {got_value!r} != {want_value!r}"
                     for name, (got_value, want_value) in checks.items() if got_value != want_value]
    return problems


@pytest.mark.parametrize("entry", TRUTH["files"], ids=[entry["file"] for entry in TRUTH["files"]])
def test_synthetic_banks_parse_to_their_ground_truth_item_for_item(entry):
    result = parse_bank_file((SYNTHETIC / entry["file"]).read_bytes(), entry["file"])
    assert synthetic_problems(entry, result) == []
    if len(entry["parts"]) > 1:  # предметы файла — по заголовкам «Дисциплина: …»
        assert [(s["title"], s["evidence"]) for s in result["subjects"]] == [
            (TRUTH["subjects"][part["subject"]]["subject"], "heading") for part in entry["parts"]]


def test_synthetic_set_covers_every_layout_and_key_source():
    files = {entry["file"] for entry in TRUTH["files"]}
    assert files == {path.name for path in SYNTHETIC.iterdir() if path.name != "truth.json"}
    sources, layouts, kinds, items = set(), set(), set(), 0
    for entry in TRUTH["files"]:
        result = parse_bank_file((SYNTHETIC / entry["file"]).read_bytes(), entry["file"])
        sources |= set(result["key_source"])
        layouts.add(result["layout"])
        kinds |= set(result["stats"]["kinds"])
        items += result["stats"]["items"]
    assert sources == {"bold", "highlight", "key_list", "marker"} and kinds == {"single", "multiple", "open"}
    assert {"table_option_rows", "numbered_paragraphs"} <= layouts and items == 682


# ── Помощники: DOCX из шаблона корпуса и маленькие документы python-docx ─────────


def docx_with_body(body: str, prolog: str = "") -> bytes:
    """DOCX с другим телом документа на основе файла корпуса (python-docx на тысячах строк таблицы медленный)."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO((FIXTURES / "13_inline_bold.docx").read_bytes())) as source, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                data = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>{prolog}'
                        f'<w:document xmlns:w="{W_NS}" xmlns:m="{M_NS}"><w:body>{body}<w:sectPr/></w:body></w:document>').encode()
            target.writestr(info, data)
    return out.getvalue()


def xml_paragraph(text: str, bold: bool = False) -> str:
    props = "<w:rPr><w:b/></w:rPr>" if bold else ""
    return f'<w:p><w:r>{props}<w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'


def xml_cell(content: str, props: str = "") -> str:
    return f"<w:tc><w:tcPr>{props}</w:tcPr>{content or '<w:p/>'}</w:tc>"


def owner_table(subject: int, items: int, topics: int, options: int = 5) -> str:
    """Таблица как у владельца: «№ | Содержание вопроса | Варианты ответов», строки тем на всю ширину,
    № и вопрос объединены вниз, вариант на строку, правильный (номер вопроса по модулю options) — жирным."""
    rows = ["<w:tr>" + "".join(xml_cell(xml_paragraph(text, bold=True)) for text in
                                ("№", "Содержание вопроса", "Варианты ответов")) + "</w:tr>"]
    number = 0
    for topic in range(1, topics + 1):
        rows.append("<w:tr>" + xml_cell(xml_paragraph(f"Тема {topic}. Раздел {topic} предмета {subject}", bold=True),
                                        '<w:gridSpan w:val="3"/>') + "</w:tr>")
        for _ in range(items // topics):
            number += 1
            for k in range(options):
                first = k == 0
                merge = '<w:vMerge w:val="restart"/>' if first else "<w:vMerge/>"
                rows.append(
                    "<w:tr>"
                    + xml_cell(xml_paragraph(str(number)) if first else "", merge)
                    + xml_cell(xml_paragraph(f"Вопрос {subject}.{number}: что делает строка count += {number}; в цикле?")
                               if first else "", merge)
                    + xml_cell(xml_paragraph(f"вариант {k + 1} к вопросу {subject}.{number}", bold=k == number % options))
                    + "</w:tr>")
    return "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/><w:gridCol/></w:tblGrid>" + "".join(rows) + "</w:tbl>"


def docx_bytes(build) -> bytes:
    doc = Document()
    build(doc)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def best_seconds(call, repeat: int = 3) -> float:
    """Лучшее время из нескольких прогонов: машина под нагрузкой полного прогона даёт шум."""
    best = None
    for _ in range(repeat):
        started = time.perf_counter()
        call()
        elapsed = time.perf_counter() - started
        best = elapsed if best is None else min(best, elapsed)
    return best


# ── Скорость ────────────────────────────────────────────────────────────────────


def test_a_240_question_bank_parses_in_under_half_a_second():
    contents = docx_with_body(owner_table(1, 240, 8))
    result = parse_bank_file(contents, "C# и .NET.docx")
    assert result["stats"]["items"] == 240 and result["key_source"] == {"bold": 240}
    assert [item["correct"] for item in result["items"]] == [[n % 5] for n in range(1, 241)]
    assert all(item["flags"] == [] and len(item["options"]) == 5 for item in result["items"])
    assert len(result["topics"]) == 8 and result["subjects"] == [{"index": 1, "title": None, "evidence": "none", "items": 240}]
    assert best_seconds(lambda: parse_bank_file(contents, "C# и .NET.docx")) < 0.5


def test_the_synthetic_240_question_csharp_table_parses_in_under_half_a_second():
    """Таблица колледжа на 240 вопросов C#: 87 с кодом по абзацу на строку, 36 — с пустыми строками внутри кода."""
    contents = (SYNTHETIC / "csharp_table.docx").read_bytes()
    result = parse_bank_file(contents, "csharp_table.docx")
    assert result["stats"]["items"] == 240 and result["key_source"] == {"bold": 240} and result["stats"]["to_check"] == 0
    assert best_seconds(lambda: parse_bank_file(contents, "csharp_table.docx")) < 0.5


SESSION_SUBJECTS = ("Основы программирования", "Базы данных", "Компьютерные сети", "Операционные системы")


def test_a_960_question_session_of_four_subjects_parses_in_under_one_and_a_half_seconds():
    body = "".join(xml_paragraph(f"Дисциплина: {title}") + owner_table(index, 240, 8)
                   for index, title in enumerate(SESSION_SUBJECTS, start=1))
    contents = docx_with_body(body)
    result = parse_bank_file(contents, "Сессия.docx")
    assert result["stats"]["items"] == 960 and result["key_source"] == {"bold": 960}
    assert result["subjects"] == [{"index": i, "title": title, "evidence": "heading", "items": 240}
                                  for i, title in enumerate(SESSION_SUBJECTS, start=1)]
    assert [topic["subject"] for topic in result["topics"]] == [s for s in range(1, 5) for _ in range(8)]
    assert [topic["index"] for topic in result["topics"]] == list(range(1, 33))
    assert all(item["flags"] == [] for item in result["items"]) and result["stats"]["gaps"] == []
    assert [item["label"] for item in result["items"][235:245]] == ["236", "237", "238", "239", "240", "1", "2", "3", "4", "5"]
    assert result["warnings"] == []
    assert best_seconds(lambda: parse_bank_file(contents, "Сессия.docx")) < 1.5


# ── Код дословно ────────────────────────────────────────────────────────────────

CODE_BANK = """1. Какой оператор вычитает и сразу присваивает результат?
а) +=
б) **-=**
в) =-
г) %=
2. Что делает метод __init__ в Python?
а) **__init__ вызывается при создании объекта**
б) __init__ удаляет объект
3. Какая директива начинает сворачиваемый блок?
а) **#region Helpers**
б) #define TRACE
в) #pragma warning disable
4. Какой оператор уменьшает число на единицу?
а) **--**
б) *
в) ++
г) ~
5. В какой строке есть оператор условного доступа?
а) **Changed?.Invoke(this, e);**
б) Changed.Invoke(this, e);
в) count -= step;
"""


def test_code_survives_verbatim_in_questions_and_options():
    result = parse_bank_text(CODE_BANK)
    assert [item["options"] for item in result["items"]] == [
        ["+=", "-=", "=-", "%="],
        ["__init__ вызывается при создании объекта", "__init__ удаляет объект"],
        ["#region Helpers", "#define TRACE", "#pragma warning disable"],
        ["--", "*", "++", "~"],
        ["Changed?.Invoke(this, e);", "Changed.Invoke(this, e);", "count -= step;"],  # «;» кода не срезается
    ]
    assert [item["correct"] for item in result["items"]] == [[1], [0], [0], [0], [0]]
    assert result["key_source"] == {"bold": 5} and result["marker"] is None
    assert all(item["flags"] == [] for item in result["items"])
    assert [item["code"] for item in result["items"]] == [True, False, True, False, True]
    owner = parse_spec(SPECS[0])  # таблица как у владельца: код в вариантах
    options = {option for item in owner["items"] for option in item["options"]}
    assert {"//", "#", "--", "#region", "+=", "=>", "::", "*", "Func<int>", "Converter<int, int>"} <= options
    assert any(item["text"] == "Что вернёт выражение items?.Count ?? 0, если items равно null?" for item in owner["items"])


def test_open_question_text_keeps_code_and_is_not_cleaned():
    result = parse_spec(next(spec for spec in SPECS if spec["file"] == "23_open_autonumber.docx"))
    texts = [item["text"] for item in result["items"]]
    assert "Что выведет Console.WriteLine(5 / 2); и почему?" in texts
    assert "Что такое __init__ в Python и есть ли похожий метод в C#?" in texts
    assert "Что делает строка count -= step; в цикле while (count > 0)?" in texts
    assert "Зачем нужен #region Helpers и когда без него лучше обойтись?" in texts
    assert all(item["kind"] == "open" and item["options"] == [] and item["marks"] == {} for item in result["items"])


# ── Источники ключа ─────────────────────────────────────────────────────────────


def _shade(run) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "FFFF00")
    run._r.get_or_add_rPr().append(shd)


MARKERS = {
    "bold": lambda run: setattr(run, "bold", True),
    "underline": lambda run: setattr(run, "underline", True),
    "italic": lambda run: setattr(run, "italic", True),
    "highlight": lambda run: setattr(run.font, "highlight_color", WD_COLOR_INDEX.YELLOW),
    "color": lambda run: setattr(run.font.color, "rgb", RGBColor(0xFF, 0x00, 0x00)),
    "shade": _shade,
}


@pytest.mark.parametrize("signal", sorted(MARKERS))
def test_the_key_is_read_from_each_formatting_signal(signal):
    def build(doc):
        for number in range(1, 5):
            doc.add_paragraph(f"{number}. Вопрос номер {number}?")
            for k in range(4):
                p = doc.add_paragraph(f"{'абвг'[k]}) ")
                run = p.add_run(f"вариант {k + 1} вопроса {number}")
                if k == number % 4:
                    MARKERS[signal](run)

    result = parse_bank_file(docx_bytes(build), "bank.docx")
    assert result["format_signal"] == signal and result["key_source"] == {signal: 4}
    assert [item["correct"] for item in result["items"]] == [[n % 4] for n in range(1, 5)]
    assert list(result["items"][0]["marks"]) == [signal]  # признаки без отметок в документе не приходят
    assert result["items"][0]["marks"][signal] == [0, 1, 0, 0]


def test_checkbox_symbols_mark_the_key():
    def build(doc):
        for number in (1, 2):
            doc.add_paragraph(f"{number}. Вопрос {number}?")
            for k in range(3):
                p = doc.add_paragraph()
                sym = OxmlElement("w:sym")
                sym.set(qn("w:font"), "Wingdings")
                sym.set(qn("w:char"), "F0FE" if k == number else "F0A8")
                p.add_run()._r.append(sym)
                p.add_run(f" вариант {k + 1}")

    result = parse_bank_file(docx_bytes(build), "bank.docx")
    assert result["key_source"] == {"checkbox": 2}
    assert [item["correct"] for item in result["items"]] == [[1], [2]]


def test_key_mode_first_option_and_none_override_the_file_and_key_text_still_applies():
    contents = (FIXTURES / "01_table_rows_bold.docx").read_bytes()
    first = parse_bank_file(contents, "bank.docx", key_mode="first_option")
    assert first["key_source"] == {"first_option": 12}
    assert all(item["correct"] == [0] and "no_key" not in item["flags"] for item in first["items"])
    none = parse_bank_file(contents, "bank.docx", key_mode="none")
    assert none["key_source"] == {} and none["stats"]["no_key"] == 12
    assert all(item["correct"] == [] and item["flags"] == ["no_key"] for item in none["items"])
    assert none["items"][0]["marks"]["bold"] == [1, 0, 0, 0]  # клиент всё равно может взять жирный сам
    keyed = parse_bank_file(contents, "bank.docx", key_mode="none", key_text="1 – б\n2. в")
    assert keyed["key_source"] == {"key_text": 2}
    assert [item["correct"] for item in keyed["items"][:3]] == [[1], [2], []]


def test_key_text_is_a_key_list_and_wins_over_the_file_with_a_conflict_flag():
    contents = (FIXTURES / "05_key_column_inline.docx").read_bytes()
    result = parse_bank_file(contents, "bank.docx", key_text="Ответы:\n1 – в; 2 – в")
    first, second = result["items"][:2]
    assert first["correct"] == [2] and first["evidence"] == "key_text" and "key_conflict" in first["flags"]
    assert second["correct"] == [2] and second["evidence"] == "key_text" and "key_conflict" not in second["flags"]
    assert bank_parse.parse_key_text("Ключ\n1 – б; 2-в\n3. а\n4б") == {"1": "б", "2": "в", "3": "а", "4": "б"}


def test_key_text_that_matches_nothing_is_reported():
    contents = (FIXTURES / "15_pasted_no_key.txt").read_bytes()
    no_pairs = parse_bank_file(contents, "bank.txt", key_text="правильные ответы смотрите в журнале")
    assert [w["code"] for w in no_pairs["warnings"]] == ["key_text_unused"]
    assert no_pairs["warnings"][0]["message"] == bank_parse.KEY_TEXT_NO_PAIRS_WARNING
    no_match = parse_bank_file(contents, "bank.txt", key_text="99 – а")
    assert no_match["warnings"][0]["message"] == bank_parse.KEY_TEXT_NO_MATCH_WARNING
    applied = parse_bank_file(contents, "bank.txt", key_text="1 – в\n2 – б\n3 – б\n4 – а")
    assert applied["warnings"] == [] and applied["key_source"] == {"key_text": 4} and applied["stats"]["no_key"] == 0


# ── Флаги, пропуски номеров, доли отметок ───────────────────────────────────────

FLAGGED = """1. Один вариант?
а) только этот
2. Девять вариантов?
а) 1
б) 2
в) 3
г) 4
д) 5
е) 6
ж) 7
з) 8
и) 9
3. Одинаковые варианты?
а) да
б) да
в) нет
4. Пустой вариант?
а) x;
б) ;
в) z.
"""


def test_validation_flags():
    result = parse_bank_text(FLAGGED)
    assert [item["flags"] for item in result["items"]] == [
        ["few_options", "no_key"], ["many_options", "no_key"], ["no_key", "same_options"], ["no_key", "empty_option"],
    ]
    assert result["stats"]["to_check"] == 4 and result["confidence"] == 0.0
    test_only = parse_bank_text("1. Вопрос без вариантов?\n2. Вопрос?\nа) да\nб) нет\n", format_mode="test")
    assert test_only["items"][0]["kind"] == "unknown" and test_only["items"][0]["flags"] == ["no_options"]


def test_long_texts_are_cut_to_the_database_limits_with_a_flag():
    question = "Очень длинный вопрос " * 220  # больше 4000 знаков
    option = "очень длинный вариант " * 30  # больше 600 знаков
    result = parse_bank_text(f"1. {question}?\nа) **{option.strip()}**\nб) короткий\n")
    item = result["items"][0]
    assert len(item["text"]) <= 4000 and question.startswith(item["text"]) and len(item["options"][0]) <= 600
    assert item["flags"] == ["too_long"] and item["correct"] == [0] and result["stats"]["to_check"] == 1


def test_numbering_gaps_are_reported_per_subject():
    messy = parse_spec(next(spec for spec in SPECS if spec["file"] == "14_messy.docx"))
    assert messy["stats"]["gaps"] == [[8, 8]]
    restart = parse_spec(next(spec for spec in SPECS if spec["file"] == "21_two_tables_restart.docx"))
    assert restart["stats"]["gaps"] == []  # нумерация заново во втором предмете — не пропуск


def test_marks_give_the_share_of_marked_characters_per_option_for_each_signal_in_the_document():
    def build(doc):
        doc.add_paragraph("1. Первый вопрос?")
        doc.add_paragraph("а) ").add_run("alpha").bold = True
        doc.add_paragraph("б) beta")
        doc.add_paragraph("в) gamma")
        doc.add_paragraph("2. Второй вопрос?")
        p = doc.add_paragraph("а) ")
        p.add_run("Stack").bold = True  # 5 из 13 видимых знаков
        p.add_run("Overflow")
        doc.add_paragraph("б) GitHub")
        doc.add_paragraph("в) ").add_run("Habr").underline = True

    result = parse_bank_file(docx_bytes(build), "bank.docx")
    first, second = result["items"]
    assert first["marks"] == {"bold": [1, 0, 0], "underline": [0, 0, 0]}
    assert second["marks"] == {"bold": [0.38, 0, 0], "underline": [0, 0, 1]}
    assert result["format_signal"] == "bold" and result["mark_share"] == 0.6
    assert first["correct"] == [0] and second["correct"] == [] and second["flags"] == ["partial_mark", "no_key"]
    assert isinstance(first["marks"]["bold"][0], int) and isinstance(second["marks"]["bold"][1], int)


@pytest.mark.parametrize(("marked", "letters", "value"), [
    (0, 5, 0), (5, 5, 1), (3, 5, 0.6), (599, 1000, 0.59), (1, 200, 0.01), (2, 3, 0.66), (0, 0, 0),
])
def test_mark_values_round_down_so_the_client_threshold_matches_the_server(marked, letters, value):
    assert bank_parse._mark_value(marked, letters) == value
    if letters:
        assert (value >= bank_parse.MARK_SHARE) == (marked / letters >= bank_parse.MARK_SHARE)


def test_marks_are_trimmed_with_list_punctuation_so_a_short_bold_option_is_still_the_key():
    def build(doc):
        for number in (1, 2):
            p = doc.add_paragraph(f"{number}. Сколько байт в типе? ")
            for k, option in enumerate(("2", "4", "8")):
                p.add_run(f"{'абв'[k]}) ")
                p.add_run(option).bold = k == number
                p.add_run("; " if k < 2 else ".")

    result = parse_bank_file(docx_bytes(build), "bank.docx")
    assert [item["options"] for item in result["items"]] == [["2", "4", "8"]] * 2
    assert [item["correct"] for item in result["items"]] == [[1], [2]]
    assert result["items"][0]["marks"]["bold"] == [0, 1, 0]


# ── Предметы и темы ─────────────────────────────────────────────────────────────


@pytest.mark.parametrize(("line", "title"), [
    ("Дисциплина: Основы программирования", "Основы программирования"),
    ("Учебная дисциплина «Базы данных»", "Базы данных"),
    ("Учебная дисциплина Базы данных", "Базы данных"),
    ("ДИСЦИПЛИНА ОСНОВЫ ПРОГРАММИРОВАНИЯ", "ОСНОВЫ ПРОГРАММИРОВАНИЯ"),
    ("Предмет: История Беларуси", "История Беларуси"),
    ("Предмет «История Беларуси»", "История Беларуси"),
    ("Курс: Компьютерные сети", "Компьютерные сети"),
    ("МДК 01.01 Разработка программных модулей", "МДК 01.01 Разработка программных модулей"),
    ("ПМ.01 Разработка модулей программного обеспечения", "ПМ.01 Разработка модулей программного обеспечения"),
    ("Модуль 2. Базы данных", "Модуль 2. Базы данных"),
    ("Тестовые задания по учебной дисциплине «Основы программирования»", "Основы программирования"),
    ("Модуль 2", None),
    ("Курс 2", None),
    ("Курс лекций по экономике", None),
    ("Дисциплина труда на предприятии", None),
    ("Предмет и метод макроэкономики.", None),
    ("Что изучает дисциплина «Экономика»?", None),
    ("Тема 1. Введение", None),
])
def test_subject_titles_from_discipline_headings(line, title):
    assert bank_parse.subject_title(line) == title


def test_word_heading_levels_make_subjects_when_topics_sit_one_level_lower():
    def build(doc):
        number = 0
        for subject in ("Основы программирования", "Базы данных"):
            doc.add_heading(subject, level=1)
            for topic in ("Введение", "Практика"):
                doc.add_heading(topic, level=2)
                for _ in range(2):
                    number += 1
                    doc.add_paragraph(f"{number}. Вопрос {number}?")
                    doc.add_paragraph("а) ").add_run("да").bold = True
                    doc.add_paragraph("б) нет")

    result = parse_bank_file(docx_bytes(build), "bank.docx")
    assert [(s["title"], s["evidence"], s["items"]) for s in result["subjects"]] == [
        ("Основы программирования", "heading_level", 4), ("Базы данных", "heading_level", 4)]
    assert [(t["title"], t["subject"], t["items"]) for t in result["topics"]] == [
        ("Введение", 1, 2), ("Практика", 1, 2), ("Введение", 2, 2), ("Практика", 2, 2)]
    assert [item["topic"] for item in result["items"]] == [1, 1, 2, 2, 3, 3, 4, 4]


def test_a_merged_table_row_with_a_discipline_code_starts_a_subject():
    def row(cells, merged=False):
        if merged:
            return "<w:tr>" + xml_cell(xml_paragraph(cells, bold=True), '<w:gridSpan w:val="3"/>') + "</w:tr>"
        return "<w:tr>" + "".join(xml_cell(xml_paragraph(text) if text else "") for text in cells) + "</w:tr>"

    rows = [row(("№", "Вопрос", "Варианты ответов"))]
    for code, title in (("МДК 01.01", "Разработка программных модулей"), ("МДК 01.02", "Поддержка программных модулей")):
        rows.append(row(f"{code} {title}", merged=True))
        rows.append(row("Тема 1. Основы", merged=True))
        for number in (1, 2):
            rows.append(row((str(number), f"Вопрос {code} номер {number}?", "верный")))
            rows.append(row(("", "", "неверный")))
    table = "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/><w:gridCol/></w:tblGrid>" + "".join(rows) + "</w:tbl>"
    result = parse_bank_file(docx_with_body(table), "bank.docx", key_mode="first_option")
    assert [(s["title"], s["evidence"]) for s in result["subjects"]] == [
        ("МДК 01.01 Разработка программных модулей", "heading"), ("МДК 01.02 Поддержка программных модулей", "heading")]
    assert [(t["title"], t["subject"]) for t in result["topics"]] == [("Тема 1. Основы", 1), ("Тема 1. Основы", 2)]
    assert [item["subject"] for item in result["items"]] == [1, 1, 2, 2] and result["warnings"] == []


def numbered_questions(first: int, count: int) -> list[str]:
    lines = []
    for number in range(first, first + count):
        lines += [f"{number}. Вопрос номер {number}?", "а) **верно**", "б) неверно"]
    return lines


def test_numbering_restart_after_a_plain_heading_splits_subjects_and_asks():
    text = "\n".join(numbered_questions(1, 3) + ["Микроэкономика"] + numbered_questions(1, 2))
    result = parse_bank_text(text)
    assert [(s["title"], s["evidence"], s["items"]) for s in result["subjects"]] == [
        (None, "restart", 3), ("Микроэкономика", "restart", 2)]
    assert result["warnings"] == [{"code": "subjects_restart", "message": "Похоже, в файле несколько предметов. Разделить?"}]
    assert [item["options"] for item in result["items"]][2] == ["верно", "неверно"]  # заголовок не приклеился к варианту
    assert result["stats"]["items"] == 5 and result["key_source"] == {"bold": 5}


def test_numbering_restart_after_a_topic_heading_is_topic_numbering_not_a_new_subject():
    text = "\n".join(["Тема 1. Введение"] + numbered_questions(1, 3) + ["Тема 2. Практика"] + numbered_questions(1, 3))
    result = parse_bank_text(text)
    assert result["subjects"] == [{"index": 1, "title": None, "evidence": "none", "items": 6}]
    assert [(t["title"], t["subject"], t["items"]) for t in result["topics"]] == [
        ("Тема 1. Введение", 1, 3), ("Тема 2. Практика", 1, 3)]
    assert result["warnings"] == []


def test_more_than_300_topics_are_merged_into_the_last_one_with_a_warning():
    lines = []
    for number in range(1, 311):
        lines += [f"Тема {(number - 1) // 10 + 1}.{(number - 1) % 10 + 1}", f"{number}. Вопрос номер {number}?"]
    result = parse_bank_text("\n".join(lines))
    assert len(result["topics"]) == 300 and result["topics"][-1]["items"] == 11
    assert [item["topic"] for item in result["items"]][-12:] == [299] + [300] * 11
    assert [w["code"] for w in result["warnings"]] == ["topics_merged"]
    assert "310 тем" in result["warnings"][0]["message"]


# ── Открытые вопросы и format ───────────────────────────────────────────────────


def test_a_plain_numbered_exam_list_is_open_questions():
    text = "Вопросы к экзамену\nРаздел 1. Введение\n" + "\n".join(f"{n}. Тема вопроса номер {n} для устного ответа." for n in range(1, 21))
    result = parse_bank_text(text)
    assert result["stats"]["items"] == 20 and result["stats"]["kinds"] == {"open": 20}
    assert all(item["options"] == [] and item["correct"] == [] and item["flags"] == [] for item in result["items"])
    assert result["stats"]["no_key"] == 0 and result["skipped"] == ["Вопросы к экзамену"]
    assert [t["title"] for t in result["topics"]] == ["Раздел 1. Введение"]


def test_a_header_line_ending_with_a_colon_over_a_long_numbered_list_is_not_a_question():
    text = "Перечень вопросов к экзамену:\n" + "\n".join(f"{n}. Вопрос номер {n}." for n in range(1, 13))
    result = parse_bank_text(text)
    assert [item["label"] for item in result["items"]] == [str(n) for n in range(1, 13)]
    assert result["stats"]["kinds"] == {"open": 12}


def test_format_open_turns_gift_options_into_question_text():
    result = parse_bank_file((FIXTURES / "10_gift.txt").read_bytes(), "bank.txt", format_mode="open")
    assert result["stats"]["kinds"] == {"open": 10} and result["key_source"] == {}
    assert result["items"][0]["text"].endswith("\n862 г.\n988 г.\n1054 г.\n1147 г.")


def test_format_test_rejects_a_list_without_options():
    with pytest.raises(BankError) as raised:
        parse_bank_file((FIXTURES / "24_open_list.txt").read_bytes(), "bank.txt", format_mode="test")
    assert (raised.value.status, raised.value.code, raised.value.detail) == (422, "no_items", bank_parse.NO_TEST_ITEMS_DETAIL)


# ── Входные файлы, пределы, безопасность ────────────────────────────────────────


@pytest.mark.parametrize(("name", "contents", "status", "code"), [
    ("Банк.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, 415, "old_doc"),
    ("Банк.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, 415, "old_doc"),
    ("Банк.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + "EncryptedPackage".encode("utf-16-le"), 422, "unreadable"),
    ("Банк.pdf", b"%PDF-1.4", 415, "unsupported"),
    ("Банк.docx", b"%PDF-1.4 renamed", 415, "unsupported"),
    ("Банк.xls", b"anything", 415, "unsupported"),
    ("Банк.csv", b"1;2;3", 415, "unsupported"),
    ("Банк", "1. Вопрос?".encode(), 415, "unsupported"),
    ("Банк.docx", b"not a zip at all", 422, "unreadable"),
    ("Банк.xlsx", b"PK\x03\x04 broken", 422, "unreadable"),
    ("Банк.txt", b"", 422, "empty_file"),
    ("Банк.txt", "Просто текст без вопросов.".encode(), 422, "no_items"),
])
def test_file_kinds_and_broken_files(name, contents, status, code):
    with pytest.raises(BankError) as raised:
        parse_bank_file(contents, name)
    assert (raised.value.status, raised.value.code) == (status, code)


def test_texts_in_windows_1251_and_markdown_files_are_read():
    contents = "1. Сколько дней в неделе?\nа) 5\nб) 7\nОтвет: б\n".encode("cp1251")
    result = parse_bank_file(contents, "Тест.txt")
    assert result["items"][0]["options"] == ["5", "7"] and result["items"][0]["correct"] == [1]
    markdown = parse_bank_file("1. Вопрос?\nа) **да**\nб) нет\n".encode(), "Тест.md")
    assert markdown["items"][0]["correct"] == [0] and markdown["key_source"] == {"bold": 1}


def test_a_word_table_pasted_as_text_is_read_as_a_table():
    rows = ["№\tСодержание вопроса\tВарианты ответов", "Тема 1. Директивы\t\t"]
    for number in (1, 2):
        rows += [f"{number}\tКакая директива номер {number}?\t#region", "\t\t#define", "\t\t#if"]
    result = parse_bank_text("\n".join(rows))
    assert result["layout"] == "table_option_rows" and [t["title"] for t in result["topics"]] == ["Тема 1. Директивы"]
    assert [item["options"] for item in result["items"]] == [["#region", "#define", "#if"]] * 2


def test_more_than_2000_items_or_2_000_000_characters_are_refused():
    with pytest.raises(BankError) as many:
        parse_bank_text("\n".join(f"{n}. Вопрос {n}?" for n in range(1, 2002)))
    assert (many.value.status, many.value.code, many.value.detail) == (
        413, "too_many", "В файле больше 2000 вопросов. Разделите его на части и загрузите по очереди.")
    full = parse_bank_text("\n".join(f"{n}. Вопрос {n}?" for n in range(1, 2001)))
    assert full["stats"]["items"] == 2000 and full["items"][-1]["label"] == "2000" and full["stats"]["gaps"] == []
    year = parse_bank_text("1. В каком году Грюнвальдская битва?\nа) 1385.\n1410. год битвы\nб) 1569.")
    assert year["stats"]["items"] == 1  # «1410.» не продолжает нумерацию: не вопрос
    with pytest.raises(BankError) as long:
        parse_bank_text("а" * 2_000_001)
    assert (long.value.status, long.value.code) == (413, "too_much_text")


def test_long_runs_of_spaces_and_lines_do_not_slow_the_parser_down():
    """Разбор линейный: «[ \\t]+\\n» (parser.clean_extracted_text) и поиск маркера по всей строке были квадратичными."""
    spaced = "1. Вопрос?\nа) x" + " " * 1_000_000 + "+\nб) z (верно)\nв) **y**"
    assert best_seconds(lambda: parse_bank_text(spaced), repeat=1) < 2
    assert best_seconds(lambda: parse_bank_file(spaced.encode(), "bank.txt"), repeat=1) < 2
    wrapped = "1. Вопрос?\nа) начало\n" + "\n".join("хвост" for _ in range(50_000)) + "\nб) конец"
    result = parse_bank_text(wrapped)
    assert len(result["items"][0]["options"][0]) <= 600 and "too_long" in result["items"][0]["flags"]
    with pytest.raises(BankError) as many:
        parse_bank_text("\n".join(f"{n % 9 + 1}. Вопрос?\nа) да\nб) нет\nЗаголовок {n}" for n in range(9000)))
    assert many.value.code == "too_many"


def test_archives_over_the_unpacked_limits_are_not_unpacked(monkeypatch):
    contents = (FIXTURES / "01_table_rows_bold.docx").read_bytes()
    monkeypatch.setattr(bank_docx, "MAX_XML_PART_BYTES", 10_000)
    with pytest.raises(BankError) as part:
        parse_bank_file(contents, "bank.docx")
    assert (part.value.status, part.value.code) == (413, "too_much_text")
    monkeypatch.setattr(bank_docx, "MAX_XML_PART_BYTES", 1_000)
    with pytest.raises(BankError) as xlsx:
        parse_bank_file((FIXTURES / "11_columns_key.xlsx").read_bytes(), "bank.xlsx")
    assert (xlsx.value.status, xlsx.value.code) == (413, "too_much_text")
    monkeypatch.setattr(bank_docx, "MAX_XML_PART_BYTES", 40 * 1024 * 1024)
    monkeypatch.setattr(bank_docx, "MAX_ARCHIVE_BYTES", 50_000)
    with pytest.raises(bank_docx.ArchiveTooLarge):
        bank_docx.check_archive(contents)


def test_a_zip_bomb_is_refused_before_it_is_unpacked():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as bomb:
        bomb.writestr("[Content_Types].xml", "<Types/>")
        bomb.writestr("word/document.xml", b"<" + b" " * (41 * 1024 * 1024))  # 41 МБ пробелов сжимаются в ~40 КБ
    contents = out.getvalue()
    assert len(contents) < 200_000
    with pytest.raises(BankError) as raised:
        parse_bank_file(contents, "bank.docx")
    assert (raised.value.status, raised.value.code) == (413, "too_much_text")


BILLION_LAUGHS = (
    '<!DOCTYPE w:document [<!ENTITY a "РАЗДУТО">'
    + "".join(f'<!ENTITY {chr(98 + i)} "{("&" + chr(97 + i) + ";") * 10}">' for i in range(8))
    + '<!ENTITY ext SYSTEM "file:///etc/hostname">]>'
)


def test_xml_entities_in_a_docx_are_never_expanded():
    body = ('<w:p><w:r><w:t>1. Вопрос &i; &ext;?</w:t></w:r></w:p>'
            + xml_paragraph("а) да", bold=True) + xml_paragraph("б) нет"))
    started = time.perf_counter()
    try:
        result = parse_bank_file(docx_with_body(body, prolog=BILLION_LAUGHS), "bank.docx")
    except BankError as exc:  # отказ — тоже безопасный исход
        assert exc.code == "unreadable"
    else:
        dumped = json.dumps(result, ensure_ascii=False)
        assert "РАЗДУТО" not in dumped and len(dumped) < 20_000
    assert time.perf_counter() - started < 2


def test_a_declared_entity_is_left_unexpanded_and_the_file_still_parses():
    body = ('<w:p><w:r><w:t>1. Вопрос &boom;?</w:t></w:r></w:p>'
            + xml_paragraph("а) да", bold=True) + xml_paragraph("б) нет"))
    result = parse_bank_file(docx_with_body(body, prolog='<!DOCTYPE w:document [<!ENTITY boom "РАЗДУТО">]>'), "bank.docx")
    assert result["items"][0]["text"] == "Вопрос" and result["items"][0]["correct"] == [0]


def test_xml_entities_in_an_xlsx_are_never_expanded():
    source = (FIXTURES / "11_columns_key.xlsx").read_bytes()
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as book, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in book.infolist():
            data = book.read(info.filename)
            if info.filename == "xl/sharedStrings.xml":
                text = data.decode("utf-8")
                head, rest = text.split("?>", 1)
                data = (head + "?>" + BILLION_LAUGHS.replace("w:document", "sst")
                        + rest.replace("<t>Вопрос</t>", "<t>Вопрос &i; &ext;</t>", 1)).encode("utf-8")
            target.writestr(info, data)
    try:
        result = parse_bank_file(out.getvalue(), "bank.xlsx")
    except BankError as exc:
        assert exc.code == "unreadable"
    else:
        assert "РАЗДУТО" not in json.dumps(result, ensure_ascii=False)


def test_xlsx_cells_keep_numbers_line_breaks_and_formula_strings():
    import xlsxwriter

    buffer = io.BytesIO()
    book = xlsxwriter.Workbook(buffer, {"in_memory": True})
    sheet = book.add_worksheet()
    bold = book.add_format({"bold": True})
    red = book.add_format({"font_color": "#FF0000"})
    sheet.write_row(0, 0, ["№", "Вопрос", "Варианты ответов"])
    sheet.write_number(1, 0, 1)
    sheet.write_string(1, 1, "Сколько будет 0 + 0?")
    sheet.write_rich_string(1, 2, "а) 0\n", red, "б) 1", "\nв) 2")  # варианты в одной ячейке через Alt+Enter
    sheet.write_number(2, 0, 2)
    sheet.write_string(2, 1, "Какой оператор вычитает и сразу присваивает результат?")
    sheet.write_rich_string(2, 2, "а) += ", red, "б) -=", " в) =-")  # варианты в строку, ключ красным
    sheet.write_formula(3, 1, '="Вопрос "&"из формулы?"', None, "Вопрос из формулы?")
    sheet.write_number(3, 0, 3)
    sheet.write_string(3, 2, "а) да\nб) нет", bold)
    book.close()
    result = parse_bank_file(buffer.getvalue(), "bank.xlsx")
    items = result["items"]
    assert [item["options"] for item in items] == [["0", "1", "2"], ["+=", "-=", "=-"], ["да", "нет"]]
    assert items[2]["text"] == "Вопрос из формулы?" and items[0]["where"] == "лист 1, строка 2"
    assert result["format_signal"] == "color" and [item["correct"] for item in items] == [[1], [1], []]
    assert items[2]["flags"] == ["no_key"] and items[2]["marks"]["bold"] == [1, 1]  # жирное всё — не ключ


# ── Ошибки, найденные на синтетических банках ───────────────────────────────────


def owner_like_table(questions: list[tuple[str, list[str], set[int]]], code: dict[int, list[str]] | None = None) -> str:
    """Таблица «№ | Содержание вопроса | Варианты ответов», вариант на строку без букв, ключ жирным; code[k] —
    строки кода под вопросом k отдельными абзацами Consolas ('' — пустой абзац)."""
    def cell(paragraphs: str, props: str = "") -> str:
        return xml_cell(paragraphs, props)

    def mono(text: str) -> str:
        if not text:
            return "<w:p/>"
        return (f'<w:p><w:r><w:rPr><w:rFonts w:ascii="Consolas" w:hAnsi="Consolas"/></w:rPr>'
                f'<w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>')

    rows = ["<w:tr>" + "".join(cell(xml_paragraph(t)) for t in ("№", "Содержание вопроса", "Варианты ответов")) + "</w:tr>"]
    for k, (stem, options, keys) in enumerate(questions, start=1):
        for i, option in enumerate(options):
            merge = '<w:vMerge w:val="restart"/>' if i == 0 else "<w:vMerge/>"
            stem_cell = xml_paragraph(stem) + "".join(mono(line) for line in (code or {}).get(k, [])) if i == 0 else ""
            rows.append("<w:tr>" + cell(xml_paragraph(str(k)) if i == 0 else "", merge) + cell(stem_cell, merge)
                        + cell(xml_paragraph(option, bold=i in keys)) + "</w:tr>")
    return "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/><w:gridCol/></w:tblGrid>" + "".join(rows) + "</w:tbl>"


# (1) Слова «по порядку», «последовательности», «соответствия» вида вопроса не меняют: вид — по устройству.

KIND_WORDS = """1. Какой метод LINQ проецирует каждый элемент последовательности в новую форму?
а) OrderBy
б) **Select**
в) Where
2. Что вернёт вызов делегата, в который по порядку добавлены методы, возвращающие 1, 2 и 3?
а) 1
б) **3**
в) 6
3. Какое соединение вернёт строки, для которых нет соответствия в другой таблице?
а) INNER JOIN
б) **FULL OUTER JOIN**
в) CROSS JOIN
"""


@pytest.mark.parametrize("variant", ["bold", "key_list", "marker", "no_key"])
def test_order_and_match_words_in_a_question_with_options_do_not_change_its_kind(variant):
    text = {
        "bold": KIND_WORDS,
        "key_list": KIND_WORDS.replace("**", "") + "Ключ: 1 – б, 2 – б, 3 – б",
        "marker": "\n".join(("+ " if "**" in line else "- ") + line[3:].strip("*") if line[:2] in ("а)", "б)", "в)") else line
                            for line in KIND_WORDS.splitlines()),
        "no_key": KIND_WORDS.replace("**", ""),
    }[variant]
    result = parse_bank_text(text)
    assert [item["kind"] for item in result["items"]] == ["single"] * 3
    assert [item["options"] for item in result["items"]][2] == ["INNER JOIN", "FULL OUTER JOIN", "CROSS JOIN"]
    if variant == "no_key":
        assert all(item["correct"] == [] and item["flags"] == ["no_key"] for item in result["items"])
    else:
        assert [item["correct"] for item in result["items"]] == [[1], [1], [1]]
        assert all(item["flags"] == [] for item in result["items"])


def test_order_and_match_come_from_the_task_and_the_key_not_from_single_words():
    stages = "а) компоновка\nб) лексический анализ\nв) синтаксический анализ\n"
    sizes = "а) 1 байт\nб) 4 байта\nв) 8 байт\n"
    text = (f"1. Расположите этапы компиляции в правильном порядке:\n{stages}"
            f"2. Установите соответствие между типом и размером: А. int Б. long В. byte\n{sizes}"
            f"3. Расположите этапы компиляции в правильном порядке:\n{stages}Ответ: б, в, а\n"
            f"4. Установите соответствие между типом и размером: А. int Б. long В. byte\n{sizes}Ответ: А-2, Б-3, В-1\n"
            f"5. В каком порядке выполняются этапы компиляции?\n{stages}Ответ: б\n")
    items = parse_bank_text(text)["items"]
    assert [item["kind"] for item in items] == ["order", "match", "order", "match", "single"]
    assert items[0]["flags"] == ["no_key"] and items[1]["flags"] == []  # задание без ключа: вид по заданию
    assert items[2]["order"] == [1, 2, 0] and items[2]["evidence"] == "answer_line"  # ключ-перестановка
    assert items[3]["answer"] == "А-2, Б-3, В-1" and items[3]["correct"] == []  # ключ-пары
    assert items[4]["correct"] == [1]  # вопрос с «?» и ключ-буква — выбор
    bold = stages.replace("б) лексический анализ", "б) **лексический анализ**")
    marked = parse_bank_text(f"1. Расположите этапы компиляции в правильном порядке:\n{bold}2. Какой этап первый?\n{bold}")
    assert [(item["kind"], item["correct"]) for item in marked["items"]] == [("single", [1])] * 2  # отмеченный вариант — выбор


def test_whole_words_only_for_kind_hints():
    assert bank_parse.OPEN_HINT_RE.search("Дополните предложение: цикл с ___ выполняется хотя бы раз.")
    assert not bank_parse.OPEN_HINT_RE.search("Какая дополнительная память нужна сортировке слиянием?")
    tested = parse_bank_text("1. Сколько дополнительной памяти нужно сортировке слиянием?\n2. Вопрос?\nа) да\nб) нет\n",
                             format_mode="test")
    assert tested["items"][0]["kind"] == "unknown" and tested["items"][0]["flags"] == ["no_options"]
    assert not bank_parse.order_task("Что проверяет контрольная последовательность FCS?")
    assert not bank_parse.match_task("Какой командой посмотреть таблицу соответствия IP- и MAC-адресов?")
    assert bank_parse.order_task("Установите правильную последовательность этапов компиляции:")
    assert bank_parse.match_task("Соотнесите слово с переводом:")


# (2) «3.5», адрес, «c.Count» — текст варианта, а не подпись; подпись срезается, только если подписи идут по порядку.


def test_an_option_that_starts_like_a_number_or_a_letter_label_keeps_its_text():
    questions = [
        ("Что будет выведено при делении двух переменных типа int?", ["3,5", "4", "3", "3,0", "3.5"], {2}),
        ("Какой адрес можно назначить узлу сети 10.1.1.192/28?", ["10.1.1.200", "10.1.1.0", "0.0.0.0"], {0}),
        ("Какое свойство вернёт число элементов списка?", ["c.Count", "a.Length", "d.Size()"], {0}),
        ("Сколько будет 2 + 3?", ["5) пять", "4", "3"], {0}),
        ("Пронумерованные варианты?", ["1) один", "2) два", "3) три"], {1}),
    ]
    items = parse_bank_file(docx_with_body(owner_like_table(questions)), "bank.docx")["items"]
    assert [item["options"] for item in items] == [
        ["3,5", "4", "3", "3,0", "3.5"], ["10.1.1.200", "10.1.1.0", "0.0.0.0"], ["c.Count", "a.Length", "d.Size()"],
        ["5) пять", "4", "3"],  # «5)» не первым — не подпись
        ["один", "два", "три"],  # 1), 2), 3) по порядку — подписи
    ]
    assert [item["correct"] for item in items] == [[2], [0], [0], [0], [1]]
    assert bank_parse.labels_in_sequence(["А", "Б", "C"]) and bank_parse.labels_in_sequence(["3", "1", "2"])
    assert not bank_parse.labels_in_sequence(["3"]) and not bank_parse.labels_in_sequence(["а", "", "в"])


# (3) «;» в конце каждого варианта — часть оператора: SQL и C# остаются как написаны.


def test_a_semicolon_at_the_end_of_every_option_is_kept():
    text = ("1. Какой запрос выберет товары с ценой от 100 до 500?\n"
            "а) SELECT * FROM Goods WHERE price IN (100, 500);\nб) **SELECT * FROM Goods WHERE price BETWEEN 100 AND 500;**\n"
            "2. Какое объявление создаёт массив из 10 целых чисел?\n"
            "а) int arr[10];\nб) int[10] arr;\nв) **int[] arr = new int[10];**\n"
            "3. Перечислите виды инфляции:\nа) ползучая;\nб) **галопирующая;**\nв) гиперинфляция.\n")
    items = parse_bank_text(text)["items"]
    assert items[0]["options"][1] == "SELECT * FROM Goods WHERE price BETWEEN 100 AND 500;"
    assert items[1]["options"] == ["int arr[10];", "int[10] arr;", "int[] arr = new int[10];"]
    assert items[2]["options"] == ["ползучая", "галопирующая", "гиперинфляция"]  # «x; y; z.» — знаки списка
    assert [item["correct"] for item in items] == [[1], [2], [1]]


# (4) Список ключей одной строкой, в несколько строк и таблицей.

THREE_QUESTIONS = ("1. Первый вопрос?\nа) один\nб) два\nв) три\n2. Второй вопрос?\nа) один\nб) два\nв) три\n"
                   "3. Третий вопрос — выберите все верные ответы?\nа) один\nб) два\nв) три\n")


@pytest.mark.parametrize("key", [
    "Ключ: 1 – б, 2 – а, 3 – а, в",
    "Ответы: 1-б; 2-а; 3-а, в",
    "Ключ к тесту — 1б 2а 3ав",
    "1б 2а 3ав",
    "Ключ: 1 – б, 2 – а,\n3 – а, в",
    "Правильные ответы:\n1 – б, 2 – а, 3 – а, в",
])
def test_a_key_list_on_one_line_or_several_lines_is_read(key):
    result = parse_bank_text(THREE_QUESTIONS + "\n" + key)
    assert [item["correct"] for item in result["items"]] == [[1], [0], [0, 2]]
    assert result["key_source"] == {"key_list": 3} and result["stats"]["key_lists"] == 1
    assert [item["kind"] for item in result["items"]] == ["single", "single", "multiple"]
    assert all(item["flags"] == [] for item in result["items"])


def test_a_key_list_in_a_docx_paragraph_or_table_is_read():
    def build(paragraph_key=None, rows=None):
        def doc_builder(doc):
            for line in THREE_QUESTIONS.splitlines():
                doc.add_paragraph(line)
            if paragraph_key:
                doc.add_paragraph(paragraph_key)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        table.cell(r, c).text = value
        return docx_bytes(doc_builder)

    variants = {
        "paragraph": build(paragraph_key="Ключ: 1 – б, 2 – а, 3 – а, в"),
        "across": build(rows=[["№ вопроса", "1", "2", "3"], ["Ответ", "б", "а", "а, в"]]),
        "pairs": build(rows=[["№", "Ответ", "№", "Ответ"], ["1", "б", "2", "а"], ["3", "а, в", "", ""]]),
        "cell": build(rows=[["Ключ: 1 – б, 2 – а, 3 – а, в"]]),
    }
    for name, contents in variants.items():
        result = parse_bank_file(contents, "bank.docx")
        assert [item["correct"] for item in result["items"]] == [[1], [0], [0, 2]], name
        assert result["stats"]["items"] == 3 and result["key_source"] == {"key_list": 3}, name


def test_an_answer_line_is_still_the_answer_of_its_own_question():
    result = parse_bank_text("1. Вопрос?\nа) да\nб) нет\nОтвет: б\n2. Соотнесите слово с переводом:\nа) окно\nб) стена\n"
                             "Ответ: 1-б, 2-а\n")
    first, second = result["items"]
    assert first["correct"] == [1] and first["evidence"] == "answer_line"
    assert second["kind"] == "match" and second["answer"] == "1-б, 2-а" and result["stats"]["key_lists"] == 0


# (5) Код в TXT с «+»/«-»: варианты — только отмеченные строки, строки до них — текст вопроса как написан.

PLUS_CODE = (
    "25. Сколько раз будет выведено слово hello?\r\n#include <stdio.h>\r\n#include <unistd.h>\r\nint main(void) {\r\n"
    "    fork();\r\n\r\n    fork();\r\n    printf(\"hello\\n\");\r\n}\r\n- 2\r\n+ 4\r\n- 8\r\n\r\n"
    "26. Может ли группа staff изменить файл?\r\n$ ls -l report.txt\r\n-rw-r----- 1 anna staff 2048 report.txt\r\n"
    "- Да\r\n+ Нет, только чтение\r\n\r\n"
    "27. Что выведет сценарий без аргументов?\r\ncase \"$1\" in\r\n    start) echo \"Запуск\" ;;\r\n"
    "    *)     echo \"Использование: $0 start\" ;;\r\nesac\r\n+ Использование: ./service.sh start\r\n- Запуск\r\n"
)


def test_code_lines_between_the_question_and_its_marked_options_belong_to_the_question():
    result = parse_bank_text(PLUS_CODE)
    assert result["marker"] == "plus_minus" and [item["label"] for item in result["items"]] == ["25", "26", "27"]
    first, second, third = result["items"]
    assert first["text"] == ("Сколько раз будет выведено слово hello?\n#include <stdio.h>\n#include <unistd.h>\n"
                             "int main(void) {\n    fork();\n\n    fork();\n    printf(\"hello\\n\");\n}")
    assert first["options"] == ["2", "4", "8"] and first["correct"] == [1]
    assert second["text"].endswith("\n$ ls -l report.txt\n-rw-r----- 1 anna staff 2048 report.txt")
    assert second["options"] == ["Да", "Нет, только чтение"] and second["correct"] == [1]
    assert third["text"].splitlines()[2:4] == ["    start) echo \"Запуск\" ;;", "    *)     echo \"Использование: $0 start\" ;;"]
    assert third["options"] == ["Использование: ./service.sh start", "Запуск"] and third["correct"] == [0]
    assert all(item["flags"] == [] for item in result["items"])
    # вопросы без номеров через пустую строку (MyTestX): пустая строка внутри кода не делит вопрос надвое
    blocks = parse_bank_text("Что выведет код?\nint x = 1;\n\nConsole.WriteLine(x);\n+ 1\n- 0\n\n"
                             "Какой оператор прибавляет и присваивает?\n+ +=\n- -=\n\nТретий вопрос?\n+ да\n- нет\n")
    assert blocks["layout"] == "blocks" and blocks["stats"]["items"] == 3
    assert blocks["items"][0]["text"] == "Что выведет код?\nint x = 1;\n\nConsole.WriteLine(x);"
    assert [item["options"] for item in blocks["items"]] == [["1", "0"], ["+=", "-="], ["да", "нет"]]


def test_code_with_braces_and_double_colons_does_not_make_a_text_gift():
    csharp = ("1. Какое значение value выведет фрагмент?\nstatic void Change(int x)\n{\n    x = 100;\n}\n\n"
              "int value = 5;\nChange(value);\nConsole.WriteLine(global::System.Math.Abs(value));\n+ 5\n- 100\n\n"
              "2. Что выведет сценарий?\necho ${#names[@]} ${HOME}\n+ 3 /home/anna\n- 0\n")
    result = parse_bank_text(csharp)
    assert result["layout"] == "numbered_paragraphs" and result["marker"] == "plus_minus"
    assert result["items"][0]["text"].endswith("{\n    x = 100;\n}\n\nint value = 5;\nChange(value);"
                                               "\nConsole.WriteLine(global::System.Math.Abs(value));")
    assert [item["correct"] for item in result["items"]] == [[0], [0]]
    assert bank_parse.looks_like_gift("::Q1:: Сколько? {=4 ~5}") and bank_parse.looks_like_gift("Да? {T}\nНет? {F}")
    assert bank_parse.looks_like_gift("Год? {#1410}\nГород? {=Минск ~Брест}")
    assert not bank_parse.looks_like_gift("int[] a = { 1, 2 };\nif (x) { y = 1; }\nstd::cout << List{T};")


# (6) Пустая строка внутри кода (код абзацами, пустой абзац между ними) остаётся.


def test_blank_lines_inside_code_written_one_paragraph_per_line_are_kept():
    code = {1: ["static void Change(int x)", "{", "    x = 100;", "}", "", "int value = 5;", "Change(value);", ""]}
    table = owner_like_table([("Что выведет фрагмент?", ["5", "100"], {0})], code)
    item = parse_bank_file(docx_with_body(table), "bank.docx")["items"][0]
    assert item["text"] == ("Что выведет фрагмент?\nstatic void Change(int x)\n{\n    x = 100;\n}\n\n"
                            "int value = 5;\nChange(value);")  # пустой абзац после кода — не строка кода
    assert item["options"] == ["5", "100"] and item["correct"] == [0]

    def build(doc):
        doc.add_paragraph("1. Что выведет фрагмент?")
        for line in ("int a = 1;", "", "Console.WriteLine(a);"):
            run = doc.add_paragraph().add_run(line)
            run.font.name = "Courier New"
        doc.add_paragraph("а) 1")
        doc.add_paragraph("б) 0")
        doc.add_paragraph("Ключ: 1 – а, 2 – б, 3 – в")
    typed = parse_bank_file(docx_bytes(build), "bank.docx")["items"][0]
    assert typed["text"] == "Что выведет фрагмент?\nint a = 1;\n\nConsole.WriteLine(a);" and typed["correct"] == [0]


# (7) Повторы вариантов — с учётом регистра (и кода).


def test_same_options_are_case_sensitive():
    result = parse_bank_text("1. Что выведет фрагмент?\nа) **Hello**\nб) hello\nв) HELLO\n"
                             "2. Одинаковые варианты?\nа) **да**\nб) да\nв) нет\n")
    assert [item["flags"] for item in result["items"]] == [[], ["same_options"]]


# (8) «#include», «#!/bin/bash» — строки кода, а не маркер «#».


def test_hash_lines_of_code_are_not_answer_markers():
    def build(doc):
        doc.add_paragraph("25. Сколько раз будет выведено слово hello?")
        for line in ("#include <stdio.h>", "#define N 2", "int main(void) {", "    fork();", "}"):
            doc.add_paragraph().add_run(line).font.name = "Courier New"
        for letter, option in zip("абв", ("2", "4", "8")):
            run = doc.add_paragraph().add_run(f"{letter}) {option}")
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW if option == "4" else None
        doc.add_paragraph("26. Что выведет сценарий?")
        for line in ("#!/bin/bash", "echo \"Привет, $1\""):
            doc.add_paragraph().add_run(line).font.name = "Courier New"
        for letter, option in zip("абв", ("Привет, Анна", "Привет, $1", "Ошибка")):
            run = doc.add_paragraph().add_run(f"{letter}) {option}")
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW if letter == "а" else None
    result = parse_bank_file(docx_bytes(build), "bank.docx")
    first, second = result["items"]
    assert first["text"].splitlines()[1:3] == ["#include <stdio.h>", "#define N 2"] and first["options"] == ["2", "4", "8"]
    assert second["text"].splitlines()[1] == "#!/bin/bash" and len(second["options"]) == 3
    assert [item["correct"] for item in result["items"]] == [[1], [0]] and result["marker"] is None
    assert not bank_parse.GLUED_MARK_RE.match("#include <stdio.h>") and bank_parse.GLUED_MARK_RE.match("#Полоцк")


# (9) Строка таблицы маршрутизации или вывода с числом в начале — текст, а не вопрос.


def test_numbers_that_do_not_continue_the_numbering_are_text():
    text = ("124. Вопрос перед таблицей?\nа) **да**\nб) нет\n"
            "125. Через какой шлюз будет отправлен пакет для адреса 10.20.5.7?\n"
            "Сеть           Маска            Шлюз\n0.0.0.0        0.0.0.0          192.168.1.1\n"
            "10.20.0.0      255.255.0.0      192.168.1.3\n192) строка вывода\n"
            "а) 192.168.1.1\nб) **192.168.1.3**\n126. Вопрос после таблицы?\nа) да\nб) **нет**\n")
    result = parse_bank_text(text)
    assert [item["label"] for item in result["items"]] == ["124", "125", "126"]
    table = result["items"][1]
    assert table["text"].splitlines()[1:] == ["Сеть           Маска            Шлюз", "0.0.0.0        0.0.0.0          192.168.1.1",
                                              "10.20.0.0      255.255.0.0      192.168.1.3", "192) строка вывода"]
    assert table["options"] == ["192.168.1.1", "192.168.1.3"] and table["correct"] == [1]
    assert result["stats"]["gaps"] == [] and all(item["flags"] == [] for item in result["items"])


def test_topic_question_numbers_and_numbered_options_under_code():
    topic = parse_bank_text("Тема 1\n1.1. Что такое алгоритм?\nа) **набор шагов**\nб) число\n1.2. Что такое программа?\nx = 1\n"
                            "а) **текст**\nб) число\n2.1. Что такое цикл?\nа) **повтор**\nб) число\n")
    assert [(item["label"], item["text"], item["correct"]) for item in topic["items"]] == [
        ("1.1", "Что такое алгоритм?", [0]), ("1.2", "Что такое программа?\nx = 1", [0]), ("2.1", "Что такое цикл?", [0])]
    assert topic["stats"]["gaps"] == [] and [s["evidence"] for s in topic["subjects"]] == ["none"]
    numbered = parse_bank_text("1. Что выведет код?\nx = 1\n1) 1\n2) 0\nОтвет: 1\n2. Второй вопрос?\n1) да\n2) нет\nОтвет: 2\n")
    assert [(item["label"], item["text"], item["options"], item["correct"]) for item in numbered["items"]] == [
        ("1", "Что выведет код?\nx = 1", ["1", "0"], [0]), ("2", "Второй вопрос?", ["да", "нет"], [1])]


# (10) Открытые вопросы в TXT: код под вопросом — его текст, строки и отступы как написаны.

OPEN_CODE = """Вопросы к экзамену по дисциплине «Основы алгоритмизации»
Раздел 1. Циклы
1. Выполните трассировку приведённого фрагмента для n = 4.
s := 0
нц для i от 1 до n
    s := s + i * i
кц
2. Найдите ошибку в приведённом фрагменте:
i = 0
while i < 10:
    print(i)
i += 1
3. Какую последовательность вызовов создаст f(4)?
def f(n):
    if n <= 1:
        return 1
    return n * f(n - 1)
4. Объясните разницу между += и -= при подписке
на событие и при отписке от него.
"""


def test_open_questions_in_text_keep_their_code_line_by_line():
    result = parse_bank_text(OPEN_CODE)
    assert result["stats"]["kinds"] == {"open": 4} and all(item["options"] == [] for item in result["items"])
    texts = [item["text"] for item in result["items"]]
    assert texts[0] == "Выполните трассировку приведённого фрагмента для n = 4.\ns := 0\nнц для i от 1 до n\n    s := s + i * i\nкц"
    assert texts[1] == "Найдите ошибку в приведённом фрагменте:\ni = 0\nwhile i < 10:\n    print(i)\ni += 1"
    assert texts[2].splitlines()[1:] == ["def f(n):", "    if n <= 1:", "        return 1", "    return n * f(n - 1)"]
    assert texts[3] == "Объясните разницу между += и -= при подписке на событие и при отписке от него."  # перенос строки
    assert [(s["title"], s["evidence"]) for s in result["subjects"]] == [("Основы алгоритмизации", "heading")]
    # format «test»: строки под вопросом с «?» — варианты; код после «.» и «:» — всё равно строками
    tested = parse_bank_text(OPEN_CODE, format_mode="test")["items"]
    assert tested[0]["text"] == texts[0] and tested[2]["options"][:2] == ["def f(n):", "if n <= 1:"]


# (11) Флаг «code» (проверка удобства 2026-10-10, m1): моноширинный шрифт самого вопроса и строки SQL и оболочки.
# Раньше код узнавали только по знакам C в вопросе и по шрифту вариантов: 90 из 203 вопросов с SQL и командами
# показывались шрифтом заголовка.

# Код, которого флаг не видит намеренно: таблица в TXT (без шрифта, без знаков кода) и Python в открытых вопросах TXT.
CODE_FLAG_LIMITS = {
    ("db_plus.txt", "db", 43),
    ("four_subjects_slice.txt", "networks", 125),
    *(("algo_open.txt", "algo", n) for n in (11, 18, 21, 53, 62, 66)),
}


def test_every_code_question_of_the_synthetic_banks_is_flagged_as_code():
    misses = set()
    for entry in TRUTH["files"]:
        result = parse_bank_file((SYNTHETIC / entry["file"]).read_bytes(), entry["file"])
        rows = []
        for part in entry["parts"]:
            by_number = {item["n"]: item for item in TRUTH["subjects"][part["subject"]]["items"]}
            rows += [(part["subject"], by_number[n]) for n in (part["pick"] or sorted(by_number))]
        assert len(rows) == len(result["items"])
        for (subject, want), got in zip(rows, result["items"]):
            if want["code"] and not got["code"]:
                misses.add((entry["file"], subject, want["n"]))
    assert misses == CODE_FLAG_LIMITS


def test_no_line_of_a_question_without_code_reads_as_sql_or_shell():
    for subject in TRUTH["subjects"].values():
        for item in subject["items"]:
            if not item["code"]:
                assert not bank_parse.code_in_lines(item["text"]), (subject["subject"], item["n"], item["text"])


@pytest.mark.parametrize("line", [
    "SELECT *", "FROM Students", "WHERE course = 2;", "ORDER BY name DESC;", "GROUP BY group_id", "LEFT JOIN Groups g ON s.group_id = g.id;",
    "CREATE TABLE Orders (", "CREATE UNIQUE INDEX ux_users_email", "ALTER TABLE Students", "DROP TABLE Groups;",
    "INSERT INTO Students (student_id, full_name, group_id)", "VALUES (1, 'Иванов Иван', 99);", "UPDATE Employees",
    "SET TRANSACTION ISOLATION LEVEL REPEATABLE READ;", 'SET session:42 "user=ivanov" EX 3600', "BEGIN;", "ROLLBACK;",
    "select name from students where id = 1", "uname -r", "ps aux | grep nginx", "kill -9 2741", "$ ls -l report.txt",
    "#!/bin/bash", "@echo off", "C:\\>ping 8.8.8.8", "chmod 755 deploy.sh", "sudo chown -R www-data:www-data /var/www/site",
    'grep -q "ERROR" app.log && echo "Есть ошибки" || echo "Ошибок нет"', "docker ps -a", "history | tail -5",
    "Get-Process | Where-Object { $_.CPU -gt 100 }", "echo $?",
])
def test_sql_and_shell_lines_are_code(line):
    assert bank_parse.sql_or_shell_line(line)


@pytest.mark.parametrize("line", [
    "Какой оператор SQL удаляет таблицу?", "Выберите все верные ответы.", "FROM — это часть запроса, верно?",
    "Where is the station?", "Select the correct answer.", "Find the odd one out", "Type of memory: ROM",
    "Set-up of the network", "Read-Only Memory (ROM)", "ON/OFF", "Windows XP", "TCP/IP", "x = 5",
    "ls?", "",
])
def test_words_and_questions_are_not_sql_or_shell(line):
    assert not bank_parse.sql_or_shell_line(line)


def test_the_monospace_of_a_question_flags_code_without_signs_of_code():
    pascal = {1: ["x := 5", "writeln(x + 1)"]}
    table = owner_like_table([("Что выведет программа на Паскале?", ["6", "5"], {0})], pascal)
    item = parse_bank_file(docx_with_body(table), "bank.docx")["items"][0]
    assert item["text"] == "Что выведет программа на Паскале?\nx := 5\nwriteln(x + 1)" and item["code"] is True

    def build(code_font: str | None, inline: bool):
        def make(doc):
            question = doc.add_paragraph("1. Что вернёт ")
            call = question.add_run("len(s)" if inline else "длина строки")
            if inline:
                call.font.name = "Consolas"
            question.add_run(" для пустой строки?")
            for line in ("s = ''", "print(len(s))"):
                run = doc.add_paragraph().add_run(line)
                if code_font:
                    run.font.name = code_font
            doc.add_paragraph("а) 0")
            doc.add_paragraph("б) 1")
            doc.add_paragraph("Ключ: 1 – а")
        return make
    # Код строками под вопросом моноширинным шрифтом; код внутри строки вопроса; без шрифта и без знаков — не код.
    assert parse_bank_file(docx_bytes(build("Courier New", False)), "bank.docx")["items"][0]["code"] is True
    assert parse_bank_file(docx_bytes(build(None, True)), "bank.docx")["items"][0]["code"] is True
    plain = parse_bank_file(docx_bytes(build(None, False)), "bank.docx")["items"][0]
    assert plain["text"].startswith("Что вернёт длина строки") and plain["code"] is False


def test_a_question_line_typed_in_a_monospace_font_is_not_code_by_its_font():
    def build(doc):
        doc.add_paragraph().add_run("1. Сколько дней в неделе?").font.name = "Courier New"
        for text in ("а) 5", "б) 7", "Ключ: 1 – б"):
            doc.add_paragraph(text)
    item = parse_bank_file(docx_bytes(build), "bank.docx")["items"][0]
    assert item["text"] == "Сколько дней в неделе?" and item["options"] == ["5", "7"] and item["code"] is False


def test_sql_and_shell_under_a_question_in_text_flag_code():
    result = parse_bank_text("1. Что вернёт запрос?\nSELECT name\nFROM Students\nа) **имена**\nб) ничего\n"
                             "2. Что делает команда?\nchmod 755 deploy.sh\nа) **права**\nб) удаляет\n"
                             "3. Choose the right form:\nWhere is the station?\nа) **is**\nб) are\n")
    assert [item["code"] for item in result["items"]] == [True, True, False]


# ── Строка вопроса на всю ширину, под ней варианты по ячейкам (владелец, 2026-10-14) ─────────


def stem_table(questions: list[tuple[str, list[str], int]], width: int = 5, bold_stem: bool = False) -> str:
    """Таблица файлов колледжа «Вариант 4»: строка «N.Вопрос» объединена на всю ширину, под ней строка из width
    ячеек «А)…», «Б)…»; верный вариант — жирный (correct = -1: без ключа)."""
    rows = []
    for number, (question, options, correct) in enumerate(questions, start=1):
        stem = "".join(xml_paragraph(line, bold=bold_stem) for line in [f"{number}.{question.splitlines()[0]}",
                                                                       *question.splitlines()[1:]])
        rows.append("<w:tr>" + xml_cell(stem, f'<w:gridSpan w:val="{width}"/>') + "</w:tr>")
        rows.append("<w:tr>" + "".join(xml_cell(xml_paragraph(f"{'АБВГДЕ'[k]}){option}", bold=k == correct))
                                       for k, option in enumerate(options)) + "</w:tr>")
    grid = "<w:gridCol/>" * width
    return f"<w:tbl><w:tblPr/><w:tblGrid>{grid}</w:tblGrid>" + "".join(rows) + "</w:tbl>"


STEM_QUESTIONS = [
    ("Что такое Session в ASP.NET?", ["Серверное состояние пользователя", "Клиентский файл cookie", "База данных",
                                       "HTML-тег", "Тип CSS"], 0),
    ("Дан алгоритм в виде блок-схемы.\nНайти А, В, С, D, если изначально:\nА=0, В=0, C=5, D=10;",
     ["0, 0, 5, 0", "0, 0, 5, 10", "5, 5, 10, 10", "0, 0, 0, 5", "5, 0, 0, 5"], 3),
    ("Что возвращает метод Read() класса XmlReader?", ["true или false", "Строку", "Число", "XmlNode", "Ничего"], 0),
]


def test_a_question_row_across_the_table_takes_the_option_cells_below_it():
    body = xml_paragraph("ТЕСТОВЫЙ ВАРИАНТ 4", bold=True) + stem_table(STEM_QUESTIONS)
    result = parse_bank_file(docx_with_body(body), "variant4.docx")
    items = result["items"]
    assert [(i["label"], i["text"], i["options"], i["correct"]) for i in items] == [
        (str(n), q, opts, [c]) for n, (q, opts, c) in enumerate(STEM_QUESTIONS, start=1)]
    assert result["topics"] == [] and result["key_source"] == {"bold": 3} and result["layout"] == "table_stem_rows"
    assert all(not set(i["flags"]) & bank_parse.PROBLEM_FLAGS for i in items)


def test_a_bold_question_row_without_marked_options_is_a_bank_without_a_key():
    questions = [(q, opts, -1) for q, opts, _ in STEM_QUESTIONS]
    result = parse_bank_file(docx_with_body(xml_paragraph("Вариант 1", bold=True) + stem_table(questions, bold_stem=True)),
                             "variant1.docx")
    assert [(i["text"], i["options"], i["correct"]) for i in result["items"]] == [(q, opts, []) for q, opts, _ in questions]
    assert result["stats"]["no_key"] == 3 and result["topics"] == []


def test_a_numbered_topic_row_above_ordinary_question_rows_stays_a_topic():
    # «1. Введение» на всю ширину над строками «№ | вопрос | вариант» — тема, а не вопрос (под ней нет ряда вариантов)
    def row(cells, merged=False):
        if merged:
            return "<w:tr>" + xml_cell(xml_paragraph(cells, bold=True), '<w:gridSpan w:val="3"/>') + "</w:tr>"
        return "<w:tr>" + "".join(xml_cell(xml_paragraph(text) if text else "") for text in cells) + "</w:tr>"

    rows = [row(("№", "Вопрос", "Варианты ответов")), row("1. Введение в предмет", merged=True)]
    for number in (1, 2):
        rows.append(row((str(number), f"Вопрос номер {number}?", "верный")))
        rows.append(row(("", "", "неверный")))
    table = "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/><w:gridCol/></w:tblGrid>" + "".join(rows) + "</w:tbl>"
    result = parse_bank_file(docx_with_body(table), "bank.docx", key_mode="first_option")
    assert [t["title"] for t in result["topics"]] == ["1. Введение в предмет"]
    assert [(i["text"], i["options"]) for i in result["items"]] == [
        ("Вопрос номер 1?", ["верный", "неверный"]), ("Вопрос номер 2?", ["верный", "неверный"])]
