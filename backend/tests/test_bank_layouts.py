"""Раскладки, которые разбор узнаёт после «дикого» корпуса (fixtures/banks/wild, bank_eval.py): по тесту на раскладку.

Нумерация «№1.», «Q1.», «Question 1.», «Пытанне 1.», римская, «1 – …», «Вопрос 3» строкой над текстом; вопросы без
номеров подряд; варианты строкой под вопросом, по два в строке, «a<tab>…»; отметки перед вариантом («+в)», «*в)», «☑/☐»,
«[x]», «- [x]»); ключ по вариантам в самом файле; колонтитулы и подписи; «(Ответ: б)» в строке вопроса; подсказка
«Выберите один правильный ответ:»; страница попытки Moodle. И что прежнее не сломано: ключ «1 – б» остаётся ключом.
"""
import io
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from bank_parse import parse_bank_file, parse_bank_text  # noqa: E402

OPTIONS = "а) Минск\nб) **Гомель**\nв) Брест\n"


def stems(result):
    return [item["text"] for item in result["items"]]


def keys(result):
    return [item["correct"] for item in result["items"]]


@pytest.mark.parametrize("first, second", [
    ("№1. Первый вопрос?", "№2. Второй вопрос?"),
    ("№ 1 Первый вопрос?", "№ 2 Второй вопрос?"),
    ("Q1. Первый вопрос?", "Q2. Второй вопрос?"),
    ("Question 1. Первый вопрос?", "Question 2. Второй вопрос?"),
    ("Пытанне 1. Первый вопрос?", "Пытанне 2. Второй вопрос?"),
    ("I. Первый вопрос?", "II. Второй вопрос?"),
    ("1 – Первый вопрос?", "2 – Второй вопрос?"),
])
def test_question_numbers_with_words_signs_roman_numerals_and_dashes(first, second):
    result = parse_bank_text(f"{first}\n{OPTIONS}{second}\n{OPTIONS}")
    assert stems(result) == ["Первый вопрос?", "Второй вопрос?"]
    assert [item["label"] for item in result["items"]] == ["1", "2"]
    assert keys(result) == [[1], [1]]


def test_roman_numbers_go_on_past_three():
    text = "".join(f"{numeral}. Вопрос {n}?\n{OPTIONS}" for n, numeral in enumerate(["I", "II", "III", "IV", "V", "VI", "IX"], 1))
    assert [item["label"] for item in parse_bank_text(text)["items"]] == ["1", "2", "3", "4", "5", "6", "9"]


@pytest.mark.parametrize("blank", ["", "\n"])
def test_a_question_number_on_its_own_line_heads_the_text_below_it(blank):
    text = f"Вопрос 1\nПервый вопрос?\n{OPTIONS}{blank}Вопрос №2\nВторой вопрос?\n1) раз\n2) два\n{blank}"
    result = parse_bank_text(text + "Ответы: 1-б, 2-2")
    assert stems(result) == ["Первый вопрос?", "Второй вопрос?"]
    assert [item["options"] for item in result["items"]] == [["Минск", "Гомель", "Брест"], ["раз", "два"]]
    assert keys(result) == [[1], [1]]


def test_questions_without_numbers_follow_each_other_without_blank_lines():
    result = parse_bank_text(f"Первый вопрос?\n{OPTIONS}Второй вопрос?\n{OPTIONS}Третий вопрос?\n{OPTIONS}")
    assert stems(result) == ["Первый вопрос?", "Второй вопрос?", "Третий вопрос?"]
    assert keys(result) == [[1], [1], [1]]


def test_a_number_with_a_dash_is_not_taken_from_a_key_line():
    result = parse_bank_text(f"1. Первый вопрос?\n{OPTIONS}2. Второй вопрос?\n{OPTIONS}\nОтветы:\n1 – а\n2 – в")
    assert stems(result) == ["Первый вопрос?", "Второй вопрос?"]


def test_options_on_one_line_under_the_question_with_semicolons_or_tabs():
    result = parse_bank_text("1. Столица?\nа) Минск; б) Гомель; в) Брест\n2. Река?\nа) Днепр\tб) Сож\tв) Неман\nОтветы: 1-а, 2-в")
    assert [item["options"] for item in result["items"]] == [["Минск", "Гомель", "Брест"], ["Днепр", "Сож", "Неман"]]
    assert keys(result) == [[0], [2]]


def test_options_two_per_line_are_read_in_letter_order():
    text = "1. Сколько?\nа) 180\t\tв) 360\nб) 270\t\tг) 400\n\n2. Где?\nа) тут  в) там\nб) здесь  г) нигде\n\nОтветы: 1в, 2а"
    result = parse_bank_text(text)
    assert [item["options"] for item in result["items"]] == [["180", "270", "360", "400"], ["тут", "здесь", "там", "нигде"]]
    assert keys(result) == [[2], [0]]


def test_a_letter_and_a_tab_label_an_option():
    result = parse_bank_text("1. Capital?\na\tToronto\nb\t**Ottawa**\nc\tMontreal\n2. Planet?\na\tVenus\nb\t**Mercury**\nc\tMars")
    assert [item["options"] for item in result["items"]] == [["Toronto", "Ottawa", "Montreal"], ["Venus", "Mercury", "Mars"]]
    assert keys(result) == [[1], [1]]


@pytest.mark.parametrize("marked, plain", [("+в) ", "{}) "), ("*в) ", "{}) "), ("☑ в) ", "☐ {}) "), ("[x] ", "[ ] "),
                                           ("- [x] ", "- [ ] ")])
def test_a_mark_before_the_option_is_the_key(marked, plain):
    def option(letter, text):
        return (marked if letter == "в" else plain.format(letter)) + text

    lines = []
    for n in (1, 2):
        lines.append(f"{n}. Вопрос {n}?")
        lines += [option(letter, f"вариант {letter}") for letter in "абвг"]
    result = parse_bank_text("\n".join(lines))
    assert [item["options"] for item in result["items"]] == [["вариант а", "вариант б", "вариант в", "вариант г"]] * 2
    assert keys(result) == [[2], [2]]
    assert set(result["key_source"]) == {"checkbox"}


def test_markdown_task_list_options_under_a_question_without_a_question_mark():
    result = parse_bank_file("1. Отметьте хвойные.\n   - [x] сосна\n   - [ ] берёза\n   - [x] ель\n".encode(), "test.md")
    assert result["items"][0]["options"] == ["сосна", "берёза", "ель"]
    assert result["items"][0]["correct"] == [0, 2] and result["items"][0]["kind"] == "multiple"


def test_keys_of_two_variants_in_one_file_go_to_their_own_variant():
    body = "Вариант 1\n1. Вопрос А?\nа) да\nб) нет\n2. Вопрос Б?\nа) да\nб) нет\n\nВариант 2\n1. Вопрос В?\nа) да\nб) нет\n" \
           "2. Вопрос Г?\nа) да\nб) нет\n\nОтветы\nВариант 1: 1а, 2б\nВариант 2: 1б, 2а\n"
    result = parse_bank_text(body)
    assert keys(result) == [[0], [1], [1], [0]]
    assert [(s["title"], s["evidence"]) for s in result["subjects"]] == [("Вариант 1", "restart"), ("Вариант 2", "restart")]
    one_key = body.replace("Вариант 2: 1б, 2а\n", "")
    assert keys(parse_bank_text(one_key)) == [[0], [1], [], []]  # ключ варианта 1 не достаётся варианту 2


def test_a_key_table_with_a_column_per_variant_inside_the_test_file():
    doc = Document()
    for variant in (1, 2):
        doc.add_paragraph(f"Вариант {variant}")
        for n in (1, 2):
            doc.add_paragraph(f"{n}. Вопрос {variant}.{n}?")
            doc.add_paragraph("а) да")
            doc.add_paragraph("б) нет")
    doc.add_paragraph("Ключ")
    table = doc.add_table(rows=3, cols=3)
    for r, row in enumerate([["№", "Вариант 1", "Вариант 2"], ["1", "а", "б"], ["2", "б", "а"]]):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    buffer = io.BytesIO()
    doc.save(buffer)
    assert keys(parse_bank_file(buffer.getvalue(), "bank.docx")) == [[0], [1], [1], [0]]


def test_page_furniture_and_signatures_do_not_stick_to_the_last_option():
    text = ("Тест по естествознанию для 9 класса\n\n1. Вопрос один?\nа) раз\nб) два\n\n2\nТест по естествознанию для 9 класса\n\n"
            "в) три\n2. Вопрос два?\nа) раз\nб) два\n\nСтраница 1 из 3\n\n3. Вопрос три?\nа) раз\nб) два\n\n"
            "Составитель: преподаватель ________ П. П. Петров\n")
    result = parse_bank_text(text)
    assert [item["options"] for item in result["items"]] == [["раз", "два", "три"], ["раз", "два"], ["раз", "два"]]


def test_an_answer_in_brackets_at_the_end_of_the_question_line_is_its_key():
    result = parse_bank_text("1. Кто написал «Бородино»? (Ответ: б)\nа) Пушкин\nб) Лермонтов\n2. Сколько? (правильный ответ — а)\nа) 1\nб) 2")
    assert stems(result) == ["Кто написал «Бородино»?", "Сколько?"]
    assert keys(result) == [[1], [0]] and result["key_source"] == {"answer_line": 2}


def test_choose_one_hint_is_not_part_of_the_question():
    result = parse_bank_text("Задание №1. Какой прибор измеряет давление?\nВыберите один правильный ответ:\nа) термометр\nб) **барометр**")
    assert stems(result) == ["Какой прибор измеряет давление?"] and keys(result) == [[1]]


def test_a_moodle_attempt_review_page_copied_as_text():
    lines = []
    for n, (stem, options, right) in enumerate([("Какой газ преобладает?", ["кислород", "азот"], "азот"),
                                                ("Сколько планет?", ["8", "9"], "8")], start=1):
        lines += [f"Вопрос {n}", "Верно", "Баллов: 1,00 из 1,00", "Отметить вопрос", "Текст вопроса", stem, "Выберите один ответ:"]
        lines += [f"{'abcd'[k]}. {option}" for k, option in enumerate(options)]
        lines += ["Отзыв", f"Правильный ответ: {right}"]
    result = parse_bank_text("\n".join(lines))
    assert stems(result) == ["Какой газ преобладает?", "Сколько планет?"]
    assert [item["options"] for item in result["items"]] == [["кислород", "азот"], ["8", "9"]]
    assert keys(result) == [[1], [0]]
