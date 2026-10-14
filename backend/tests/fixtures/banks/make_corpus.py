"""Синтетический корпус банков тестовых вопросов с эталоном для tests/test_bank_parse.py (без чужих файлов).

Пересобрать (из корня репозитория; нужны python-docx, Pillow и xlsxwriter — последний только для этого скрипта):

    python backend/tests/fixtures/banks/make_corpus.py backend/tests/fixtures/banks

Пишет файлы корпуса и expected.json: у каждого файла вопросы (label, topic, kind, text, options, correct…),
темы и флаги, которые разбор обязан поставить. Содержание придумано; в вопросах по C# — код, который должен
дойти до варианта без изменений («#region», «+=», «-=», «--», «*», «=>», «items?.Count ?? 0», «count -= step;»).
Файлы колледжа сюда не кладутся: дерево выпуска уходит в Space на Hugging Face.
"""
from __future__ import annotations

import copy
import io
import json
import os
import sys

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

RUS = 'абвгдежзик'

# ── item sets ───────────────────────────────────────────────────────────────────


def q(text, options, correct, kind='single', **extra):
    return {'text': text, 'options': options, 'correct': correct, 'kind': kind, **extra}


CSHARP = [
    ('Тема 1. Директивы и пространства имён', [
        q('Какие символы открывают однострочный комментарий в C#?', ['//', '#', '--', '%'], [0]),
        q('Какая директива сворачивает блок кода в именованную область?', ['#define', '#region', '#pragma', '#line'], [1]),
        q('Какое ключевое слово подключает пространство имён?', ['import', 'include', 'using', 'require'], [2]),
        q('Какой оператор подписывает обработчик на событие?', ['=', '->', '<<', '+='], [3]),
    ]),
    ('Тема 2. Делегаты и события', [
        q('Какой оператор отделяет параметры лямбда-выражения от его тела?', ['=>', '->', '::', '*'], [0]),
        q('Что вернёт выражение items?.Count ?? 0, если items равно null?',
          ['0', 'null', 'исключение NullReferenceException', 'пустой список'], [0]),
        q('Какой делегат подходит для метода без возвращаемого значения?', ['Func<int>', 'Action', 'Predicate<int>', 'Converter<int, int>'], [1]),
        q('Что содержит переменная делегата без присвоенного метода?', ['0', 'false', 'null', '"empty"'], [2]),
    ]),
    ('Тема 3. Строки', [
        q('Какой класс используется для частых изменений строки?', ['StringWriter', 'StringBuilder', 'CharBuffer', 'TextArray'], [1]),
        q('Каким символом помечают интерполированную строку?', ['@', '#', '$', '%'], [2]),
        q('Какой метод объединяет массив строк с разделителем?', ['String.Split', 'String.Join', 'String.Merge', 'String.Concat'], [1]),
        q('Что означает префикс @ перед строкой "C:\\temp\\new"?', ['экранирование отключено', 'строка шифруется', 'строка в верхнем регистре', 'строка неизменяема'], [0]),
    ]),
]

HISTORY = [
    ('Раздел 1. Древняя Русь', [
        q('В каком году произошло Крещение Руси?', ['862 г.', '988 г.', '1054 г.', '1147 г.'], [1]),
        q('Кто из князей принял Русскую Правду?', ['Олег', 'Игорь', 'Ярослав Мудрый', 'Святослав'], [2]),
        q('Выберите все города Полоцкой земли:', ['Полоцк', 'Новгород', 'Витебск', 'Киев', 'Минск'], [0, 2, 4], 'multiple'),
    ]),
    ('Раздел 2. ВКЛ', [
        q('В каком году состоялась Грюнвальдская битва?', ['1385 г.', '1410 г.', '1569 г.', '1648 г.'], [1]),
        q('Как назывался свод законов ВКЛ 1588 года?', ['Судебник', 'Статут', 'Уложение', 'Правда'], [1]),
        q('Какое государство образовалось по Люблинской унии?', ['Речь Посполитая', 'Московское царство', 'Орда', 'Ливонский орден'], [0]),
    ]),
]


def flat(sets):
    out = []
    number = 0
    for topic_index, (title, items) in enumerate(sets, start=1):
        for item in items:
            number += 1
            out.append({**copy.deepcopy(item), 'label': str(number), 'topic': topic_index})
    return out


def expected(name, sets, items, flags=None, notes='', layout='', key='', extra=None):
    return {
        'file': name,
        'layout': layout,
        'key': key,
        'topics': [title for title, _ in sets],
        'items': [{k: item[k] for k in ('label', 'topic', 'subject', 'kind', 'text', 'options', 'correct', 'order', 'answer', 'pairs', 'key_text')
                   if k in item}
                  for item in items],
        'flags': flags or {},
        'notes': notes,
        **(extra or {}),
    }


# ── docx helpers ────────────────────────────────────────────────────────────────

def set_cell_merge(cell, value):
    tc_pr = cell._tc.get_or_add_tcPr()
    merge = OxmlElement('w:vMerge')
    if value == 'restart':
        merge.set(qn('w:val'), 'restart')
    tc_pr.append(merge)


def set_span(cell, span):
    tc_pr = cell._tc.get_or_add_tcPr()
    grid = OxmlElement('w:gridSpan')
    grid.set(qn('w:val'), str(span))
    tc_pr.append(grid)


def mark_run(run, signal):
    if signal == 'bold':
        run.bold = True
    elif signal == 'underline':
        run.underline = True
    elif signal == 'red':
        run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    elif signal == 'highlight':
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    elif signal == 'italic':
        run.italic = True


def add_omml(paragraph, text):
    m = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    omath = OxmlElement('m:oMath')
    r = OxmlElement('m:r')
    t = OxmlElement('m:t')
    t.text = text
    r.append(t)
    omath.append(r)
    paragraph._p.append(omath)


def tiny_png() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new('RGB', (12, 12), (200, 30, 30)).save(buffer, format='PNG')
    return buffer.getvalue()


def save(doc, path):
    doc.save(path)


# ── renderers ───────────────────────────────────────────────────────────────────


def docx_table_rows(path, sets, signal='bold'):
    """The owner's layout: header row, topic rows spanning 3 columns, № and question merged down, one option per row."""
    doc = Document()
    doc.add_paragraph('')
    items = flat(sets)
    rows = 1 + len(sets) + sum(len(item['options']) for item in items)
    table = doc.add_table(rows=rows, cols=3)
    header = table.rows[0].cells
    for cell, text in zip(header, ['№', 'Содержание вопроса', 'Варианты ответов']):
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
    r = 1
    by_topic = {}
    for item in items:
        by_topic.setdefault(item['topic'], []).append(item)
    for topic_index, (title, _) in enumerate(sets, start=1):
        row = table.rows[r].cells
        merged = row[0].merge(row[2])
        merged.paragraphs[0].add_run(title).bold = True
        r += 1
        for item in by_topic[topic_index]:
            for k, option in enumerate(item['options']):
                cells = table.rows[r].cells
                if k == 0:
                    cells[0].paragraphs[0].add_run(item['label'])
                    cells[1].paragraphs[0].add_run(item['text'])
                    set_cell_merge(cells[0], 'restart')
                    set_cell_merge(cells[1], 'restart')
                else:
                    set_cell_merge(cells[0], 'continue')
                    set_cell_merge(cells[1], 'continue')
                run = cells[2].paragraphs[0].add_run(option)
                if k in item['correct']:
                    mark_run(run, signal)
                r += 1
    save(doc, path)
    return items


def docx_paragraphs(path, sets, signal='bold', stems_bold=False, break_item=None):
    doc = Document()
    doc.add_paragraph('Тестовые задания по учебной дисциплине «Основы программирования»')
    items = flat(sets)
    for topic_index, (title, _) in enumerate(sets, start=1):
        doc.add_paragraph(title)
        for item in [i for i in items if i['topic'] == topic_index]:
            p = doc.add_paragraph()
            run = p.add_run(f"{item['label']}. {item['text']}")
            run.bold = stems_bold
            for k, option in enumerate(item['options']):
                p = doc.add_paragraph()
                p.add_run(f'{RUS[k]}) ')
                run = p.add_run(option)
                if k in item['correct'] and item['label'] != break_item:
                    mark_run(run, signal)
    save(doc, path)
    return items


def docx_autonumber(path, sets, signal='underline', two_marked=None):
    """Questions and options in one Word list: level 0 decimal «1.», level 1 russianLower «а)»."""
    doc = Document()
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), '77')
    for level, fmt, text in ((0, 'decimal', '%1.'), (1, 'russianLower', '%2)')):
        lvl = OxmlElement('w:lvl')
        lvl.set(qn('w:ilvl'), str(level))
        for tag, value in (('w:start', '1'), ('w:numFmt', fmt), ('w:lvlText', text)):
            el = OxmlElement(tag)
            el.set(qn('w:val'), value)
            lvl.append(el)
        abstract.append(lvl)
    first = numbering.find(qn('w:num'))
    if first is not None:
        first.addprevious(abstract)
    else:
        numbering.append(abstract)
    num = OxmlElement('w:num')
    num.set(qn('w:numId'), '77')
    ref = OxmlElement('w:abstractNumId')
    ref.set(qn('w:val'), '77')
    num.append(ref)
    numbering.append(num)

    def numbered(p, level):
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement('w:numPr')
        ilvl = OxmlElement('w:ilvl')
        ilvl.set(qn('w:val'), str(level))
        nid = OxmlElement('w:numId')
        nid.set(qn('w:val'), '77')
        num_pr.append(ilvl)
        num_pr.append(nid)
        p_pr.append(num_pr)

    items = flat(sets)
    for topic_index, (title, _) in enumerate(sets, start=1):
        doc.add_paragraph(title)
        for item in [i for i in items if i['topic'] == topic_index]:
            p = doc.add_paragraph(item['text'])
            numbered(p, 0)
            for k, option in enumerate(item['options']):
                p = doc.add_paragraph()
                run = p.add_run(option)
                if k in item['correct'] or (item['label'] == two_marked and k == (item['correct'][0] + 1) % len(item['options'])):
                    mark_run(run, signal)
                numbered(p, 1)
    save(doc, path)
    return items


def docx_options_in_cell(path, sets, signal='red', formula_item=None, image_item=None):
    """SPMI-like table: № п/п | Вопрос | Варианты ответов (options «1. …» as paragraphs in one cell)."""
    doc = Document()
    items = flat(sets)
    for topic_index, (title, _) in enumerate(sets, start=1):
        doc.add_paragraph(f'{topic_index}. {title.split(". ", 1)[-1]}. I категория:')
        topic_items = [i for i in items if i['topic'] == topic_index]
        table = doc.add_table(rows=1 + len(topic_items), cols=3)
        for cell, text in zip(table.rows[0].cells, ['№ п/п', 'Вопрос', 'Варианты ответов']):
            cell.paragraphs[0].add_run(text).bold = True
        for r, item in enumerate(topic_items, start=1):
            cells = table.rows[r].cells
            cells[0].paragraphs[0].add_run(item['label'])
            cells[1].paragraphs[0].add_run(item['text'])
            if item['label'] == formula_item:
                add_omml(cells[1].paragraphs[0], 'pv=RT')
            if item['label'] == image_item:
                cells[1].add_paragraph().add_run().add_picture(io.BytesIO(tiny_png()))
            for k, option in enumerate(item['options']):
                p = cells[2].paragraphs[0] if k == 0 else cells[2].add_paragraph()
                p.add_run(f'{k + 1}. ')
                run = p.add_run(option + (';' if k < len(item['options']) - 1 else '.'))
                if k in item['correct']:
                    mark_run(run, signal)
    save(doc, path)
    return items


def docx_key_column(path, sets, bad_key_item=None):
    """№ | Вопрос | Варианты ответов (inline «а) … б) …») | Правильный ответ."""
    doc = Document()
    items = flat(sets)
    table = doc.add_table(rows=1 + len(items) + len(sets), cols=4)
    for cell, text in zip(table.rows[0].cells, ['№', 'Вопрос', 'Варианты ответов', 'Правильный ответ']):
        cell.paragraphs[0].add_run(text).bold = True
    r = 1
    for topic_index, (title, _) in enumerate(sets, start=1):
        row = table.rows[r].cells
        row[0].merge(row[3]).paragraphs[0].add_run(title)
        r += 1
        for item in [i for i in items if i['topic'] == topic_index]:
            cells = table.rows[r].cells
            cells[0].paragraphs[0].add_run(item['label'])
            cells[1].paragraphs[0].add_run(item['text'])
            cells[2].paragraphs[0].add_run(' '.join(f'{RUS[k]}) {o};' for k, o in enumerate(item['options'])).rstrip(';'))
            key = ', '.join(RUS[k] for k in item['correct'])
            if item['label'] == bad_key_item:
                key = RUS[len(item['options'])]  # a letter past the last option
            cells[3].paragraphs[0].add_run(key)
            r += 1
    save(doc, path)
    return items


def docx_key_list(path, sets, latin_item=None):
    doc = Document()
    items = flat(sets)
    for topic_index, (title, _) in enumerate(sets, start=1):
        doc.add_paragraph(title)
        for item in [i for i in items if i['topic'] == topic_index]:
            doc.add_paragraph(f"{item['label']}. {item['text']}")
            for k, option in enumerate(item['options']):
                doc.add_paragraph(f'{RUS[k]}) {option}')
    doc.add_paragraph('Ключи к тесту')
    pairs = []
    for item in items:
        letters = ''.join(RUS[k] for k in item['correct'])
        if item['label'] == latin_item:
            letters = {'а': 'a', 'в': 'B', 'е': 'e'}.get(letters, letters)  # typed in the Latin layout
        pairs.append(f"{item['label']} – {letters}")
    doc.add_paragraph('; '.join(pairs))
    save(doc, path)
    return items


def docx_numbered_both(path, sets, extra_items):
    """ORGMA-like: «1. Вопрос» and options «1. вариант» both typed numbers, lowercase options, key table at the end."""
    doc = Document()
    doc.add_paragraph('Банк тестовых заданий')
    items = flat(sets) + extra_items
    for item in items:
        doc.add_paragraph(f"{item['label']}. {item['text']}")
        for k, option in enumerate(item['options']):
            doc.add_paragraph(f'{k + 1}. {option}')
        doc.add_paragraph('')
    doc.add_paragraph('Эталоны ответов')
    table = doc.add_table(rows=1 + len(items), cols=2)
    table.rows[0].cells[0].paragraphs[0].add_run('№ ТЗ')
    table.rows[0].cells[1].paragraphs[0].add_run('Ответ')
    for r, item in enumerate(items, start=1):
        table.rows[r].cells[0].paragraphs[0].add_run(item['label'])
        table.rows[r].cells[1].paragraphs[0].add_run(item.get('key_text') or ', '.join(str(k + 1) for k in item['correct']))
    save(doc, path)
    return items


def docx_inline_options(path, sets):
    doc = Document()
    items = flat(sets)
    for item in items:
        p = doc.add_paragraph()
        p.add_run(f"{item['label']}. {item['text']} ")
        for k, option in enumerate(item['options']):
            p.add_run(f'{RUS[k]}) ')
            run = p.add_run(option)
            if k in item['correct']:
                run.bold = True
            p.add_run('; ' if k < len(item['options']) - 1 else '.')
    save(doc, path)
    return items


def docx_messy(path):
    """What a hand-made file looks like: mixed numbering, a header, a wrapped option, a partial bold, a gap, a duplicate,
    a picture, an «Ответ:» line that contradicts the bold, bold spaces."""
    doc = Document()
    doc.add_paragraph('УО «Колледж» Цикловая комиссия программирования')
    doc.add_paragraph('Тестовые задания для самоподготовки')
    doc.add_paragraph('')
    items = []

    def add(label, text, options, correct, *, bold=None, partial=None, image=False, answer=None, gap_note=None, wrap=None):
        style_label = {'1': '1.', '2': '2)', '3': 'Вопрос 3.'}.get(label, f'{label}.')
        p = doc.add_paragraph(f'{style_label} {text}')
        if image:
            doc.add_paragraph().add_run().add_picture(io.BytesIO(tiny_png()))
        for k, option in enumerate(options):
            p = doc.add_paragraph()
            p.add_run(f'{RUS[k]}) ')
            if k == wrap:
                head, tail = option.split(' ', 1)
                r = p.add_run(head)
                r.bold = k in (bold if bold is not None else correct)
                p2 = doc.add_paragraph()
                r2 = p2.add_run(tail)
                r2.bold = k in (bold if bold is not None else correct)
                continue
            if partial is not None and k == partial[0]:
                p.add_run(option[: partial[1]]).bold = True
                p.add_run(option[partial[1]:])
            else:
                run = p.add_run(option)
                run.bold = k in (bold if bold is not None else correct)
            p.add_run('  ').bold = True  # bold spaces after every option
        if answer:
            doc.add_paragraph(f'Ответ: {answer}')
        doc.add_paragraph('')
        items.append({'label': label, 'topic': None, 'kind': 'single', 'text': text, 'options': options, 'correct': correct,
                      'image': image})

    add('1', 'Что выведет Console.WriteLine(5 / 2);?', ['2', '2.5', '3', 'ошибку компиляции'], [0])
    add('2', 'Какой тип хранит один символ Юникода?', ['byte', 'char', 'string', 'short'], [1])
    add('3', 'Какой модификатор делает поле доступным только внутри класса?', ['public', 'internal', 'private', 'protected'], [2])
    add('4', 'Как называется сайт вопросов и ответов для программистов?', ['StackOverflow', 'GitHub', 'Habr', 'MSDN'], [], partial=(0, 5))
    add('5', 'Что показано на рисунке?', ['блок-схема цикла', 'диаграмма классов', 'таблица истинности'], [0], image=True)
    add('6', 'Какой цикл выполнится хотя бы один раз?', ['for', 'while', 'do … while', 'foreach'], [2], answer='в')
    add('7', 'Какое значение по умолчанию у поля типа int?', ['null', '0', '-1', 'не определено'], [1], answer='а')
    add('9', 'Какой оператор проверяет равенство значений?', ['=', '==', '===', ':='], [1])
    add('10', 'Какой тип хранит один символ Юникода?', ['byte', 'char', 'string', 'short'], [1])
    add('11', 'Какое исключение возникает при делении целого числа на ноль?',
        ['DivideByZeroException', 'ArithmeticOverflowException в режиме checked при переполнении', 'NullReferenceException', 'FormatException'], [0], wrap=1)
    save(doc, path)
    flags = {'4': ['partial_mark', 'no_key'], '5': ['has_image'], '7': ['key_conflict'], '10': ['duplicate']}
    return items, flags


# ── сессия: несколько предметов в файле, открытые вопросы ────────────────────────

SESSION = [('Основы программирования', CSHARP), ('История Беларуси', HISTORY)]

OPEN_QUESTIONS = [
    ('Раздел 1. Язык C#', [
        'Что выведет Console.WriteLine(5 / 2); и почему?',
        'Объясните разницу между += и -= при подписке на событие.',
        'Зачем нужен #region Helpers и когда без него лучше обойтись?',
        'Что такое __init__ в Python и есть ли похожий метод в C#?',
    ]),
    ('Раздел 2. События и делегаты', [
        'Что делает строка count -= step; в цикле while (count > 0)?',
        'Чем делегат отличается от события.',
        'Опишите шаблон «издатель — подписчик».',
    ]),
]


def open_item(text, **extra):
    return {'text': text, 'options': [], 'correct': [], 'kind': 'open', **extra}


def flat_subjects(subjects):
    """[(название, sets)] → вопросы: номера с 1 в каждом предмете, темы сквозные по файлу, subject — с 1."""
    out, topic = [], 0
    for subject, (_, sets) in enumerate(subjects, start=1):
        number = 0
        for title, items in sets:
            if title:
                topic += 1
            for item in items:
                number += 1
                out.append({**copy.deepcopy(item), 'label': str(number), 'topic': topic if title else None, 'subject': subject})
    return out


def topic_subjects(subjects):
    return [subject for subject, (_, sets) in enumerate(subjects, start=1) for title, _ in sets if title]


def add_rows_table(doc, sets, signal='bold'):
    """Таблица владельца в документ doc: заголовок, строки тем на всю ширину (если у темы есть название),
    № и вопрос объединены вниз, вариант на строку; номера с 1."""
    rows = 1 + sum(1 for title, _ in sets if title) + sum(len(item['options']) for _, items in sets for item in items)
    table = doc.add_table(rows=rows, cols=3)
    for cell, text in zip(table.rows[0].cells, ['№', 'Содержание вопроса', 'Варианты ответов']):
        cell.paragraphs[0].add_run(text).bold = True
    r, number = 1, 0
    for title, items in sets:
        if title:
            row = table.rows[r].cells
            row[0].merge(row[2]).paragraphs[0].add_run(title).bold = True
            r += 1
        for item in items:
            number += 1
            for k, option in enumerate(item['options']):
                cells = table.rows[r].cells
                if k == 0:
                    cells[0].paragraphs[0].add_run(str(number))
                    cells[1].paragraphs[0].add_run(item['text'])
                set_cell_merge(cells[0], 'restart' if k == 0 else 'continue')
                set_cell_merge(cells[1], 'restart' if k == 0 else 'continue')
                run = cells[2].paragraphs[0].add_run(option)
                if k in item['correct']:
                    mark_run(run, signal)
                r += 1


def docx_two_subjects(path):
    """Два предмета: «Дисциплина: …» над каждой таблицей владельца, номера в каждой с 1."""
    doc = Document()
    for title, sets in SESSION:
        doc.add_paragraph(f'Дисциплина: {title}')
        add_rows_table(doc, sets)
        doc.add_paragraph('')
    save(doc, path)
    return flat_subjects(SESSION)


def docx_restart_tables(path):
    """Две таблицы без заголовков и без строк тем; во второй нумерация снова с 1."""
    parts = [(None, [('', CSHARP[0][1] + CSHARP[1][1])]), (None, [('', HISTORY[0][1] + HISTORY[1][1])])]
    doc = Document()
    for index, (_, sets) in enumerate(parts):
        if index:
            doc.add_paragraph('')
        add_rows_table(doc, sets)
    save(doc, path)
    return flat_subjects(parts)


def txt_subjects(path):
    """TXT: строки «Предмет: …» и «Предмет «…»», темы, вопросы с вариантами «а)» и строкой «Ответ: …»."""
    parts = [('Основы программирования', 'Предмет: Основы программирования', CSHARP[:2]),
             ('История Беларуси', 'Предмет «История Беларуси»', HISTORY)]
    lines = []
    for _, heading, sets in parts:
        lines += [heading, '']
        number = 0
        for title, items in sets:
            lines.append(title)
            for item in items:
                number += 1
                lines.append(f"{number}. {item['text']}")
                lines += [f'{RUS[k]}) {option}' for k, option in enumerate(item['options'])]
                lines.append('Ответ: ' + ', '.join(RUS[k] for k in item['correct']))
            lines.append('')
    txt(path, '\n'.join(lines) + '\n')
    return flat_subjects([(title, sets) for title, _, sets in parts]), [(title, sets) for title, _, sets in parts]


def add_decimal_list(doc, num_id='78'):
    """Список Word: уровень 0 — «1.», уровень 1 — «а)»; возвращает функцию, которая делает абзац пунктом списка."""
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), num_id)
    for level, fmt, text in ((0, 'decimal', '%1.'), (1, 'russianLower', '%2)')):
        lvl = OxmlElement('w:lvl')
        lvl.set(qn('w:ilvl'), str(level))
        for tag, value in (('w:start', '1'), ('w:numFmt', fmt), ('w:lvlText', text)):
            el = OxmlElement(tag)
            el.set(qn('w:val'), value)
            lvl.append(el)
        abstract.append(lvl)
    first = numbering.find(qn('w:num'))
    if first is not None:
        first.addprevious(abstract)
    else:
        numbering.append(abstract)
    num = OxmlElement('w:num')
    num.set(qn('w:numId'), num_id)
    ref = OxmlElement('w:abstractNumId')
    ref.set(qn('w:val'), num_id)
    num.append(ref)
    numbering.append(num)

    def numbered(p, level=0):
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement('w:numPr')
        ilvl = OxmlElement('w:ilvl')
        ilvl.set(qn('w:val'), str(level))
        nid = OxmlElement('w:numId')
        nid.set(qn('w:val'), num_id)
        num_pr.append(ilvl)
        num_pr.append(nid)
        p_pr.append(num_pr)

    return numbered


def docx_open_autonumber(path):
    """Список вопросов к экзамену без вариантов: автонумерация Word сквозная через два раздела."""
    doc = Document()
    numbered = add_decimal_list(doc)
    doc.add_paragraph('Вопросы к экзамену по дисциплине «Основы программирования»')
    for title, questions in OPEN_QUESTIONS:
        doc.add_paragraph(title)
        for text in questions:
            numbered(doc.add_paragraph(text))
    save(doc, path)
    sets = [(title, [open_item(text) for text in questions]) for title, questions in OPEN_QUESTIONS]
    return flat_subjects([('Основы программирования', sets)]), sets


def txt_open_list(path):
    """Список открытых вопросов в TXT: шапка, разделы, перенос строки, «3)» вместо «3.», строка «Ответ: …»."""
    (title1, first), (title2, second) = OPEN_QUESTIONS
    lines = ['Вопросы к экзамену', title1, f'1. {first[0]}',
             '2. Объясните разницу между += и -= при подписке', 'на событие и при отписке от него.',
             f'3) {first[2]}', f'4. {first[3]}', title2]
    lines += [f'{n}. {text}' for n, text in enumerate(second, start=5)]
    lines += ['8. Сколько байт занимает int?', 'Ответ: 4']
    txt(path, '\n'.join(lines) + '\n')
    texts1 = [first[0], 'Объясните разницу между += и -= при подписке на событие и при отписке от него.', first[2], first[3]]
    sets = [(title1, [open_item(t) for t in texts1]),
            (title2, [open_item(t) for t in second] + [open_item('Сколько байт занимает int?', answer='4')])]
    return flat_subjects([(None, sets)]), sets


MIXED = [
    q('Какое ключевое слово запрещает менять поле после конструктора?', ['readonly', 'const', 'sealed', 'static'], [0]),
    open_item('Объясните, зачем нужен #region Helpers.'),
    q('Какое ключевое слово подключает пространство имён?', ['import', 'include', 'using', 'require'], [2]),
    open_item('Что вернёт выражение items?.Count ?? 0, если список не создан?'),
    q('Какой тип лучше подходит для денежных сумм?', ['decimal', 'float', 'double', 'int'], [0]),
]


def docx_mixed(path):
    """Тест и открытые вопросы в одном списке: у вопросов с вариантами ключ жирным, у открытых вариантов нет."""
    doc = Document()
    for number, item in enumerate(MIXED, start=1):
        doc.add_paragraph(f"{number}. {item['text']}")
        for k, option in enumerate(item['options']):
            p = doc.add_paragraph()
            p.add_run(f'{RUS[k]}) ')
            p.add_run(option).bold = k in item['correct']
    save(doc, path)
    return flat_subjects([(None, [('', MIXED)])])


SUBPOINTS = [
    ('Экономические системы:', ['традиционная;', 'рыночная;', 'смешанная.'], 'а'),
    ('Функции денег.', [], ''),
    ('Виды инфляции по темпам роста:', ['ползучая;', 'галопирующая;', 'гиперинфляция.'], 'а'),
    ('Кривая Филлипса:', ['краткосрочная;', 'долгосрочная.'], '1'),
    ('Функции Центрального банка.', [], ''),
]


def txt_open_subpoints(path):
    """Открытые вопросы с подпунктами «а)», «1)»: в format «auto» — вопросы с вариантами без ключа, в «open» — открытые."""
    lines = []
    for number, (text, points, style) in enumerate(SUBPOINTS, start=1):
        lines.append(f'{number}. {text}')
        lines += [f"{RUS[k] if style == 'а' else k + 1}) {point}" for k, point in enumerate(points)]
    txt(path, '\n'.join(lines) + '\n')
    as_test, as_open = [], []
    for number, (text, points, style) in enumerate(SUBPOINTS, start=1):
        if points:
            as_test.append({'label': str(number), 'topic': None, 'subject': 1, 'kind': 'single', 'text': text,
                            'options': [point.rstrip(';.') for point in points], 'correct': []})
        else:
            as_test.append({'label': str(number), 'topic': None, 'subject': 1, **open_item(text)})
        written = [f"{RUS[k] if style == 'а' else k + 1}) {point}" for k, point in enumerate(points)]
        as_open.append({'label': str(number), 'topic': None, 'subject': 1, **open_item('\n'.join([text, *written]))})
    return as_test, as_open


def txt(path, text):
    with open(path, 'w', encoding='utf-8') as handle:
        handle.write(text)


def mytest_text(items):
    blocks = []
    for item in items:
        if item['kind'] == 'order':
            lines = [item['text']] + [f"{item['order'][k]}. {o}" for k, o in enumerate(item['options'])]
        elif item['kind'] == 'open':
            lines = [item['text'], f"= {item['answer']}"]
        else:
            lines = [item['text']] + [('+ ' if k in item['correct'] else '- ') + o for k, o in enumerate(item['options'])]
        blocks.append('\n'.join(lines))
    return '\n\n'.join(blocks) + '\n'


def indigo_text(items):
    blocks = []
    for item in items:
        if item['kind'] == 'multiple':
            lines = [item['text']] + [('#' if k in item['correct'] else '') + o for k, o in enumerate(item['options'])]
        elif item['kind'] == 'open':
            lines = [item['text'], f"[{item['answer']}]"]
        elif item['kind'] == 'match':
            lines = [item['text']] + [f'{a} = {b}' for a, b in item['pairs']]
        elif item['kind'] == 'order':
            lines = [item['text']] + item['options']
        else:
            lines = [item['text']] + [('*' if k in item['correct'] else '') + o for k, o in enumerate(item['options'])]
        blocks.append('\n'.join(lines))
    return '\n\n'.join(blocks) + '\n'


def aiken_text(items):
    blocks = []
    for item in items:
        lines = [item['text']] + [f"{'ABCDEFGH'[k]}. {o}" for k, o in enumerate(item['options'])]
        lines.append(f"ANSWER: {'ABCDEFGH'[item['correct'][0]]}")
        blocks.append('\n'.join(lines))
    return '\n\n'.join(blocks) + '\n'


def gift_escape(text):
    for ch in '\\~=#{}:':
        text = text.replace(ch, '\\' + ch)
    return text


def gift_text(sets):
    out = []
    for title, items in sets:
        out.append(f'$CATEGORY: {title}\n')
        for n, item in enumerate(items, start=1):
            if item['kind'] == 'multiple':
                share = 100 // len(item['correct'])
                answers = ' '.join((f'~%{share}%' if k in item['correct'] else '~%-100%') + gift_escape(o) for k, o in enumerate(item['options']))
            else:
                answers = ' '.join(('=' if k in item['correct'] else '~') + gift_escape(o) for k, o in enumerate(item['options']))
            out.append(f'::{title[:8]} {n}:: {gift_escape(item["text"])} {{{answers}}}\n')
    return '\n'.join(out)


def xlsx_columns(path, sets):
    import xlsxwriter

    items = flat(sets)
    book = xlsxwriter.Workbook(path)
    sheet = book.add_worksheet('Тест')
    bold = book.add_format({'bold': True})
    sheet.write_row(0, 0, ['№', 'Вопрос', 'А', 'Б', 'В', 'Г', 'Д', 'Ответ'], bold)
    r = 1
    for topic_index, (title, _) in enumerate(sets, start=1):
        sheet.write(r, 0, title, bold)
        r += 1
        for item in [i for i in items if i['topic'] == topic_index]:
            sheet.write_number(r, 0, int(item['label']))
            sheet.write(r, 1, item['text'])
            for k, option in enumerate(item['options']):
                if option.isdigit():
                    sheet.write_number(r, 2 + k, int(option))  # numeric cell: must come back as «0», not «0.0»
                else:
                    sheet.write_string(r, 2 + k, option)
            sheet.write(r, 7, ', '.join('АБВГД'[k] for k in item['correct']))
            r += 1
    book.close()
    return items


# ── main ────────────────────────────────────────────────────────────────────────


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
    os.makedirs(out, exist_ok=True)
    specs = []

    def emit(name, sets, items, **kw):
        specs.append(expected(name, sets, items, **kw))

    emit('01_table_rows_bold.docx', CSHARP, docx_table_rows(f'{out}/01_table_rows_bold.docx', CSHARP),
         layout='table_option_rows', key='bold',
         notes='The owner\'s layout. Header row and topic rows are bold too: they must not count as keys. Code options exact.')
    emit('02_paragraphs_bold_stems.docx', CSHARP, docx_paragraphs(f'{out}/02_paragraphs_bold_stems.docx', CSHARP, 'bold', stems_bold=True, break_item='7'),
         layout='paragraph_list', key='bold', flags={'7': ['no_key']},
         notes='Every question is bold too; the key is the bold option. Item 7 has no bold option. '
               'The header «Тестовые задания по учебной дисциплине «…»» names the one subject.',
         extra={'subjects': [{'title': 'Основы программирования', 'evidence': 'heading'}]})
    emit('03_autonumber_underline.docx', HISTORY, docx_autonumber(f'{out}/03_autonumber_underline.docx', HISTORY, 'underline', two_marked='2'),
         layout='paragraph_list', key='underline', flags={'2': ['several_marked']},
         notes='Word list: level 0 decimal, level 1 russianLower; labels exist only in numbering.xml. Item 2 has two underlined options.')
    items4 = docx_options_in_cell(f'{out}/04_options_in_cell_red.docx', CSHARP, 'red', formula_item='2', image_item='6')
    in_cell_topics = [(f'{n}. {title.split(". ", 1)[-1]}. I категория:', None) for n, (title, _) in enumerate(CSHARP, start=1)]
    emit('04_options_in_cell_red.docx', in_cell_topics, items4, layout='table_options_in_cell', key='color',
         flags={'2': ['has_formula'], '6': ['has_image']},
         notes='Options «1. …;» as paragraphs of one cell, key in red. Trailing «;» and «.» are cut from options. Topic headings above each table.')
    emit('05_key_column_inline.docx', HISTORY, docx_key_column(f'{out}/05_key_column_inline.docx', HISTORY, bad_key_item='5'),
         layout='table_options_in_cell', key='key_column', flags={'5': ['no_key', 'key_invalid']},
         notes='Options inline in one paragraph; key column «б» / «а, в, д» (multiple). Item 5 key letter is past the last option.')
    emit('06_key_list_end.docx', HISTORY, docx_key_list(f'{out}/06_key_list_end.docx', HISTORY, latin_item='2'),
         layout='paragraph_list', key='key_list', flags={'2': ['key_letter_ambiguous']},
         notes='No formatting; «Ключи к тесту» then «1 – б; 2 – в; …». Item 2 key typed as Latin «B» (looks like «В», but B is also the 2nd Latin letter): take «в» and flag it.')
    order = {'text': 'Расставьте месяцы по порядку:', 'options': ['март', 'январь', 'февраль'], 'order': [3, 1, 2], 'correct': [], 'kind': 'order'}
    open_q = {'text': 'Каким по счёту идёт месяц май?', 'options': [], 'answer': '5', 'correct': [], 'kind': 'open'}
    mt_sets = [('', CSHARP[0][1] + [HISTORY[0][1][2]])]
    mt_items = flat(mt_sets) + [{**order, 'label': '', 'topic': None}, {**open_q, 'label': '', 'topic': None}]
    txt(f'{out}/07_mytest_plusminus.txt', mytest_text(mt_items))
    emit('07_mytest_plusminus.txt', [], [{**i, 'label': ''} for i in mt_items], layout='marker_text', key='marker:+/-',
         notes='MyTestXPro text: blank line between items, «+»/«-» before options. Ordering by numbers, open by «=». No numbers on questions.')
    match_q = {'text': 'Соотнесите слово с переводом:', 'options': [], 'pairs': [('wall', 'стена'), ('window', 'окно'), ('door', 'дверь')], 'correct': [], 'kind': 'match'}
    in_items = [*flat([('', CSHARP[1][1])]), {**HISTORY[0][1][2], 'label': '', 'topic': None},
                {**open_q, 'text': 'Сколько дней в високосном году?', 'answer': '366', 'label': '', 'topic': None},
                {**match_q, 'label': '', 'topic': None}]
    txt(f'{out}/08_indigo.txt', indigo_text(in_items))
    emit('08_indigo.txt', [], [{**i, 'label': ''} for i in in_items], layout='marker_text', key='marker:*/#',
         notes='Indigo: «*» single, «#» multiple, «[ответ]» open, «a = b» matching. The «*» option of item 1 (an operator) is content: '
               'no text follows it.')
    aiken_items = flat([('', CSHARP[0][1])])
    txt(f'{out}/09_aiken.txt', aiken_text(aiken_items))
    emit('09_aiken.txt', [], [{**i, 'label': ''} for i in aiken_items], layout='aiken', key='ANSWER:',
         notes='Moodle Aiken: Latin capitals «A.», «ANSWER: B». Options like «#», «+=», «<<» are content.')
    txt(f'{out}/10_gift.txt', gift_text(HISTORY + [CSHARP[1]]))
    emit('10_gift.txt', HISTORY + [CSHARP[1]], [{**i, 'label': ''} for i in flat(HISTORY + [CSHARP[1]])], layout='gift', key='gift',
         notes='Moodle GIFT: $CATEGORY topics, ::title::, {=right ~wrong}, %33% weights for multiple, escaped \\= \\{ \\} \\# \\: \\~ in code.')
    emit('11_columns_key.xlsx', CSHARP, xlsx_columns(f'{out}/11_columns_key.xlsx', CSHARP), layout='table_option_columns', key='key_column',
         notes='XLSX: № | Вопрос | А…Д | Ответ; topic rows in column A; numeric option cells («0») stay «0».')
    ext = [
        {'label': '13', 'topic': None, 'kind': 'multiple', 'text': 'К объектно-ориентированным языкам относятся:',
         'options': ['c#', 'html', 'java', 'css'], 'correct': [0, 2]},
        {'label': '14', 'topic': None, 'kind': 'order', 'text': 'Установите правильную последовательность этапов компиляции:',
         'options': ['компоновка', 'лексический анализ', 'генерация кода', 'синтаксический анализ'], 'correct': [], 'key_text': '2, 4, 3, 1'},
        {'label': '15', 'topic': None, 'kind': 'match', 'text': 'Установите соответствие между типом и размером: А. int Б. long В. byte',
         'options': ['1 байт', '4 байта', '8 байт'], 'correct': [], 'key_text': 'А-2, Б-3, В-1'},
    ]
    both_sets = [('', CSHARP[0][1] + CSHARP[1][1] + CSHARP[2][1])]
    emit('12_numbered_both_key_table.docx', [], docx_numbered_both(f'{out}/12_numbered_both_key_table.docx', both_sets, ext),
         layout='paragraph_list', key='key_list',
         notes='Questions «N.» and options «1.»…«4.» both numbers; the next question number must not be taken for an option and back. '
               'Key table «№ ТЗ | Ответ» at the end with «1, 3», «2, 4, 3, 1», «А-2, Б-3, В-1».')
    emit('13_inline_bold.docx', HISTORY, docx_inline_options(f'{out}/13_inline_bold.docx', HISTORY), layout='inline_options', key='bold',
         notes='Options in the question paragraph «… а) x; б) y; в) z.»; the bold run inside the paragraph is the key. No topics in the file.',
         extra={'topics_in_file': False})
    messy_items, messy_flags = docx_messy(f'{out}/14_messy.docx')
    emit('14_messy.docx', [], messy_items, layout='paragraph_list', key='bold', flags=messy_flags,
         notes='Header lines skipped; «1.», «2)», «Вопрос 3.» all questions; gap: 8 missing; item 10 repeats item 2; item 11 option б '
               'continues on the next paragraph (joined); item 4 bold covers 5 of 13 letters (partial, no key); item 7 «Ответ: а» '
               'contradicts bold б (conflict, explicit wins); bold spaces ignored.', extra={'gaps': [[8, 8]]})
    paste = mytest_text([dict(i, kind='single') for i in flat([('', CSHARP[2][1])])]).replace('+ ', '').replace('- ', '')
    numbered = []
    for n, block in enumerate(paste.strip().split('\n\n'), start=1):
        lines = block.split('\n')
        numbered.append('\n'.join([f'{n}.\t{lines[0]}'] + [f'{RUS[k]})\t{o}' for k, o in enumerate(lines[1:])]))
    txt(f'{out}/15_pasted_no_key.txt', '\n'.join(numbered) + '\n')
    pasted_items = [dict(i, correct=[]) for i in flat([('', CSHARP[2][1])])]
    emit('15_pasted_no_key.txt', [], pasted_items, layout='paragraph_list', key='none',
         flags={str(n): ['no_key'] for n in range(1, 5)},
         notes='What a textarea paste from Word gives: tabs after numbers and letters, formatting lost. Expect the «no key» screen; '
               'with «first option is right» every item gets [0]; with a key file see 16.')
    key_items = flat(CSHARP)
    doc = Document()
    for item in key_items:
        doc.add_paragraph(f"{item['label']}. {item['text']}")
        for k, option in enumerate(item['options']):
            doc.add_paragraph(f'{RUS[k]}) {option}')
    save(doc, f'{out}/16_questions_only.docx')
    txt(f'{out}/16_key.txt', '\n'.join(f"{i['label']}. {RUS[i['correct'][0]]}" for i in key_items) + '\n')
    emit('16_questions_only.docx', CSHARP, key_items, layout='paragraph_list', key='key_file:16_key.txt',
         notes='Questions without any key + a separate key file «1. а». Topics are lines «Тема N» — here absent from the file '
               '(topics list empty when parsed without them).', extra={'topics_in_file': False})
    # 19: a FOS-style list as seen in a real university PDF: «+» glued after right options, several «+» without a hint
    # in the question, an item with no «+», a Latin «C)» among Cyrillic labels, initials in options, a wrapped question.
    doc = Document()
    doc.add_paragraph().add_run('1) Общая психология').bold = True
    psy = [
        ('1', 'Выделяют три стороны общения:', ['интерактивную', 'субъективную', 'коммуникативную', 'перцептивную'], [0, 2, 3], 'multiple'),
        ('2', 'Основатели гештальтпсихологии:', ['К.Коффка', 'В. Вундт', 'М. Вертхеймер', 'Дж.Уотсон'], [0, 2], 'multiple'),
        ('3', 'Если собеседник отводит взгляд в сторону, это означает, что он', ['не понимает', 'не хочет общаться', 'формулирует мысль', 'хочет закончить разговор'], [2], 'single'),
        ('4', 'Верно ли утверждение: эмоции длятся дольше чувств?', ['верно', 'неверно', 'затрудняюсь ответить'], [], 'single'),
    ]
    for label, text, options, correct, kind in psy:
        if label == '3':
            head, tail = text[:40].rsplit(' ', 1)
            doc.add_paragraph(f'{label}. {head}')
            doc.add_paragraph(text[len(head) + 1:] + ':')
        else:
            doc.add_paragraph(f'{label}. {text}')
        letters = 'АБС' if label == '4' else 'абвг'
        for k, option in enumerate(options):
            doc.add_paragraph(f'{letters[k]}) {option}' + ('+' if k in correct else ''))
        doc.add_paragraph('')
    save(doc, f'{out}/19_plus_suffix.docx')
    psy_items = [{'label': l, 'topic': 1, 'kind': k, 'text': t if l != '3' else t + ':', 'options': o, 'correct': c} for l, t, o, c, k in psy]
    emit('19_plus_suffix.docx', [('1) Общая психология', [])], psy_items, layout='blocks', key='marker:suffix+',
         flags={'4': ['no_key']},
         notes='«+» glued after the right options (seen in a real FOS PDF). Several «+» mean several right answers. Item 4 has no '
               '«+» at all; its labels are «А) Б) С)» with a Latin C. Initials «К.Коффка», «В. Вундт» are option text. '
               'Item 3 question wraps onto a second line.', extra={'topics_in_file': False})
    # 20–26: сессия из нескольких предметов, открытые вопросы, format «open»
    emit('20_two_subjects.docx', [(t, None) for _, sets in SESSION for t, _ in sets], docx_two_subjects(f'{out}/20_two_subjects.docx'),
         layout='table_option_rows', key='bold',
         notes='Two subjects, each «Дисциплина: …» above its own table of the owner layout; numbering restarts at 1. '
               'Topics are numbered across the file and belong to their subject.',
         extra={'subjects': [{'title': title, 'evidence': 'heading'} for title, _ in SESSION], 'topic_subjects': topic_subjects(SESSION)})
    emit('21_two_tables_restart.docx', [], docx_restart_tables(f'{out}/21_two_tables_restart.docx'), layout='table_option_rows', key='bold',
         notes='Two tables, no headings and no topic rows; the second restarts at 1: two subjects without titles, evidence «restart», '
               'and a warning so the client can ask whether to split.',
         extra={'subjects': [{'title': None, 'evidence': 'restart'}] * 2, 'warnings': ['subjects_restart']})
    txt_items, txt_parts = txt_subjects(f'{out}/22_subjects.txt')
    emit('22_subjects.txt', [(t, None) for _, sets in txt_parts for t, _ in sets], txt_items, layout='paragraph_list', key='answer_line',
         notes='TXT with «Предмет: …» and «Предмет «…»» lines; topics, «а)» options and «Ответ: …» lines; numbering restarts per subject.',
         extra={'subjects': [{'title': title, 'evidence': 'heading'} for title, _ in txt_parts], 'topic_subjects': topic_subjects(txt_parts)})
    open_items, open_sets = docx_open_autonumber(f'{out}/23_open_autonumber.docx')
    emit('23_open_autonumber.docx', open_sets, open_items, layout='paragraph_list', key='none',
         notes='Exam questions without options in one Word list (numbering runs through both sections): open items, text verbatim '
               '(«+=», «-=», «#region Helpers», «__init__», «count -= step;»). The header names the subject.',
         extra={'subjects': [{'title': 'Основы программирования', 'evidence': 'heading'}]})
    list_items, list_sets = txt_open_list(f'{out}/24_open_list.txt')
    emit('24_open_list.txt', list_sets, list_items, layout='paragraph_list', key='none',
         notes='Open questions in TXT: a header line that is not a question, sections, a wrapped question, «3)» numbering, '
               'and an «Ответ: 4» line that becomes the answer of an open item.')
    emit('25_mixed.docx', [], docx_mixed(f'{out}/25_mixed.docx'), layout='paragraph_list', key='bold',
         notes='Test items (bold key) and open items in one numbered list: every item keeps its own kind.')
    as_test, as_open = txt_open_subpoints(f'{out}/26_open_subpoints.txt')
    emit('26_open_subpoints.txt', [], as_test, layout='paragraph_list', key='none',
         flags={str(n): ['no_key'] for n, (_, points, _) in enumerate(SUBPOINTS, start=1) if points},
         notes='format «auto»: sub-points «а)» and «1)» of open questions look like options: test items without a key.')
    emit('26_open_subpoints.txt', [], as_open, layout='paragraph_list', key='none',
         notes='format «open»: the same file, every item open, sub-points kept in the question text as written.',
         extra={'format': 'open'})
    with open(f'{out}/expected.json', 'w', encoding='utf-8') as handle:
        json.dump(specs, handle, ensure_ascii=False, indent=1)
    print(f'{len(specs)} files')


if __name__ == '__main__':
    main()
