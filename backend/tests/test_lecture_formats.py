"""Форматы лекций: PPTX (python-pptx) и распознавание сканов и фото (ocr.py).

Gemini подменён (без сети): подделка узнаёт страницы PDF по ширине — у страницы i ширина
300 + i пунктов — и «распознаёт» каждую строкой «Страница N. …». Агенты анализа — подменённый
BaseAgent.run из test_analyze_stream.
"""
import asyncio
import io
import json
import struct
import sys
import zlib
from pathlib import Path
from types import SimpleNamespace

import pytest
from lxml import etree
from PyPDF2 import PdfReader, PdfWriter
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import exams_api  # noqa: E402
import media_parser  # noqa: E402
import ocr  # noqa: E402
from agents import base  # noqa: E402
from agents.base import AIError, BaseAgent  # noqa: E402
from parser import PPT_LEGACY_DETAIL, PPTX_NO_TEXT_DETAIL, UnsupportedFormat, extract_text  # noqa: E402
from ratelimit import RATE_LIMIT_DETAIL  # noqa: E402
from test_analyze_stream import SUMMARY, answer_with, frames, kinds  # noqa: E402

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DGM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


# ── Файлы для тестов ─────────────────────────────────────────────────────────────


def deck_bytes(prs) -> bytes:
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def lecture_deck() -> bytes:
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[0])  # титульный: заголовок и подзаголовок
    first.shapes.title.text = "Лекция 3. Денежный рынок"
    first.placeholders[1].text = "Макроэкономика"

    second = prs.slides.add_slide(prs.slide_layouts[5])  # только заголовок
    second.shapes.title.text = "Спрос на деньги"
    # Правая колонка добавлена первой и на 0,05 дюйма ниже левой — читается всё равно после левой.
    right = second.shapes.add_textbox(Inches(5), Inches(2.05), Inches(4), Inches(1))
    right.text_frame.text = "Правая колонка"
    left = second.shapes.add_textbox(Inches(0.5), Inches(2), Inches(4), Inches(1))
    left.text_frame.text = "Мотивы спроса:"
    left.text_frame.add_paragraph().text = "трансакционный\vспекулятивный"  # \v — перенос строки внутри абзаца
    below = second.shapes.add_textbox(Inches(0.5), Inches(3), Inches(4), Inches(1))
    below.text_frame.text = "Под колонками"
    table = second.shapes.add_table(3, 3, Inches(0.5), Inches(4), Inches(8), Inches(1.5)).table
    for column, text in enumerate(["Мотив", "Зависит от", ""]):
        table.cell(0, column).text = text
    for column, text in enumerate(["Трансакционный", "дохода", "Y"]):
        table.cell(1, column).text = text
    table.cell(2, 0).merge(table.cell(2, 2))
    table.cell(2, 0).text = "Итого: спрос растёт с доходом"
    table.cell(2, 1).text = "под объединением"
    group = second.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(4), Inches(6), Inches(2), Inches(0.5)).text_frame.text = "Группа: справа"
    group.shapes.add_textbox(Inches(1), Inches(6), Inches(2), Inches(0.5)).text_frame.text = "Группа: слева"
    second.notes_slide.notes_text_frame.text = "Сказать про ставку.\nПривести пример."
    for placeholder in (
        f'<p:sp xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:nvSpPr><p:cNvPr id="40" name="Нижний колонтитул"/><p:cNvSpPr/>'
        f'<p:nvPr><p:ph type="ftr" sz="quarter" idx="11"/></p:nvPr></p:nvSpPr><p:spPr/>'
        f"<p:txBody><a:bodyPr/><a:p><a:r><a:t>Москва 2026</a:t></a:r></a:p></p:txBody></p:sp>",
        f'<p:sp xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:nvSpPr><p:cNvPr id="41" name="Номер слайда"/><p:cNvSpPr/>'
        f'<p:nvPr><p:ph type="sldNum" sz="quarter" idx="12"/></p:nvPr></p:nvSpPr><p:spPr/>'
        f"<p:txBody><a:bodyPr/><a:p><a:r><a:t>2</a:t></a:r></a:p></p:txBody></p:sp>",
    ):
        second.shapes._spTree.append(etree.fromstring(placeholder))

    prs.slides.add_slide(prs.slide_layouts[6])  # пустой слайд — пропускается
    fourth = prs.slides.add_slide(prs.slide_layouts[6])
    fourth.notes_slide.notes_text_frame.text = "Только заметки"
    return deck_bytes(prs)


LECTURE_DECK_TEXT = """Слайд 1
Лекция 3. Денежный рынок
Макроэкономика

Слайд 2
Спрос на деньги
Мотивы спроса:
трансакционный
спекулятивный
Правая колонка
Под колонками
Мотив | Зависит от
Трансакционный | дохода | Y
Итого: спрос растёт с доходом
Группа: слева
Группа: справа
Заметки: Сказать про ставку.
Привести пример.

Слайд 4
Заметки: Только заметки"""


def smartart_and_chart_deck() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Обращение денег"
    data = (
        f'<dgm:dataModel xmlns:dgm="{DGM_NS}" xmlns:a="{A_NS}"><dgm:ptLst>'
        '<dgm:pt modelId="0" type="doc"><dgm:t><a:bodyPr/><a:p><a:r><a:t>корень</a:t></a:r></a:p></dgm:t></dgm:pt>'
        '<dgm:pt modelId="1"><dgm:t><a:bodyPr/><a:p><a:r><a:t>Эмиссия</a:t></a:r></a:p></dgm:t></dgm:pt>'
        '<dgm:pt modelId="2" type="parTrans"><dgm:t><a:bodyPr/><a:p><a:r><a:t>связь</a:t></a:r></a:p></dgm:t></dgm:pt>'
        '<dgm:pt modelId="3"><dgm:t><a:bodyPr/><a:p><a:r><a:t>Банки </a:t></a:r><a:r><a:t>и население</a:t></a:r></a:p>'
        "</dgm:t></dgm:pt>"
        '<dgm:pt modelId="4" type="pres"/>'
        "</dgm:ptLst></dgm:dataModel>"
    ).encode()
    part = Part(
        PackURI("/ppt/diagrams/data1.xml"),
        "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
        prs.part.package,
        data,
    )
    rel_id = slide.part.relate_to(part, f"{R_NS}/diagramData")
    frame = (
        f'<p:graphicFrame xmlns:p="{P_NS}" xmlns:a="{A_NS}" xmlns:r="{R_NS}"><p:nvGraphicFramePr>'
        '<p:cNvPr id="10" name="Схема 1"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
        '<p:xfrm><a:off x="457200" y="1828800"/><a:ext cx="6000000" cy="2000000"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="{DGM_NS}"><dgm:relIds xmlns:dgm="{DGM_NS}" r:dm="{rel_id}" r:lo="" r:qs="" '
        'r:cs=""/></a:graphicData></a:graphic></p:graphicFrame>'
    )
    slide.shapes._spTree.append(etree.fromstring(frame))
    chart_data = CategoryChartData()
    chart_data.categories = ["2024", "2025"]
    chart_data.add_series("Денежная масса", (1.0, 2.0))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(5), Inches(6), Inches(2), chart_data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Денежная масса М2"
    return deck_bytes(prs)


def pictures_only_deck() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(1))
    return deck_bytes(prs)


def png_bytes() -> bytes:
    """PNG 1×1 без сторонних библиотек."""

    def chunk(kind: bytes, payload: bytes) -> bytes:
        return struct.pack(">I", len(payload)) + kind + payload + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(b"\x00\xff\xff\xff")) + chunk(b"IEND", b"")


def scan_pdf(pages: int) -> bytes:
    """PDF без текстового слоя; ширина страницы i — 300 + i пунктов (по ней подделка узнаёт страницу)."""
    writer = PdfWriter()
    for index in range(pages):
        writer.add_blank_page(width=300 + index, height=842)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def text_pdf(page_texts: list[str]) -> bytes:
    """Маленький PDF с текстовым слоем (Helvetica, латиница): PyPDF2 читает его текст."""
    objects: dict[int, bytes] = {3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"}
    kids = []
    number = 4
    for text in page_texts:
        content = f"BT /F1 12 Tf 72 712 Td ({text}) Tj ET".encode("latin-1")
        objects[number + 1] = b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream"
        objects[number] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {number + 1} 0 R >>"
        ).encode()
        kids.append(number)
        number += 2
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(f'{kid} 0 R' for kid in kids)}] /Count {len(kids)} >>".encode()
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = {}
    for key in sorted(objects):
        offsets[key] = out.tell()
        out.write(f"{key} 0 obj\n".encode() + objects[key] + b"\nendobj\n")
    xref = out.tell()
    size = max(objects) + 1
    out.write(f"xref\n0 {size}\n0000000000 65535 f \n".encode())
    for key in range(1, size):
        out.write(f"{offsets[key]:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def page_numbers(data: bytes) -> list[int]:
    return [int(float(page.mediabox.width)) - 300 + 1 for page in PdfReader(io.BytesIO(data)).pages]


def page_text(number: int) -> str:
    """Одна страница — больше main.LECTURE_MIN_CHARS букв: такой скан анализируется и с одной страницы."""
    return (
        f"Страница {number}. Деньги выполняют функцию меры стоимости: цены всех товаров выражают в деньгах, "
        "и по ним товары сравнивают между собой. Деньги служат и средством обращения: товар продают за деньги, "
        "а на вырученные деньги покупают другой товар."
    )


def transcribe(call):
    """Обычный ответ: каждая страница куска — своей строкой."""
    return "\n\n".join(page_text(n) for n in page_numbers(call["data"])), "STOP"


# ── Подменённый Gemini ───────────────────────────────────────────────────────────


class FakeGemini:
    """google.genai.Client для ocr.py: запоминает вызовы, отвечает behaviour(call) -> (text, finish) | Exception."""

    def __init__(self, behaviour=transcribe):
        self.behaviour = behaviour
        self.calls: list[dict] = []
        self.active = 0
        self.max_active = 0

    def install(self, monkeypatch, keys="gemini-key-1", models="gemini-2.5-flash-lite"):
        monkeypatch.setenv("GEMINI_API_KEY", keys)
        if models is None:
            monkeypatch.delenv("GEMINI_MODELS", raising=False)
        else:
            monkeypatch.setenv("GEMINI_MODELS", models)
        fake = self

        class Client:
            def __init__(self, api_key=None, **kwargs):
                async def generate_content(*, model, contents, config):
                    part, prompt = contents
                    call = {
                        "key": api_key,
                        "model": model,
                        "config": config,
                        "mime": part.inline_data.mime_type,
                        "data": part.inline_data.data,
                        "prompt": prompt,
                    }
                    fake.calls.append(call)
                    fake.active += 1
                    fake.max_active = max(fake.max_active, fake.active)
                    try:
                        await asyncio.sleep(0.01)
                        result = fake.behaviour(call)
                        if asyncio.iscoroutine(result):
                            result = await result
                    finally:
                        fake.active -= 1
                    if isinstance(result, BaseException):
                        raise result
                    text, finish = result
                    return SimpleNamespace(text=text, candidates=[SimpleNamespace(finish_reason=finish)])

                self.aio = SimpleNamespace(models=SimpleNamespace(generate_content=generate_content))

        monkeypatch.setattr(base.genai, "Client", Client)
        base._GENAI_CLIENTS.clear()
        base._MODEL_SKIP.clear()
        return self


def gemini_error(status: int, name: str, message: str = "provider text"):
    from google.genai import errors as genai_errors

    return genai_errors.ClientError(status, {"error": {"code": status, "message": message, "status": name}})


@pytest.fixture(autouse=True)
def clean_env(monkeypatch):
    for name in ("GEMINI_API_KEY", "GEMINI_MODELS", "GROQ_API_KEY", "MAX_UPLOAD_BYTES", "OCR_REQUEST_TIMEOUT_SECONDS",
                 "OCR_TOTAL_TIMEOUT_SECONDS", "MODEL_SKIP_SECONDS"):
        monkeypatch.delenv(name, raising=False)
    base._GENAI_CLIENTS.clear()
    base._MODEL_SKIP.clear()
    yield
    base._GENAI_CLIENTS.clear()
    base._MODEL_SKIP.clear()


@pytest.fixture
def api(monkeypatch):
    import main
    from fastapi.testclient import TestClient

    async def fake_user(request):
        token = request.headers.get("Authorization", "Bearer user-a").split(" ", 1)[1]
        return {"id": token}

    monkeypatch.setattr(main, "require_authenticated_user", fake_user)
    monkeypatch.setattr(main.limiter, "enabled", False)
    monkeypatch.setenv("AI_AGENT_STAGGER_SECONDS", "0")
    BaseAgent._run_cache.clear()
    client = TestClient(main.app, headers={"Authorization": "Bearer user-a"})
    client.main = main
    yield client
    BaseAgent._run_cache.clear()


def assert_plain_russian(response, status, detail=None):
    assert response.status_code == status, response.text
    body = response.json()
    assert set(body) <= {"detail", "code"}
    assert "provider text" not in body["detail"] and "Error" not in body["detail"]
    if detail is not None:
        assert body["detail"] == detail


# ── PPTX ──────────────────────────────────────────────────────────────────────────


def test_pptx_text_by_slides_with_titles_reading_order_tables_groups_and_notes():
    assert extract_text(lecture_deck(), "Лекция.pptx") == LECTURE_DECK_TEXT


def test_pptx_smartart_points_and_chart_title():
    assert extract_text(smartart_and_chart_deck(), "deck.pptx") == (
        "Слайд 1\nОбращение денег\nЭмиссия\nБанки и население\nДенежная масса М2"
    )


def test_pptx_without_text_and_broken_pptx():
    assert extract_text(pictures_only_deck(), "slides.pptx") == ""
    with pytest.raises(ValueError):
        extract_text(b"PK\x03\x04 not a presentation", "slides.pptx")


def test_old_ppt_is_a_clear_russian_error(api):
    with pytest.raises(UnsupportedFormat) as info:
        extract_text(b"\xd0\xcf\x11\xe0", "Лекция.ppt")
    assert str(info.value) == PPT_LEGACY_DETAIL and "PPTX или PDF" in PPT_LEGACY_DETAIL

    old = ("Лекция.ppt", b"\xd0\xcf\x11\xe0", "application/vnd.ms-powerpoint")
    response = api.post("/api/analyze", files={"file": old})
    assert response.status_code == 400 and response.json() == {"detail": PPT_LEGACY_DETAIL}
    assert_plain_russian(api.post("/api/extract", files={"file": old}), 415, PPT_LEGACY_DETAIL)
    # Список вопросов презентацией не бывает: там общий ответ о форматах.
    assert_plain_russian(api.post("/api/exams/parse", files={"file": old}), 415, exams_api.LIST_FORMAT_DETAIL)


def test_extract_reads_a_pptx_lecture_and_guesses_the_title_from_the_first_slide(api):
    response = api.post("/api/extract", files={"file": ("Презентация1.pptx", lecture_deck(), PPTX_TYPE)})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["text"] == LECTURE_DECK_TEXT and body["title_guess"] == "Лекция 3. Денежный рынок"
    assert "ocr" not in body and "pages" not in body
    assert exams_api.title_guess("Слайд 1\nЗаметки: план\nСлайд 2\nЭластичность спроса\nТекст", "x.pptx") == "Эластичность спроса"


def test_pptx_without_text_says_to_save_it_as_pdf(api, monkeypatch):
    assert_plain_russian(
        api.post("/api/extract", files={"file": ("slides.pptx", pictures_only_deck(), PPTX_TYPE)}), 422, PPTX_NO_TEXT_DETAIL
    )
    monkeypatch.setenv("GEMINI_API_KEY", "gemini-key-1")
    calls = answer_with(monkeypatch, {"summarizer": SUMMARY})
    items = frames(api.post("/api/analyze", files={"file": ("slides.pptx", pictures_only_deck(), PPTX_TYPE)}).text)
    assert items == [{"type": "error", "code": "no_text", "message": PPTX_NO_TEXT_DETAIL}]
    assert calls == []


def test_analyze_accepts_pptx(api, monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "gemini-key-1")
    answer_with(monkeypatch, {"summarizer": SUMMARY})
    response = api.post(
        "/api/analyze",
        files={"file": ("Лекция.pptx", lecture_deck(), PPTX_TYPE)},
        data={"enabled_agents": json.dumps(["summarizer"])},
    )
    items = frames(response.text)
    assert kinds(items)[0] == "text_extracted" and kinds(items)[-1] == "done"
    assert items[0]["text"] == LECTURE_DECK_TEXT and "ocr" not in items[0]


# ── Когда PDF считается сканом ───────────────────────────────────────────────────


def test_thin_text_layer_rules():
    assert ocr.text_layer_is_thin("", 1)
    assert ocr.text_layer_is_thin("а" * 199, None)
    assert not ocr.text_layer_is_thin("а" * 200, None)
    assert ocr.text_layer_is_thin("а" * 239 + " \n" * 50, 3)  # пробелы не считаются: 239 < 80 × 3
    assert not ocr.text_layer_is_thin("а" * 240, 3)


def test_scan_pages_only_for_a_pdf_without_text_and_only_with_a_gemini_key(monkeypatch):
    scan = scan_pdf(3)
    assert ocr.scan_pages(scan, "scan.pdf", "") is None  # нет ключа: распознавать нечем, всё как раньше
    monkeypatch.setenv("GEMINI_API_KEY", "gemini-key-1")
    assert ocr.scan_pages(scan, "scan.pdf", "") == 3
    assert ocr.scan_pages(scan, "Скан.PDF", "12\n13\n14") == 3  # одни номера страниц — тоже скан
    assert ocr.scan_pages(scan, "notes.txt", "") is None
    typed = text_pdf(["Money demand depends on income and the interest rate. " * 3] * 2)
    text = extract_text(typed, "typed.pdf")
    assert len(text) > 300 and ocr.scan_pages(typed, "typed.pdf", text) is None
    assert ocr.scan_pages(b"%PDF-1.4 broken", "broken.pdf", "") is None


def test_a_scan_over_the_page_cap_is_a_clear_error_but_a_long_pdf_with_some_text_stays_as_is(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "gemini-key-1")
    with pytest.raises(ocr.OcrError) as info:
        ocr.scan_pages(scan_pdf(61), "book.pdf", "")
    assert info.value.status == 413 and info.value.code == "too_large"
    assert info.value.detail == (
        "В скане 61 страница, а распознать за раз можно до 60. Разделите PDF на части и загрузите их по очереди."
    )
    assert ocr.scan_pages(scan_pdf(61), "book.pdf", "Глава 1 " * 40) is None
    assert ocr.scan_pages(scan_pdf(60), "book.pdf", "") == 60


# ── Распознавание: запрос, куски, MAX_TOKENS, модели ─────────────────────────────


def test_a_short_scan_goes_inline_in_one_request_with_the_strict_prompt_and_thinking_off(monkeypatch):
    fake = FakeGemini().install(monkeypatch)
    scan = scan_pdf(3)
    text = asyncio.run(ocr.recognize_scan(scan, 3, rate_key="user:a"))
    assert text == "\n\n".join(page_text(n) for n in (1, 2, 3))  # один кусок — без отметок страниц
    [call] = fake.calls
    assert call["mime"] == "application/pdf" and call["data"] == scan  # байты в запросе, без Files API
    assert call["prompt"] == ocr.OCR_PROMPT
    for rule in ("дословно", "Заголовок", "списки", "« | »", "LaTeX", "[неразборчиво]", "без вступления"):
        assert rule in ocr.OCR_PROMPT
    config = call["config"]
    assert config.thinking_config.thinking_budget == 0
    assert config.max_output_tokens == 3 * ocr.OCR_TOKENS_PER_PAGE
    assert config.temperature == ocr.OCR_TEMPERATURE and config.response_mime_type == "text/plain"
    assert config.http_options.timeout == 120_000


def test_a_long_scan_is_split_into_chunks_of_15_pages_two_at_a_time_and_joined_with_page_markers(monkeypatch):
    async def slow(call):
        await asyncio.sleep(0.05)
        return transcribe(call)

    fake = FakeGemini(slow).install(monkeypatch)
    text = asyncio.run(ocr.recognize_scan(scan_pdf(40), 40, rate_key=None))
    assert sorted(page_numbers(call["data"]) for call in fake.calls) == [
        list(range(1, 16)), list(range(16, 31)), list(range(31, 41))
    ]
    assert fake.max_active == 2
    blocks = text.split("\n\n[Стр.")
    assert text.startswith("[Стр. 1–15]\n" + page_text(1)) and len(blocks) == 3
    assert "\n\n[Стр. 16–30]\n" + page_text(16) in text and "\n\n[Стр. 31–40]\n" + page_text(31) in text
    assert text.index(page_text(15)) < text.index("[Стр. 16–30]") < text.index(page_text(40))


def test_max_tokens_splits_a_chunk_until_it_fits_and_marks_a_page_that_never_fits(monkeypatch):
    def behaviour(call):
        pages = page_numbers(call["data"])
        if 7 in pages:
            return page_text(7)[:20], "MAX_TOKENS"  # седьмая не влезает никогда
        if len(pages) > 4:
            return page_text(pages[0]), "MAX_TOKENS"
        return transcribe(call)

    fake = FakeGemini(behaviour).install(monkeypatch)
    text = asyncio.run(ocr.recognize_scan(scan_pdf(15), 15, rate_key=None))
    asked = [page_numbers(call["data"]) for call in fake.calls]
    pages = lambda first, last: list(range(first, last + 1))  # noqa: E731
    assert asked == [
        pages(1, 15), pages(1, 7), pages(1, 3), pages(4, 7), pages(4, 5), pages(6, 7), [6], [7],
        pages(8, 15), pages(8, 11), pages(12, 15),
    ]
    for number in range(1, 16):
        if number != 7:
            assert text.count(page_text(number)) == 1
    markers = [line for line in text.splitlines() if line.startswith("[Стр.")]
    assert markers == [
        "[Стр. 1–3]", "[Стр. 4–5]", "[Стр. 6]", "[Стр. 7: распознана не полностью]", "[Стр. 8–11]", "[Стр. 12–15]"
    ]
    assert f"[Стр. 7: распознана не полностью]\n{page_text(7)[:20]}\n\n[Стр. 8–11]" in text


def test_a_page_refused_by_every_model_is_marked_and_the_rest_is_kept(monkeypatch):
    def behaviour(call):
        if 2 in page_numbers(call["data"]):
            return "", "RECITATION"
        return transcribe(call)

    FakeGemini(behaviour).install(monkeypatch, models="gemini-2.5-flash-lite,gemini-2.5-flash")
    text = asyncio.run(ocr.recognize_scan(scan_pdf(3), 3, rate_key=None))
    assert "[Стр. 2: не распознана]" in text and page_text(1) in text and page_text(3) in text


def test_models_come_from_gemini_models_with_the_thinking_plan_of_each(monkeypatch):
    def behaviour(call):
        if call["model"] == "gone-model":
            return gemini_error(404, "NOT_FOUND")
        return transcribe(call)

    fake = FakeGemini(behaviour).install(monkeypatch, models="gone-model,gemini-2.0-flash")
    asyncio.run(ocr.recognize_scan(scan_pdf(3), 3, rate_key=None))
    assert [call["model"] for call in fake.calls] == ["gone-model", "gemini-2.0-flash"]
    config = fake.calls[1]["config"]
    # gemini-2.0 не думает и отвечает не длиннее 8192 токенов: три страницы упираются в этот предел.
    assert config.thinking_config is None and config.max_output_tokens == ocr.OCR_NO_THINKING_OUTPUT_CAP
    asyncio.run(ocr.recognize_scan(scan_pdf(1), 1, rate_key=None))
    assert fake.calls[2]["model"] == "gemini-2.0-flash"  # недоступная модель пропускается
    assert fake.calls[2]["config"].max_output_tokens == ocr.OCR_TOKENS_PER_PAGE

    fake = FakeGemini().install(monkeypatch, models="gemini-3-flash-preview")
    asyncio.run(ocr.recognize_scan(scan_pdf(1), 1, rate_key=None))
    config = fake.calls[0]["config"]
    assert config.temperature is None and config.thinking_config is not None  # gemini-3: температура по умолчанию

    fake = FakeGemini().install(monkeypatch, models=None)
    asyncio.run(ocr.recognize_scan(scan_pdf(1), 1, rate_key=None))
    assert fake.calls[0]["model"] == base.gemini_models_env()[0]


def test_a_model_that_rejects_the_thinking_config_is_asked_again_without_it(monkeypatch):
    def behaviour(call):
        if call["config"].thinking_config is not None:
            return gemini_error(400, "INVALID_ARGUMENT", "thinking is not supported for this model")
        return transcribe(call)

    fake = FakeGemini(behaviour).install(monkeypatch)
    assert page_text(1) in asyncio.run(ocr.recognize_scan(scan_pdf(1), 1, rate_key=None))
    assert [call["config"].thinking_config is None for call in fake.calls] == [False, True]


def test_provider_failures_are_russian_ai_errors(monkeypatch):
    FakeGemini(lambda call: gemini_error(429, "RESOURCE_EXHAUSTED")).install(monkeypatch, keys="k1,k2")
    with pytest.raises(AIError) as info:
        asyncio.run(ocr.recognize_scan(scan_pdf(2), 2, rate_key=None))
    assert info.value.status_code == 429 and info.value.detail == ocr.OCR_BUSY_DETAIL

    async def hang(call):
        await asyncio.sleep(5)

    monkeypatch.setenv("OCR_REQUEST_TIMEOUT_SECONDS", "0.05")
    FakeGemini(hang).install(monkeypatch)
    with pytest.raises(AIError) as info:
        asyncio.run(ocr.recognize_scan(scan_pdf(2), 2, rate_key=None))
    assert info.value.status_code == 504 and info.value.detail == ocr.OCR_TIMEOUT_DETAIL
    monkeypatch.delenv("OCR_REQUEST_TIMEOUT_SECONDS")

    FakeGemini(lambda call: RuntimeError("provider text")).install(monkeypatch)
    with pytest.raises(AIError) as info:
        asyncio.run(ocr.recognize_scan(scan_pdf(2), 2, rate_key=None))
    assert info.value.status_code == 502 and info.value.detail == ocr.OCR_FAILED_DETAIL


def test_a_long_scan_stops_every_chunk_after_the_first_failure(monkeypatch):
    def behaviour(call):
        if 1 in page_numbers(call["data"]):
            return gemini_error(429, "RESOURCE_EXHAUSTED")
        return transcribe(call)

    fake = FakeGemini(behaviour).install(monkeypatch)
    with pytest.raises(AIError):
        asyncio.run(ocr.recognize_scan(scan_pdf(45), 45, rate_key=None))
    assert len(fake.calls) <= 3


def test_nothing_recognised_is_a_no_text_error(monkeypatch):
    for answer in ("", "[неразборчиво]\n\n[неразборчиво]"):
        FakeGemini(lambda call, answer=answer: (answer, "STOP")).install(monkeypatch)
        with pytest.raises(ocr.OcrError) as info:
            asyncio.run(ocr.recognize_scan(scan_pdf(2), 2, rate_key=None))
        assert (info.value.status, info.value.code, info.value.detail) == (422, "no_text", ocr.OCR_EMPTY_DETAIL)


def test_a_fenced_answer_loses_its_fence():
    assert ocr._clean_answer("```markdown\n# Лекция 1\nТекст\n```") == "# Лекция 1\nТекст"
    assert ocr._clean_answer("Текст\r\n\r\n\r\n\r\nещё") == "Текст\n\nещё"


# ── Маршруты: /api/extract, /api/exams/parse, /api/analyze ────────────────────────


def test_extract_recognises_a_scan_says_so_and_spends_no_model_of_the_analysis(api, monkeypatch):
    fake = FakeGemini().install(monkeypatch)
    analysis = answer_with(monkeypatch, {})
    response = api.post("/api/extract", files={"file": ("Лекция 3 Денежный рынок.pdf", scan_pdf(3), "application/pdf")})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["ocr"] is True and body["pages"] == 3
    assert body["text"].startswith(page_text(1)) and body["length"] == len(body["text"])
    assert body["title_guess"] == "Лекция 3 Денежный рынок" and body["warnings"] == []
    assert len(fake.calls) == 1 and analysis == []


def test_extract_page_markers_never_become_the_title(api, monkeypatch):
    def behaviour(call):
        if 1 in page_numbers(call["data"]):
            return "Лекция 7. Кредитная система\n\nКредит — движение ссудного капитала.", "STOP"
        return transcribe(call)

    FakeGemini(behaviour).install(monkeypatch)
    body = api.post("/api/extract", files={"file": ("scan.pdf", scan_pdf(20), "application/pdf")}).json()
    assert body["text"].startswith("[Стр. 1–15]\nЛекция 7. Кредитная система") and "\n\n[Стр. 16–20]\n" in body["text"]
    assert body["title_guess"] == "Лекция 7. Кредитная система" and body["pages"] == 20


def test_extract_errors_of_a_scan_are_plain_russian(api, monkeypatch):
    FakeGemini().install(monkeypatch)
    too_long = api.post("/api/extract", files={"file": ("book.pdf", scan_pdf(61), "application/pdf")})
    assert_plain_russian(too_long, 413, ocr.too_many_pages_detail(61))

    FakeGemini(lambda call: gemini_error(429, "RESOURCE_EXHAUSTED")).install(monkeypatch)
    busy = api.post("/api/extract", files={"file": ("scan.pdf", scan_pdf(2), "application/pdf")})
    assert_plain_russian(busy, 429, ocr.OCR_BUSY_DETAIL)
    assert busy.json()["code"] == "ai_busy" and busy.headers["retry-after"] == "60"

    FakeGemini(lambda call: ("", "STOP")).install(monkeypatch)
    empty = api.post("/api/extract", files={"file": ("scan.pdf", scan_pdf(2), "application/pdf")})
    assert_plain_russian(empty, 422, ocr.OCR_EMPTY_DETAIL)

    monkeypatch.delenv("GEMINI_API_KEY")  # без ключа — как раньше: скан без текста
    assert_plain_russian(
        api.post("/api/extract", files={"file": ("scan.pdf", scan_pdf(2), "application/pdf")}), 422, exams_api.NO_TEXT_DETAIL
    )


def test_parse_recognises_a_scanned_question_list_and_adds_the_notice_to_its_warnings(api, monkeypatch):
    questions = "Вопросы к экзамену\n1. Предмет и метод макроэкономики.\n2. ВВП и способы его измерения.\n3. Инфляция и её виды."
    FakeGemini(lambda call: (questions, "STOP")).install(monkeypatch)
    response = api.post("/api/exams/parse", files={"file": ("Вопросы.pdf", scan_pdf(1), "application/pdf")})
    assert response.status_code == 200, response.text
    body = response.json()
    assert [question["text"] for question in body["questions"]][0].startswith("Предмет и метод")
    assert len(body["questions"]) == 3 and body["ocr"] is True and body["pages"] == 1
    assert {"code": "ocr", "message": ocr.OCR_NOTICE} in body["warnings"]

    plain = api.post("/api/exams/parse", data={"text": questions}).json()
    assert "ocr" not in plain and all(warning["code"] != "ocr" for warning in plain["warnings"])


def test_analyze_stream_recognises_a_scan_before_the_analysis(api, monkeypatch):
    FakeGemini().install(monkeypatch)
    calls = answer_with(monkeypatch, {"summarizer": SUMMARY})
    response = api.post(
        "/api/analyze",
        files={"file": ("scan.pdf", scan_pdf(2), "application/pdf")},
        data={"enabled_agents": json.dumps(["summarizer"])},
    )
    items = [item for item in frames(response.text) if item != "KA"]
    assert kinds(items)[:3] == ["status", "text_extracted", "grounding"] and kinds(items)[-1] == "done"
    assert items[0] == {"type": "status", "message": ocr.OCR_STATUS_MESSAGE}
    extracted = items[1]
    assert extracted["ocr"] is True and extracted["pages"] == 2 and extracted["text"].startswith(page_text(1))
    assert calls == ["summarizer"]


def test_analyze_stream_errors_of_a_scan(api, monkeypatch):
    FakeGemini().install(monkeypatch)
    answer_with(monkeypatch, {"summarizer": SUMMARY})
    items = frames(api.post("/api/analyze", files={"file": ("book.pdf", scan_pdf(61), "application/pdf")}).text)
    assert items == [{"type": "error", "code": "too_large", "message": ocr.too_many_pages_detail(61)}]

    FakeGemini(lambda call: gemini_error(429, "RESOURCE_EXHAUSTED")).install(monkeypatch)
    items = frames(api.post("/api/analyze", files={"file": ("scan.pdf", scan_pdf(2), "application/pdf")}).text)
    assert kinds(items) == ["status", "error"]
    assert items[-1] == {"type": "error", "code": "busy", "message": api.main.OCR_BUSY_MESSAGE}
    assert "лимит не списан" in api.main.OCR_BUSY_MESSAGE

    FakeGemini(lambda call: ("", "STOP")).install(monkeypatch)
    items = frames(api.post("/api/analyze", files={"file": ("scan.pdf", scan_pdf(2), "application/pdf")}).text)
    assert items[-1] == {"type": "error", "code": "no_text", "message": ocr.OCR_EMPTY_DETAIL}


def test_ocr_has_its_own_per_user_limit_and_ordinary_files_do_not_count(api, monkeypatch):
    FakeGemini().install(monkeypatch)
    monkeypatch.setattr(api.main.limiter, "enabled", True)
    api.main.limiter.reset()
    scan = ("scan.pdf", scan_pdf(1), "application/pdf")
    lecture = ("lecture.txt", ("Деньги — особый товар. " * 30).encode(), "text/plain")
    try:
        assert api.post("/api/extract", files={"file": lecture}).status_code == 200
        codes = [api.post("/api/extract", files={"file": scan}).status_code for _ in range(6)]
        limited = api.post("/api/extract", files={"file": scan})
        other_user = api.post("/api/extract", files={"file": scan}, headers={"Authorization": "Bearer user-b"})
    finally:
        api.main.limiter.reset()
    assert codes == [200] * 6
    assert_plain_russian(limited, 429, ocr.rate_limit_detail())
    assert int(limited.headers["retry-after"]) >= 1 and limited.json()["detail"] != RATE_LIMIT_DETAIL
    assert "6 файлов" in limited.json()["detail"]
    assert other_user.status_code == 200


# ── Фото страниц ──────────────────────────────────────────────────────────────────


def test_a_photo_is_recognised_inline_with_the_ocr_prompt_and_marked_as_recognised(api, monkeypatch):
    board = (
        "Лекция 5. Инфляция\nИнфляция — рост общего уровня цен. Её измеряют индексом потребительских цен, "
        "который сравнивает стоимость одной и той же корзины товаров в разные годы. Центральный банк сдерживает "
        "инфляцию ключевой ставкой, а высокая инфляция обесценивает сбережения."
    )
    fake = FakeGemini(lambda call: (board, "STOP")).install(monkeypatch, models="model-from-env")
    calls = answer_with(monkeypatch, {"summarizer": SUMMARY})
    response = api.post(
        "/api/analyze/media",
        files={"file": ("доска.png", png_bytes(), "image/png")},
        data={"enabled_agents": json.dumps(["summarizer"])},
    )
    items = [item for item in frames(response.text) if item != "KA"]
    assert kinds(items)[:2] == ["status", "text_extracted"] and kinds(items)[-1] == "done"
    assert items[1]["ocr"] is True and items[1]["pages"] == 1 and items[1]["text"].startswith("Лекция 5. Инфляция")
    [call] = fake.calls
    assert call["mime"] == "image/png" and call["data"] == png_bytes() and call["prompt"] == ocr.OCR_PROMPT
    # Незнакомое семейство моделей думает по умолчанию: к ответу добавлен запас на мысли.
    assert call["model"] == "model-from-env"
    assert call["config"].max_output_tokens == ocr.OCR_IMAGE_OUTPUT_TOKENS + base.THINKING_ALLOWANCE
    assert calls == ["summarizer"]


def test_a_photo_type_comes_from_its_bytes_first(api, monkeypatch):
    jpeg = b"\xff\xd8\xff\xe0 jfif photo"
    # Клиент пережимает фото в JPEG и оставляет прежнее имя: решают байты, а не «.png».
    assert ocr.image_mime("photo.png", "image/jpeg", jpeg) == "image/jpeg"
    assert ocr.image_mime("scan.jpg", "application/octet-stream", png_bytes()) == "image/png"
    assert ocr.image_mime("scan.jpg", "application/octet-stream", b"not an image") == "image/jpeg"
    assert ocr.image_mime("talk.mp4", "video/mp4", b"\x00\x00\x00\x18ftypmp42") is None

    fake = FakeGemini(lambda call: ("Текст с доски", "STOP")).install(monkeypatch)
    answer_with(monkeypatch, {"summarizer": SUMMARY})
    api.post("/api/analyze/media", files={"file": ("photo.png", jpeg, "image/jpeg")}, data={"enabled_agents": '["summarizer"]'})
    assert [call["mime"] for call in fake.calls] == ["image/jpeg"]


def test_a_photo_without_text_or_over_the_limit_is_a_readable_error(api, monkeypatch):
    FakeGemini(lambda call: ("", "STOP")).install(monkeypatch)
    items = frames(api.post("/api/analyze/media", files={"file": ("photo.jpg", b"\xff\xd8\xff photo", "image/jpeg")}).text)
    assert items[-1] == {"type": "error", "code": "no_text", "message": ocr.OCR_EMPTY_DETAIL}

    FakeGemini().install(monkeypatch)
    monkeypatch.setattr(api.main.limiter, "enabled", True)
    api.main.limiter.reset()
    try:
        for _ in range(6):
            ocr.check_ocr_rate("user:user-a")
        items = frames(api.post("/api/analyze/media", files={"file": ("photo.jpg", b"\xff\xd8\xff", "image/jpeg")}).text)
    finally:
        api.main.limiter.reset()
    assert items[-1] == {"type": "error", "code": "busy", "message": ocr.rate_limit_detail()}


def test_media_parser_reads_a_photo_inline_and_takes_models_from_env(monkeypatch, tmp_path):
    fake = FakeGemini(lambda call: ("Текст с фото", "STOP")).install(monkeypatch, models="model-a")
    path = tmp_path / "page.jpg"
    path.write_bytes(b"\xff\xd8\xff photo")
    assert asyncio.run(media_parser.transcribe_media_file(str(path), "image/jpeg")) == "Текст с фото"
    assert [(call["model"], call["mime"], call["prompt"]) for call in fake.calls] == [("model-a", "image/jpeg", ocr.OCR_PROMPT)]


def test_media_parser_video_uses_gemini_models_from_env(monkeypatch, tmp_path):
    monkeypatch.setenv("GEMINI_API_KEY", "key-one")
    monkeypatch.setenv("GEMINI_MODELS", "env-model-1,env-model-2")
    models = []

    class FakeFiles:
        def upload(self, file=None, config=None):
            return SimpleNamespace(name="files/1", state=SimpleNamespace(name="ACTIVE"))

        def delete(self, name=None):
            pass

    class FakeClient:
        def __init__(self, api_key=None, **kwargs):
            self.files = FakeFiles()

            async def generate_content(model=None, contents=None, **kwargs):
                models.append(model)
                assert contents[1] == media_parser.MEDIA_TRANSCRIPTION_PROMPT
                if model == "env-model-1":
                    raise RuntimeError("down")
                return SimpleNamespace(text="Расшифровка лекции")

            self.aio = SimpleNamespace(models=SimpleNamespace(generate_content=generate_content))

    monkeypatch.setattr(media_parser.genai, "Client", FakeClient)
    path = tmp_path / "talk.mp4"
    path.write_bytes(b"video")
    assert asyncio.run(media_parser.transcribe_media_file(str(path), "video/mp4")) == "Расшифровка лекции"
    assert models == ["env-model-1", "env-model-2"]
