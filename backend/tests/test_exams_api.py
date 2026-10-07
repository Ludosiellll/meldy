"""«Билеты» API (exams_api.py): подключение к приложению, 401 без сессии, лимиты по
пользователю, ошибки по-русски без текста исключений, /api/extract без ИИ, поток SSE
разметки с keepalive и снятие блокировки, если поток так и не начался.

Supabase — подменённый транспорт httpx из test_exams_map, модель — подменённый BaseAgent.run.
"""
import asyncio
import io
import json
import re
import sys
from pathlib import Path

import httpx
import pytest
from docx import Document
from fastapi import HTTPException
from PyPDF2 import PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import exams_api  # noqa: E402
import exams_map  # noqa: E402
import supabase_rest  # noqa: E402
from agents.base import AIError, BaseAgent  # noqa: E402
from ratelimit import RATE_LIMIT_DETAIL, user_or_ip_key  # noqa: E402
from test_exams_map import (  # noqa: E402
    ANON,
    Q_BACKPROP,
    Q_PHILLIPS,
    Q_POLICY,
    RULES,
    TOKENS,
    URL,
    Clock,
    FakeSupabase,
    Judge,
    course,
    honest,
    pg_ts,
    T0,
)

LIST_TEXT = "Вопросы к экзамену\nРаздел 1. Введение\n1. Предмет и метод макроэкономики.\n2. ВВП и способы его измерения.\n"


def assert_plain_russian(response, status, detail=None):
    assert response.status_code == status, response.text
    body = response.json()
    assert set(body) <= {"detail", "code"}
    assert re.search(r"[А-Яа-яЁё]", body["detail"]) and len(body["detail"]) <= 200
    assert not re.search(r"Error|Exception|Traceback|http|Ошибка чтения|internal", body["detail"])
    if detail is not None:
        assert body["detail"] == detail


def docx_bytes(paragraphs=(), table=None):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        grid = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                grid.cell(r, c).text = cell
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def blank_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


@pytest.fixture
def api(monkeypatch):
    import main
    from fastapi.testclient import TestClient

    async def fake_user(request):
        token = request.headers.get("authorization", "").removeprefix("Bearer ").strip()
        if token not in TOKENS:
            raise HTTPException(status_code=401, detail="Требуется войти в аккаунт")
        return {"id": TOKENS[token]}

    monkeypatch.setattr(main, "require_authenticated_user", fake_user)
    monkeypatch.setattr(main.limiter, "enabled", False)
    monkeypatch.setenv("SUPABASE_URL", URL)
    monkeypatch.setenv("SUPABASE_KEY", ANON)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("MAX_UPLOAD_BYTES", raising=False)
    monkeypatch.setattr(exams_map, "STAGGER_SECONDS", 0.0)
    clock = Clock()
    monkeypatch.setattr(exams_map, "_now", clock)
    fake = FakeSupabase()
    fake.clock = clock
    monkeypatch.setattr(supabase_rest, "_transport", httpx.MockTransport(fake.handler))
    model = Judge(honest(RULES))

    async def fake_run(agent, system_prompt, user_text, **kwargs):
        return await model.run(agent, system_prompt, user_text, **kwargs)

    monkeypatch.setattr(BaseAgent, "run", fake_run)
    BaseAgent._run_cache.clear()
    client = TestClient(main.app, headers={"Authorization": "Bearer token-a"})
    client.fake = fake
    client.judge = model
    client.main = main
    yield client
    BaseAgent._run_cache.clear()


# ── Подключение, сессия, лимиты ─────────────────────────────────────────────────


def test_routes_are_declared_with_full_paths_and_mounted():
    import main

    declared = {(route.path, frozenset(route.methods)) for route in exams_api.router.routes}
    assert declared == {
        ("/api/exams/parse", frozenset({"POST"})),
        ("/api/extract", frozenset({"POST"})),
        ("/api/exams/lecture-fit", frozenset({"POST"})),
        ("/api/exams/{exam_id}/map", frozenset({"POST"})),
    }
    assert main.exams_router is exams_api.router


def test_every_endpoint_is_401_without_a_session(monkeypatch):
    import main
    from fastapi.testclient import TestClient

    fake = FakeSupabase()
    monkeypatch.setattr(supabase_rest, "_transport", httpx.MockTransport(fake.handler))
    client = TestClient(main.app)
    exam_id = "6f1c1d3e-2a6b-4f0e-9a51-1c2d3e4f5a6b"
    responses = [
        client.post("/api/exams/parse", data={"text": LIST_TEXT}),
        client.post("/api/extract", files={"file": ("lecture.txt", "Лекция".encode(), "text/plain")}),
        client.post(f"/api/exams/{exam_id}/map", json={"question_ids": None, "force": False}),
        client.post("/api/exams/lecture-fit", json={"questions": ["Инфляция: виды и причины."]}),
    ]
    for response in responses:
        assert_plain_russian(response, 401, "Требуется войти в аккаунт")
    assert fake.requests == []


def test_limits_are_counted_per_user_with_user_or_ip_key(api, monkeypatch):
    limits = {
        name.rsplit(".", 1)[-1]: [(str(limit.limit), limit.key_func) for limit in values]
        for name, values in api.main.limiter._route_limits.items()
        if name.startswith("exams_api.")
    }
    assert limits == {
        "parse_exam_questions": [("10 per 1 minute", user_or_ip_key)],
        "extract_lecture_text": [("10 per 1 minute", user_or_ip_key)],
        "map_exam": [("3 per 1 minute", user_or_ip_key)],
        "fit_lectures": [("20 per 1 minute", user_or_ip_key)],
    }

    monkeypatch.setattr(api.main.limiter, "enabled", True)
    api.main.limiter.reset()
    try:
        for _ in range(10):
            assert api.post("/api/exams/parse", data={"text": LIST_TEXT}).status_code == 200
        limited = api.post("/api/exams/parse", data={"text": LIST_TEXT})
        assert_plain_russian(limited, 429, RATE_LIMIT_DETAIL)
        other_user = api.post("/api/exams/parse", data={"text": LIST_TEXT}, headers={"Authorization": "Bearer token-b"})
        assert other_user.status_code == 200
    finally:
        api.main.limiter.reset()


# ── POST /api/exams/parse ───────────────────────────────────────────────────────


def test_parse_pasted_text_returns_the_preview_and_writes_nothing(api):
    response = api.post("/api/exams/parse", data={"text": LIST_TEXT})
    assert response.status_code == 200
    body = response.json()
    assert body["method"] == "regex"
    assert body["questions"] == [
        {"position": 1, "label": "1", "text": "Предмет и метод макроэкономики.", "section": 1, "flags": []},
        {"position": 2, "label": "2", "text": "ВВП и способы его измерения.", "section": 1, "flags": []},
    ]
    assert body["sections"] == [{"index": 1, "title": "Раздел 1. Введение"}]
    assert body["tickets"] == [] and body["source_text"] == LIST_TEXT
    assert [warning["code"] for warning in body["warnings"]] == ["few_found"]
    assert body["skipped"] == []  # нечего показывать на месте пропущенных строк
    assert api.fake.requests == [] and api.judge.calls == []
    # Пропущенная подпись приходит списком: текст и сколько вопросов перед ней.
    signed = api.post("/api/exams/parse", data={"text": LIST_TEXT + "Зав. кафедрой ____ И. И. Иванов\n"}).json()
    assert signed["skipped"] == [{"text": "Зав. кафедрой ____ И. И. Иванов", "before": 2}]
    assert [warning["code"] for warning in signed["warnings"]] == ["unparsed_lines", "few_found"]


def test_parse_a_docx_file_with_a_ticket_table(api):
    contents = docx_bytes(table=[["Билет № 1", "1. Инфляция: виды и причины. 2. Кривая Филлипса."],
                                 ["Билет № 2", "1. Безработица. 2. Инфляция: виды и причины."]])
    response = api.post("/api/exams/parse", files={"file": ("Билеты.DOCX", contents, "application/octet-stream")})
    assert response.status_code == 200
    body = response.json()
    assert [question["text"] for question in body["questions"]] == ["Инфляция: виды и причины.", "Кривая Филлипса.", "Безработица."]
    assert body["tickets"] == [{"number": 1, "positions": [1, 2]}, {"number": 2, "positions": [3, 1]}]


def test_parse_txt_and_md_files(api):
    for name in ("Вопросы.txt", "Вопросы.md"):
        response = api.post("/api/exams/parse", files={"file": (name, LIST_TEXT.encode("utf-8"), "text/plain")})
        assert response.status_code == 200 and len(response.json()["questions"]) == 2


def test_parse_input_errors(api, monkeypatch):
    both = api.post("/api/exams/parse", data={"text": LIST_TEXT}, files={"file": ("q.txt", b"1. A", "text/plain")})
    assert_plain_russian(both, 400, exams_api.BOTH_INPUTS_DETAIL)
    assert_plain_russian(api.post("/api/exams/parse", data={"text": "  \n "}), 400, exams_api.NO_INPUT_DETAIL)
    assert_plain_russian(api.post("/api/exams/parse", data={}), 400, exams_api.NO_INPUT_DETAIL)
    assert_plain_russian(api.post("/api/exams/parse", json={"text": LIST_TEXT}), 400, exams_api.NO_INPUT_DETAIL)

    nothing = api.post("/api/exams/parse", data={"text": "Вопросы к экзамену\nПеречень вопросов\n"})
    assert_plain_russian(nothing, 400)
    assert nothing.json()["detail"].startswith("Не нашли ни одного вопроса")

    too_many = "\n".join(f"{n}. Вопрос номер {n} об экономике" for n in range(1, 152))
    assert_plain_russian(api.post("/api/exams/parse", data={"text": too_many}), 400,
                         "Можно до 150 вопросов. Разбейте список на два экзамена.")
    assert_plain_russian(api.post("/api/exams/parse", data={"text": "1. Вопрос\n" + "а" * 60_000}), 413)

    xlsx = api.post("/api/exams/parse", files={"file": ("Вопросы.xlsx", b"PK\x03\x04", "application/vnd.ms-excel")})
    assert_plain_russian(xlsx, 415, exams_api.LIST_FORMAT_DETAIL)

    monkeypatch.setenv("MAX_UPLOAD_BYTES", "1000")
    big = api.post("/api/exams/parse", files={"file": ("Вопросы.txt", ("1. Вопрос\n" * 200).encode(), "text/plain")})
    assert_plain_russian(big, 413)
    assert big.json()["detail"].startswith("Файл слишком большой")
    monkeypatch.delenv("MAX_UPLOAD_BYTES")

    scan = api.post("/api/exams/parse", files={"file": ("Вопросы.pdf", blank_pdf(), "application/pdf")})
    assert_plain_russian(scan, 422, "В файле не найден текст: похоже на скан.")
    broken = api.post("/api/exams/parse", files={"file": ("Вопросы.docx", b"not a zip at all", "application/octet-stream")})
    assert_plain_russian(broken, 422, exams_api.UNREADABLE_DETAIL)
    empty = api.post("/api/exams/parse", files={"file": ("Вопросы.txt", b"", "text/plain")})
    assert_plain_russian(empty, 422)
    assert api.fake.requests == []


# ── POST /api/extract ───────────────────────────────────────────────────────────


def test_extract_returns_text_length_title_guess_and_uses_no_ai(api):
    text = "Лекция 3. Денежный рынок\n\n" + "Денежный рынок — часть финансового рынка. " * 20
    response = api.post("/api/extract", files={"file": ("lecture-3.txt", text.encode("utf-8"), "text/plain")})
    assert response.status_code == 200
    body = response.json()
    assert body == {"text": text.strip(), "length": len(text.strip()), "title_guess": "Лекция 3. Денежный рынок", "warnings": []}
    assert api.judge.calls == [] and api.fake.requests == []


def test_extract_title_guess_falls_back_to_the_file_name_and_warns_about_little_text(api):
    long_first_line = "Очень длинная первая строка лекции без переноса " * 4
    response = api.post("/api/extract", files={"file": ("Лекция 5 Инфляция.md", f"## {long_first_line}\nТекст".encode(), "text/markdown")})
    body = response.json()
    assert response.status_code == 200 and body["title_guess"] == "Лекция 5 Инфляция"
    assert body["warnings"] == [exams_api.SHORT_TEXT_WARNING] and all(isinstance(w, str) for w in body["warnings"])
    assert exams_api.title_guess("\n  \n2024\n**Тема 1. Деньги**\n", "x.txt") == "Тема 1. Деньги"
    assert exams_api.title_guess("", "Лекция.docx") == "Лекция"


LECTURE_BODY = (
    "Денежный рынок — часть финансового рынка, на которой встречаются спрос на деньги и предложение денег, "
    "а также формируется ставка процента."
)


@pytest.mark.parametrize(
    ("head", "file_name", "expected"),
    [
        # Шапка вуза перед «Лекция 3. …»: берём строку лекции.
        ("Министерство науки и высшего образования Российской Федерации\n"
         "ФГБОУ ВО «Уральский государственный экономический университет»\nКафедра финансов и кредита\n"
         "Лекция 3. Денежный рынок", "Лекция 3 Денежный рынок.pdf", "Лекция 3. Денежный рынок"),
        ("Кафедра финансов и кредита", "Лекция 3 Денежный рынок.pdf", "Лекция 3 Денежный рынок"),
        # Название курса заглавными — не название лекции: имя файла говорит больше.
        ("МАКРОЭКОНОМИКА", "Лекция 3 Денежный рынок.pdf", "Лекция 3 Денежный рынок"),
        ("МАКРОЭКОНОМИКА\nЛекция 3. Денежный рынок", "lecture.pdf", "Лекция 3. Денежный рынок"),
        ("МАКРОЭКОНОМИКА\nДенежный рынок", "lecture3.pdf", "Денежный рынок"),
        ("ЛЕКЦИЯ 3\nДЕНЕЖНЫЙ РЫНОК", "scan.pdf", "ЛЕКЦИЯ 3. ДЕНЕЖНЫЙ РЫНОК"),
        ("КОНСПЕКТ ЛЕКЦИЙ\nпо дисциплине «Макроэкономика»\nИ. И. Иванов\nМосква – 2024", "Макроэкономика_конспект.pdf",
         "Макроэкономика конспект"),
        ("Авторское право и смежные права", "x.pdf", "Авторское право и смежные права"),
    ],
)
def test_title_guess_skips_institution_headers_and_prefers_the_lecture_line(head, file_name, expected):
    assert exams_api.title_guess(f"{head}\n{LECTURE_BODY}", file_name) == expected


def test_extract_title_guess_of_a_lecture_with_a_university_header(api):
    text = "Министерство науки и высшего образования Российской Федерации\nКафедра финансов и кредита\nМАКРОЭКОНОМИКА\n" + LECTURE_BODY * 3
    response = api.post("/api/extract", files={"file": ("Лекция 3 Денежный рынок.txt", text.encode("utf-8"), "text/plain")})
    assert response.status_code == 200 and response.json()["title_guess"] == "Лекция 3 Денежный рынок"


def test_parse_runs_off_the_event_loop(api, monkeypatch):
    """Разбор до 60 000 знаков — CPU: в потоке-исполнителе он не останавливает чужие потоки SSE."""
    seen = {}
    real = exams_api.parse_question_list

    def spy(text):
        try:
            asyncio.get_running_loop()
            seen["loop"] = True  # вызван прямо в цикле событий
        except RuntimeError:
            seen["loop"] = False  # в потоке-исполнителе цикла нет
        return real(text)

    monkeypatch.setattr(exams_api, "parse_question_list", spy)
    response = api.post("/api/exams/parse", data={"text": LIST_TEXT})
    assert response.status_code == 200 and len(response.json()["questions"]) == 2
    assert seen == {"loop": False}

    def failing(text):
        raise exams_api.ParseError(400, "Не нашли ни одного вопроса.")

    monkeypatch.setattr(exams_api, "parse_question_list", failing)
    assert_plain_russian(api.post("/api/exams/parse", data={"text": LIST_TEXT}), 400, "Не нашли ни одного вопроса.")


def test_extract_html_and_docx(api):
    html = "<html><head><script>alert(1)</script></head><body><h1>Лекция 4. Инфляция</h1><p>Инфляция — рост цен.</p></body></html>"
    response = api.post("/api/extract", files={"file": ("page.html", html.encode("utf-8"), "text/html")})
    assert response.status_code == 200
    assert response.json()["title_guess"] == "Лекция 4. Инфляция" and "alert" not in response.json()["text"]
    contents = docx_bytes(["Лекция 2. Спрос на деньги", "Спрос на деньги складывается из трёх мотивов."])
    response = api.post("/api/extract", files={"file": ("Лекция.docx", contents, "application/octet-stream")})
    assert response.status_code == 200 and response.json()["title_guess"] == "Лекция 2. Спрос на деньги"


def test_extract_replaces_lone_surrogates_from_a_broken_pdf_text_layer(api, monkeypatch):
    monkeypatch.setattr(exams_api, "extract_text", lambda contents, name: "Лекция 1. Деньги\nТекст \ud835 со сбоем." * 30)
    response = api.post("/api/extract", files={"file": ("Лекция.pdf", b"%PDF-1.4", "application/pdf")})
    assert response.status_code == 200
    text = response.json()["text"]
    assert "\ud835" not in text and "Текст ? со сбоем." in text
    text.encode("utf-8")


def test_extract_errors(api, monkeypatch):
    scan = api.post("/api/extract", files={"file": ("Скан.pdf", blank_pdf(), "application/pdf")})
    assert_plain_russian(scan, 422, "В файле не найден текст: похоже на скан.")
    assert_plain_russian(api.post("/api/extract", data={"text": "x"}), 400, exams_api.NO_FILE_DETAIL)
    assert_plain_russian(api.post("/api/extract", files={"file": ("a.exe", b"MZ", "application/octet-stream")}), 415)
    broken = api.post("/api/extract", files={"file": ("Лекция.pdf", b"%PDF-1.4 broken", "application/pdf")})
    assert_plain_russian(broken, 422, exams_api.UNREADABLE_DETAIL)
    monkeypatch.setenv("MAX_UPLOAD_BYTES", "100")
    assert_plain_russian(api.post("/api/extract", files={"file": ("a.txt", b"x" * 101, "text/plain")}), 413)
    assert api.judge.calls == []


# ── POST /api/exams/{exam_id}/map ───────────────────────────────────────────────


def sse_frames(body):
    chunks = body.split("\n\n")
    assert chunks[-1] == "", "поток должен кончаться пустой строкой"
    frames = chunks[:-1]
    for frame in frames:
        assert frame == ": keepalive" or re.fullmatch(r"data: \{.*\}", frame, flags=re.S), frame
    return frames


def events_of(frames):
    return [json.loads(frame[len("data: "):]) for frame in frames if frame.startswith("data: ")]


def test_map_streams_sse_with_keepalive_and_ends_with_done(api, monkeypatch):
    monkeypatch.setenv("SSE_KEEPALIVE_SECONDS", "0.05")
    exam_id, question_ids, _ = course(api.fake, [Q_BACKPROP, Q_PHILLIPS, Q_POLICY])
    original = BaseAgent.run

    async def slow_run(agent, system_prompt, user_text, **kwargs):
        await asyncio.sleep(0.4)
        return await original(agent, system_prompt, user_text, **kwargs)

    monkeypatch.setattr(BaseAgent, "run", slow_run)
    response = api.post(f"/api/exams/{exam_id}/map", json={"question_ids": None, "force": False})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/event-stream")
    assert response.headers["cache-control"] == "no-cache" and response.headers["x-accel-buffering"] == "no"
    frames = sse_frames(response.text)
    events = events_of(frames)
    assert [event["type"] for event in events] == ["start", "lexical", "batch", "done"]
    assert ": keepalive" in frames[frames.index(next(f for f in frames if '"lexical"' in f)):]
    assert frames[-1].startswith("data: ") and json.loads(frames[-1][6:])["type"] == "done"
    done = events[-1]
    assert done["state"] == "done" and done["used_ai"] is True and done["counts"]["full"] == 2
    assert all(request.headers["authorization"] == "Bearer token-a" for request in api.fake.requests)
    assert api.fake.row("exams", exam_id)["mapping_state"] == "done"


def test_map_with_an_empty_body_and_question_ids(api):
    exam_id, (q_backprop, q_policy), _ = course(api.fake, [Q_BACKPROP, Q_POLICY])
    response = api.post(f"/api/exams/{exam_id}/map", content=b"", headers={"Content-Type": "application/json"})
    assert response.status_code == 200 and events_of(sse_frames(response.text))[-1]["type"] == "done"
    exam_id, (q_backprop, q_policy), _ = course(api.fake, [Q_BACKPROP, Q_POLICY])
    response = api.post(f"/api/exams/{exam_id}/map", json={"question_ids": [q_policy, "not-a-uuid"], "force": False})
    events = events_of(sse_frames(response.text))
    assert events[0]["to_check"] == 1
    assert [item["question_id"] for event in events for item in event.get("items", [])] == [q_policy]


def test_map_errors_before_the_stream_are_plain_json(api):
    fake = api.fake
    post = lambda exam_id, body=None: api.post(f"/api/exams/{exam_id}/map", json=body or {"question_ids": None})  # noqa: E731
    assert_plain_russian(post("not-a-uuid"), 404, "Экзамен не найден")
    foreign, _, _ = course(fake, [Q_BACKPROP], user="user-b")
    assert_plain_russian(post(foreign), 404, "Экзамен не найден")
    running, _, _ = course(fake, [Q_BACKPROP], mapping_state="running", mapping_started_at=T0)
    assert_plain_russian(post(running), 409, "Сверка с лекциями уже идёт")
    assert_plain_russian(post(fake.exam()), 400, "В экзамене нет вопросов")
    scans, _, _ = course(fake, [Q_BACKPROP], materials=(("Скан", " "),))
    response = post(scans)
    assert_plain_russian(response, 422)
    assert response.json()["detail"].startswith("Ни у одной лекции нет текста")
    huge, _, _ = course(fake, [Q_BACKPROP], materials=(("Огромная", "слово " * 500_001),))
    assert_plain_russian(post(huge), 413)

    bad_json = api.post(f"/api/exams/{running}/map", content=b"{oops", headers={"Content-Type": "application/json"})
    assert_plain_russian(bad_json, 400, exams_api.BAD_BODY_DETAIL)
    assert_plain_russian(post(running, {"question_ids": "все"}), 400, exams_api.BAD_BODY_DETAIL)
    assert_plain_russian(post(running, {"question_ids": ["6f1c1d3e-2a6b-4f0e-9a51-1c2d3e4f5a6b"] * 301}), 400)
    assert api.judge.calls == []
    assert fake.row("exams", running)["mapping_state"] == "running"  # чужую блокировку не трогали


def test_map_provider_failures_reach_the_stream_as_plain_russian(api):
    exam_id, _, _ = course(api.fake, [Q_BACKPROP, Q_POLICY])
    api.judge.script = lambda questions, call_number: AIError.busy()
    response = api.post(f"/api/exams/{exam_id}/map", json={"question_ids": None, "force": False})
    events = events_of(sse_frames(response.text))
    assert [event["type"] for event in events] == ["start", "batch_failed", "done"]
    assert events[1]["message"] == "Сервис занят — эти вопросы проверены только по словам"
    assert events[-1]["state"] == "partial" and events[-1]["used_ai"] is False


def _map_request(exam_id):
    from starlette.requests import Request

    body = json.dumps({"question_ids": None, "force": False}).encode()
    delivered = False

    async def receive():
        nonlocal delivered
        if not delivered:
            delivered = True
            return {"type": "http.request", "body": body, "more_body": False}
        return {"type": "http.disconnect"}

    scope = {
        "type": "http", "asgi": {"version": "3.0", "spec_version": "2.3"}, "http_version": "1.1", "method": "POST",
        "scheme": "http", "path": f"/api/exams/{exam_id}/map", "raw_path": f"/api/exams/{exam_id}/map".encode(),
        "query_string": b"", "root_path": "", "client": ("127.0.0.1", 5000), "server": ("testserver", 80),
        "headers": [(b"authorization", b"Bearer token-a"), (b"content-type", b"application/json")],
    }
    return scope, receive, Request(scope, receive)


def test_a_client_that_leaves_at_once_never_leaves_the_lock_behind(api):
    """Клиент ушёл сразу: Starlette отменяет поток, блокировка снимается (finally потока или фоновая задача)."""
    exam_id, _, _ = course(api.fake, [Q_BACKPROP], mapping_state="done", mapped_at=T0)

    async def go():
        scope, receive, request = _map_request(exam_id)
        response = await exams_api.map_exam(exam_id=exam_id, request=request)
        assert api.fake.row("exams", exam_id)["mapping_state"] == "running"
        # The background releases a lock the stream never took and waits for a refund in flight.
        assert response.background is not None and response.background.func is exams_api._after_mapping

        async def send(message):
            return None

        await response(scope, receive, send)
        await asyncio.sleep(0.05)

    asyncio.run(go())
    assert api.fake.row("exams", exam_id)["mapping_state"] in ("done", "failed")
    assert api.judge.calls == []


def test_a_stream_that_never_starts_gets_the_previous_state_back(api):
    exam_id, _, _ = course(api.fake, [Q_BACKPROP], mapping_state="done", mapped_at=T0)

    async def go():
        _, _, request = _map_request(exam_id)
        response = await exams_api.map_exam(exam_id=exam_id, request=request)
        assert api.fake.row("exams", exam_id)["mapping_state"] == "running"
        await response.background()  # тело ответа так и не читали

    asyncio.run(go())
    exam = api.fake.row("exams", exam_id)
    assert exam["mapping_state"] == "done" and exam["mapped_at"] == pg_ts(T0)
    assert api.judge.calls == []
