"""Разбор списка вопросов к экзамену правилами (exams_parse.py) и зеркало хэша вопроса."""
import hashlib
import io
import re
import sys
import time
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from exams_parse import (  # noqa: E402
    MAX_QUESTIONS,
    MAX_SKIPPED_LINES,
    NO_QUESTIONS_DETAIL,
    SKIPPED_LINE_CHARS,
    TOO_LONG_DETAIL,
    TOO_MANY_DETAIL,
    ParseError,
    normalize_question,
    parse_question_list,
    question_hash,
)
from parser import extract_text  # noqa: E402


def compact(result):
    """Вопросы как (позиция, номер, текст, раздел, флаги) — для точного сравнения."""
    return [(q["position"], q["label"], q["text"], q["section"], q["flags"]) for q in result["questions"]]


def codes(result):
    return [warning["code"] for warning in result["warnings"]]


# ── Зеркало SQL: meldy_normalize_question / meldy_question_hash (контракт §1.1) ───


@pytest.mark.parametrize(
    ("text", "normalized"),
    [
        ("17. Денежно-кредитная политика ЦБ: цели, инструменты!", "денежно кредитная политика цб цели инструменты"),
        ("Вопрос 3) Ёлка и ЕЛЬ", "елка и ель"),
        ("  1.2. Предмет и метод макроэкономики.", "предмет и метод макроэкономики"),
        ("1917 год: революция", "1917 год революция"),
        ("5.Инфляция: виды", "инфляция виды"),
        ("Вопрос №12. Рынок труда", "рынок труда"),
    ],
)
def test_normalize_question_matches_the_sql_vectors(text, normalized):
    assert normalize_question(text) == normalized


def test_question_hash_matches_the_sql_md5_vector():
    assert question_hash("Предмет и метод макроэкономики") == "eee8c9ac62e0050c8e2065db1ebed8cc"
    assert question_hash("1. Предмет и метод макроэкономики!") == "eee8c9ac62e0050c8e2065db1ebed8cc"
    assert question_hash("Предмет и метод макроэкономики") == hashlib.md5("предмет и метод макроэкономики".encode()).hexdigest()


def test_normalize_uses_the_explicit_letter_table_and_ascii_spaces():
    # Только A–Z, А–Я и Ё опускаются; прочие буквы — разделители, как в SQL.
    assert normalize_question("ÄBC Straße") == "bc stra e"
    assert normalize_question("ВОПРОС 7. Ёж") == "еж"
    # Номер снимается только с обычными пробелами перед ним.
    assert normalize_question(" 1. Тема") == "1 тема"
    assert normalize_question("1.  Тема") == "тема"
    # Без точки или скобки после номера номер не снимается.
    assert normalize_question("12 стульев") == "12 стульев"
    assert normalize_question("Вопрос 1234. Тема") == "вопрос 1234 тема"


# ── Пять форматов списка ────────────────────────────────────────────────────────


def test_sections_with_numbers_and_a_wrapped_line():
    result = parse_question_list(
        "Вопросы к экзамену по дисциплине «Макроэкономика»\n"
        "Раздел 1. Введение\n"
        "1. Предмет и метод макроэкономики.\n"
        "2. Система национальных счетов. ВВП и способы\n"
        "его измерения.\n"
        "Раздел 2. Денежный рынок\n"
        "3) Спрос на деньги.\n"
        "4) Предложение денег и денежный мультипликатор.\n"
    )
    assert result["method"] == "regex"
    assert compact(result) == [
        (1, "1", "Предмет и метод макроэкономики.", 1, []),
        (2, "2", "Система национальных счетов. ВВП и способы его измерения.", 1, ["joined_lines"]),
        (3, "3", "Спрос на деньги.", 2, []),
        (4, "4", "Предложение денег и денежный мультипликатор.", 2, []),
    ]
    assert result["sections"] == [
        {"index": 1, "title": "Раздел 1. Введение"},
        {"index": 2, "title": "Раздел 2. Денежный рынок"},
    ]
    assert result["tickets"] == [] and result["warnings"] == []


def test_ticket_blocks_with_local_numbering_and_a_repeated_question():
    result = parse_question_list(
        "ЭКЗАМЕНАЦИОННЫЙ БИЛЕТ № 1\n"
        "Дисциплина: Макроэкономика\n"
        "1. Инфляция: виды и причины.\n"
        "2. Кривая Филлипса.\n"
        "Билет № 2\n"
        "1. Безработица.\n"
        "2. Инфляция: виды и причины.\n"
        "Билет №3\n"
        "1. Мультипликатор автономных расходов.\n"
        "2. Кривая Филлипса!\n"
    )
    assert compact(result) == [
        (1, "", "Инфляция: виды и причины.", None, ["duplicate_removed"]),
        (2, "", "Кривая Филлипса.", None, ["duplicate_removed"]),
        (3, "", "Безработица.", None, []),
        (4, "", "Мультипликатор автономных расходов.", None, []),
    ]
    assert result["tickets"] == [
        {"number": 1, "positions": [1, 2]},
        {"number": 2, "positions": [3, 1]},
        {"number": 3, "positions": [4, 2]},
    ]


def test_unnumbered_lines_skip_the_header():
    result = parse_question_list(
        "Перечень вопросов к зачёту\n"
        "Предмет и метод макроэкономики.\n"
        "• Система национальных счетов.\n"
        "- Денежный рынок и его структура.\n"
    )
    assert compact(result) == [
        (1, "", "Предмет и метод макроэкономики.", None, []),
        (2, "", "Система национальных счетов.", None, []),
        (3, "", "Денежный рынок и его структура.", None, []),
    ]
    assert result["warnings"] == []


def test_question_n_prefix():
    result = parse_question_list(
        "Вопрос 1. Понятие нейронной сети.\n"
        "Вопрос 2. Функции активации.\n"
        "Вопрос №3. Метод обратного распространения ошибки.\n"
        "ВОПРОС 4) Регуляризация.\n"
    )
    assert compact(result) == [
        (1, "1", "Понятие нейронной сети.", None, []),
        (2, "2", "Функции активации.", None, []),
        (3, "3", "Метод обратного распространения ошибки.", None, []),
        (4, "4", "Регуляризация.", None, []),
    ]


def test_table_rows_numbered_questions_and_ticket_rows():
    numbered = parse_question_list(
        "№ п/п | Формулировка вопроса | Код компетенции\n"
        "1. | Предмет и метод макроэкономики | ОПК-1\n"
        "2 | Система национальных счетов | ОПК-2, УК-1\n"
        "3 | Денежный рынок | УК-1\n"
    )
    assert compact(numbered) == [
        (1, "1", "Предмет и метод макроэкономики", None, []),
        (2, "2", "Система национальных счетов", None, []),
        (3, "3", "Денежный рынок", None, []),
    ]
    tickets = parse_question_list(
        "Билет № 1 | 1. Инфляция: виды и причины. 2. Кривая Филлипса.\n"
        "Билет № 2 | 1. Безработица. 2. Мультипликатор.\n"
        "Билет № 3 | Денежный рынок | Инфляция: виды и причины.\n"
    )
    assert compact(tickets) == [
        (1, "", "Инфляция: виды и причины.", None, ["duplicate_removed"]),
        (2, "", "Кривая Филлипса.", None, []),
        (3, "", "Безработица.", None, []),
        (4, "", "Мультипликатор.", None, []),
        (5, "", "Денежный рынок", None, []),
    ]
    assert tickets["tickets"] == [
        {"number": 1, "positions": [1, 2]},
        {"number": 2, "positions": [3, 4]},
        {"number": 3, "positions": [5, 1]},
    ]


def test_ticket_table_with_a_section_column_keeps_global_numbers():
    # Строки таблицы под «Билет № 1»: раздел (объединённая ячейка) | номер | вопрос.
    result = parse_question_list("Билет № 1\nРаздел 1 | 1 | Предмет макроэкономики\n21 | Инфляция и её виды\n")
    assert compact(result) == [
        (1, "1", "Предмет макроэкономики", 1, []),
        (2, "21", "Инфляция и её виды", 1, []),
    ]
    assert result["sections"] == [{"index": 1, "title": "Раздел 1"}]
    assert result["tickets"] == [{"number": 1, "positions": [1, 2]}]


def test_ticket_table_with_a_header_row():
    result = parse_question_list(
        "Билет | Вопрос 1 | Вопрос 2\n1 | Инфляция и её виды | Кривая Филлипса\n2 | Безработица | Мультипликатор\n"
    )
    assert [q["text"] for q in result["questions"]] == ["Инфляция и её виды", "Кривая Филлипса", "Безработица", "Мультипликатор"]
    assert result["tickets"] == [{"number": 1, "positions": [1, 2]}, {"number": 2, "positions": [3, 4]}]


def test_tickets_that_reference_question_numbers():
    result = parse_question_list(
        "1. Предмет макроэкономики\n2. Система национальных счетов\n3. Инфляция и её виды\n"
        "Билет № 1 — вопросы 1 и 3\nБилет № 2: 2, 3\n"
    )
    assert [q["label"] for q in result["questions"]] == ["1", "2", "3"]
    assert result["tickets"] == [{"number": 1, "positions": [1, 3]}, {"number": 2, "positions": [2, 3]}]


# ── DOCX: таблица и автонумерация Word ──────────────────────────────────────────


def test_docx_with_a_table_and_auto_numbering(tmp_path):
    doc = Document()
    doc.add_paragraph("Министерство науки и высшего образования")
    doc.add_paragraph("Вопросы к экзамену по дисциплине «Макроэкономика»")
    doc.add_paragraph("Раздел 1. Введение")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "№"
    table.cell(0, 1).text = "Вопрос"
    for row, question in enumerate(["Предмет и метод макроэкономики.", "ВВП и способы его измерения."], start=1):
        table.cell(row, 0).paragraphs[0].style = doc.styles["List Number"]  # номер ставит Word, текста в ячейке нет
        table.cell(row, 1).text = question
    doc.add_paragraph("Раздел 2. Деньги")
    doc.add_paragraph("Спрос на деньги.", style="List Number")
    paragraph = doc.add_paragraph("Предложение денег", style="List Number")
    paragraph.add_run().add_break()
    paragraph.add_run("и денежный мультипликатор.")
    doc.add_paragraph("Зав. кафедрой ____________")
    path = tmp_path / "Вопросы.docx"
    doc.save(path)

    text = extract_text(path.read_bytes(), path.name)
    result = parse_question_list(text)
    assert compact(result) == [
        (1, "1", "Предмет и метод макроэкономики.", 1, []),
        (2, "2", "ВВП и способы его измерения.", 1, []),
        (3, "3", "Спрос на деньги.", 2, []),
        (4, "4", "Предложение денег и денежный мультипликатор.", 2, ["joined_lines"]),
    ]
    assert result["sections"] == [{"index": 1, "title": "Раздел 1. Введение"}, {"index": 2, "title": "Раздел 2. Деньги"}]
    assert codes(result) == ["unparsed_lines"]  # подпись заведующего пропущена — об этом сказано


def test_docx_ticket_table():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Билет № 1"
    table.cell(0, 1).text = "1. Инфляция: виды и причины. 2. Кривая Филлипса."
    table.cell(1, 0).text = "Билет № 2"
    table.cell(1, 1).text = "1. Безработица. 2. Инфляция: виды и причины."
    buffer = io.BytesIO()
    doc.save(buffer)
    result = parse_question_list(extract_text(buffer.getvalue(), "Билеты.docx"))
    assert [q["text"] for q in result["questions"]] == ["Инфляция: виды и причины.", "Кривая Филлипса.", "Безработица."]
    assert result["tickets"] == [{"number": 1, "positions": [1, 2]}, {"number": 2, "positions": [3, 1]}]


# ── Дубли, склейка, флаги, предупреждения ──────────────────────────────────────


def test_normalized_duplicates_become_one_question():
    result = parse_question_list("17. Инфляция\n18. Безработица и её виды\nИнфляция!\n")
    assert compact(result) == [
        (1, "17", "Инфляция", None, ["duplicate_removed"]),
        (2, "18", "Безработица и её виды", None, []),
    ]


def test_continuations_sub_items_and_flattened_lists():
    result = parse_question_list(
        "1. Функции денег:\n"
        "1) мера стоимости\n"
        "2) средство обращения\n"
        "2. Денежные агрегаты\n"
        "(М0, М1, М2)\n"
        "3. Предмет и метод. 4. ВВП и способы его измерения. 5. Инфляция: виды и причины.\n"
    )
    assert compact(result) == [
        (1, "1", "Функции денег: 1) мера стоимости; 2) средство обращения", None, ["joined_lines"]),
        (2, "2", "Денежные агрегаты (М0, М1, М2)", None, ["joined_lines"]),
        (3, "3", "Предмет и метод.", None, []),
        (4, "4", "ВВП и способы его измерения.", None, []),
        (5, "5", "Инфляция: виды и причины.", None, []),
    ]


def test_pdf_line_breaks_after_a_hyphen_are_joined_without_a_space():
    result = parse_question_list("1. Денежно-кредитная политика: цели и инстру-\nменты.\n2. Социально-\nэкономические последствия инфляции.\n3. Безработица.\n")
    assert [q["text"] for q in result["questions"]] == [
        "Денежно-кредитная политика: цели и инстру-менты.",
        "Социально-экономические последствия инфляции.",
        "Безработица.",
    ]


def test_numbers_on_their_own_lines_label_the_next_question_only_in_sequence():
    # Таблица из PDF: номер и вопрос — отдельными строками.
    table = parse_question_list("№\nВопрос\n1\nПредмет и метод микроэкономики\n2\nКривая производственных возможностей\n3.\nСпрос и его факторы\n")
    assert [(q["label"], q["text"]) for q in table["questions"]] == [
        ("1", "Предмет и метод микроэкономики"),
        ("2", "Кривая производственных возможностей"),
        ("3", "Спрос и его факторы"),
    ]
    assert table["warnings"] == []
    # Номер страницы не по порядку не становится номером вопроса.
    pages = parse_question_list("Предмет и метод\nКривая производственных возможностей\n7\nСпрос и его факторы\n")
    assert [q["label"] for q in pages["questions"]] == ["", "", ""]
    numbered = parse_question_list("1. Предмет и метод.\n2. Кривая.\n2\n3. Спрос и его факторы.\n")
    assert [q["label"] for q in numbered["questions"]] == ["1", "2", "3"]


def test_outline_numbers_become_sections():
    result = parse_question_list("1. Введение\n1.1. Предмет и метод\n1.2. Функции теории\n2. Деньги\n2.1. Функции денег\n")
    assert compact(result) == [
        (1, "1.1", "Предмет и метод", 1, []),
        (2, "1.2", "Функции теории", 1, []),
        (3, "2.1", "Функции денег", 2, []),
    ]
    assert result["sections"] == [{"index": 1, "title": "1. Введение"}, {"index": 2, "title": "2. Деньги"}]


def test_topics_without_questions_are_the_questions():
    result = parse_question_list("Тема 1. Инфляция и её виды\nТема 2. Безработица\nТема 3. Денежный рынок\n")
    assert compact(result) == [
        (1, "1", "Инфляция и её виды", None, []),
        (2, "2", "Безработица", None, []),
        (3, "3", "Денежный рынок", None, []),
    ]
    assert result["sections"] == []


def test_topics_under_parts_are_the_questions_and_parts_stay_sections():
    result = parse_question_list(
        "Раздел 1. Введение\n"
        "Тема 1.1. Предмет и метод макроэкономики\n"
        "Тема 1.2. Система национальных счетов\n"
        "Раздел 2. Деньги\n"
        "Тема 2.1. Денежный рынок\n"
    )
    assert compact(result) == [
        (1, "1.1", "Предмет и метод макроэкономики", 1, []),
        (2, "1.2", "Система национальных счетов", 1, []),
        (3, "2.1", "Денежный рынок", 2, []),
    ]
    assert result["sections"] == [{"index": 1, "title": "Раздел 1. Введение"}, {"index": 2, "title": "Раздел 2. Деньги"}]
    assert result["warnings"] == []


def test_lowercase_plain_list_is_not_glued_together():
    result = parse_question_list("предмет и метод макроэкономики\nввп и способы его измерения\nинфляция и её виды\n")
    assert [q["text"] for q in result["questions"]] == [
        "предмет и метод макроэкономики",
        "ввп и способы его измерения",
        "инфляция и её виды",
    ]


def test_numbers_without_a_dot_need_a_capital_letter_and_order():
    result = parse_question_list("1 Предмет макроэкономики\n2 ВВП и способы его измерения\n3 вида инфляции\n4 Безработица\n")
    assert [(q["label"], q["text"]) for q in result["questions"]] == [
        ("1", "Предмет макроэкономики"),
        ("2", "ВВП и способы его измерения"),
        ("", "3 вида инфляции"),
        ("4", "Безработица"),
    ]
    assert codes(result) == ["mixed_numbering"]


def test_long_and_short_flags_and_the_1000_char_cut():
    long_text = "Сравните модели экономического роста " * 12  # ~450 знаков
    huge_text = "слово " * 400
    result = parse_question_list(f"1. ВВП\n2. {long_text}\n3. {huge_text}\n")
    questions = result["questions"]
    assert questions[0]["flags"] == ["short"]
    assert questions[1]["flags"] == ["long"]
    assert questions[2]["flags"] == ["long"] and len(questions[2]["text"]) <= 1000


def test_header_footer_and_page_numbers_are_skipped_with_a_warning():
    result = parse_question_list(
        "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РФ\n"
        "Кафедра экономической теории\n"
        "Макроэкономика\n"
        "Вопросы к экзамену\n"
        "для студентов 2 курса\n"
        "1. Предмет и метод макроэкономики.\n"
        "2. ВВП и способы его измерения.\n"
        "Стр. 1 из 2\n"
        "3. Инфляция: виды, причины, последствия.\n"
        "2\n"
        "Зав. кафедрой ____________ И. И. Иванов\n"
    )
    assert [q["text"] for q in result["questions"]] == [
        "Предмет и метод макроэкономики.",
        "ВВП и способы его измерения.",
        "Инфляция: виды, причины, последствия.",
    ]
    assert codes(result) == ["unparsed_lines"]
    assert "1 строку" in result["warnings"][0]["message"]


def test_mixed_numbering_and_few_found_warnings():
    mixed = parse_question_list("1. Предмет и метод.\n2. ВВП и его измерение.\nИнфляция и её виды\n4. Безработица.\n")
    assert codes(mixed) == ["mixed_numbering"]
    assert [q["label"] for q in mixed["questions"]] == ["1", "2", "", "4"]
    few = parse_question_list("1. Предмет и метод макроэкономики.\n")
    assert codes(few) == ["few_found"]
    assert few["warnings"][0]["message"].startswith("Нашли всего 1 вопрос.")


def test_question_marks_keep_a_preamble_line_as_a_question():
    result = parse_question_list("Что изучает макроэкономика?\n2. ВВП и способы его измерения.\n3. Инфляция.\n")
    assert [q["text"] for q in result["questions"]][0] == "Что изучает макроэкономика?"


def test_markdown_list():
    result = parse_question_list("## Раздел 1. Введение\n1. **Предмет и метод**\n2\\. ВВП и его измерение\n")
    assert compact(result) == [
        (1, "1", "Предмет и метод", 1, []),
        (2, "2", "ВВП и его измерение", 1, []),
    ]


def test_control_characters_never_reach_the_questions_or_the_source_text():
    result = parse_question_list("1. Предмет\x00 и метод\x07.\n2. ВВП\x1b и его измерение.\n\x00")
    assert [q["text"] for q in result["questions"]] == ["Предмет и метод .", "ВВП и его измерение."]
    assert "\x00" not in result["source_text"]


def test_source_text_is_returned_as_given():
    text = "1. Предмет и метод.\r\n2. ВВП.\r\n3. Инфляция.\r\n"
    assert parse_question_list(text)["source_text"] == text


# ── Официальные листы билетов, заголовки документа, переносы из PDF ────────────

OFFICIAL_TICKETS = [
    ("Предмет и метод макроэкономики.", "Инфляция: виды и причины."),
    ("Безработица: виды и последствия.", "Денежный рынок и его структура."),
    ("Кривая Филлипса.", "Мультипликатор автономных расходов."),
]
OFFICIAL_QUESTIONS = [question for pair in OFFICIAL_TICKETS for question in pair]


def official_sheet(header_on_ticket_line=False):
    lines = []
    for number, (first, second) in enumerate(OFFICIAL_TICKETS, start=1):
        lines += ["ФГБОУ ВО «Российский экономический университет имени Г. В. Плеханова»", "Кафедра экономической теории"]
        if header_on_ticket_line:
            lines.append(f"ЭКЗАМЕНАЦИОННЫЙ БИЛЕТ № {number} по дисциплине «Макроэкономика»")
        else:
            lines += [f"ЭКЗАМЕНАЦИОННЫЙ БИЛЕТ № {number}", "по дисциплине «Макроэкономика»"]
        lines += [
            "для направления подготовки 38.03.01 «Экономика»",
            f"1. {first}",
            f"2. {second}",
            "Зав. кафедрой ____________ И. И. Иванов",
            "«___» ____________ 2025 г.",
        ]
    return "\n".join(lines)


@pytest.mark.parametrize("header_on_ticket_line", [False, True])
def test_an_official_ticket_sheet_yields_only_the_questions(header_on_ticket_line):
    result = parse_question_list(official_sheet(header_on_ticket_line))
    assert [q["text"] for q in result["questions"]] == OFFICIAL_QUESTIONS
    assert result["tickets"] == [
        {"number": 1, "positions": [1, 2]},
        {"number": 2, "positions": [3, 4]},
        {"number": 3, "positions": [5, 6]},
    ]
    # Шапки, подписи и даты между билетами пропущены — об этом сказано; нумерация не смешанная.
    assert codes(result) == ["unparsed_lines"]


def test_ticket_line_remainders_that_are_headers_are_dropped():
    result = parse_question_list(
        "Билет № 1 (весенний семестр)\n1. Предмет и метод.\n2. Инфляция.\n"
        "Билет № 2. Вариант 2\n1. Безработица.\n2. Денежный рынок.\n"
        "Билет № 3 для студентов 2 курса\n1. Кривая Филлипса.\n2. Мультипликатор.\n"
    )
    assert [q["text"] for q in result["questions"]] == [
        "Предмет и метод.", "Инфляция.", "Безработица.", "Денежный рынок.", "Кривая Филлипса.", "Мультипликатор.",
    ]
    assert [t["positions"] for t in result["tickets"]] == [[1, 2], [3, 4], [5, 6]]
    # Вопрос в самой строке билета остаётся вопросом.
    inline = parse_question_list("Билет № 1. Предмет и метод макроэкономики.\nБилет № 2. Инфляция: виды и причины.\n")
    assert [q["text"] for q in inline["questions"]] == ["Предмет и метод макроэкономики.", "Инфляция: виды и причины."]


def test_a_course_name_repeated_on_every_ticket_line_is_not_a_question():
    result = parse_question_list(
        "Билет № 1 Макроэкономика\n1. Предмет и метод.\n2. Инфляция.\n"
        "Билет № 2 Макроэкономика\n1. Безработица.\n2. Денежный рынок.\n"
    )
    assert [q["text"] for q in result["questions"]] == ["Предмет и метод.", "Инфляция.", "Безработица.", "Денежный рынок."]
    assert [t["positions"] for t in result["tickets"]] == [[1, 2], [3, 4]]
    assert codes(result) == ["unparsed_lines"]


def test_an_unrecognised_header_repeated_inside_numbered_tickets_is_dropped():
    sheet = "".join(
        f"Российский экономический университет\nБилет № {n}\n1. Вопрос {n} об инфляции.\n2. Вопрос {n} о деньгах.\n"
        for n in range(1, 4)
    )
    result = parse_question_list(sheet)
    assert len(result["questions"]) == 6 and not any("университет" in q["text"] for q in result["questions"])
    assert [t["positions"] for t in result["tickets"]] == [[1, 2], [3, 4], [5, 6]]


def test_unnumbered_lines_inside_a_numbered_ticket_warn_about_mixed_numbering():
    result = parse_question_list("Билет № 1\n1. Предмет и метод.\nИнфляция и её виды\nБилет № 2\n1. Безработица.\n2. Деньги.\n")
    assert "mixed_numbering" in codes(result)


@pytest.mark.parametrize(
    "title",
    [
        "Экзаменационные билеты",
        "Экзаменационные билеты по дисциплине «Макроэкономика»",
        "Билеты к экзамену по макроэкономике",
        "Примерные вопросы к экзамену",
        "Темы для подготовки к экзамену",
    ],
)
def test_a_document_title_is_not_question_1(title):
    tickets = parse_question_list(f"{title}\nБилет № 1\n1. Предмет и метод.\n2. Инфляция.\nБилет № 2\n1. Безработица.\n2. Деньги.\n")
    assert [q["text"] for q in tickets["questions"]] == ["Предмет и метод.", "Инфляция.", "Безработица.", "Деньги."]
    assert [t["positions"] for t in tickets["tickets"]] == [[1, 2], [3, 4]]
    plain = parse_question_list(f"{title}\nПредмет и метод макроэкономики.\nСистема национальных счетов.\nИнфляция и её виды.\n")
    assert [q["text"] for q in plain["questions"]] == [
        "Предмет и метод макроэкономики.", "Система национальных счетов.", "Инфляция и её виды.",
    ]


def test_a_course_name_in_capitals_over_an_unnumbered_list_is_skipped_with_a_warning():
    result = parse_question_list("МАКРОЭКОНОМИКА\nПредмет и метод макроэкономики.\nСистема национальных счетов.\nИнфляция и её виды.\n")
    assert [q["text"] for q in result["questions"]] == [
        "Предмет и метод макроэкономики.", "Система национальных счетов.", "Инфляция и её виды.",
    ]
    assert codes(result) == ["unparsed_lines"]
    # Аббревиатура и список, набранный заглавными целиком, остаются вопросами.
    assert len(parse_question_list("ВВП\nИнфляция и её виды.\nБезработица.\n")["questions"]) == 3
    assert len(parse_question_list("ИНФЛЯЦИЯ\nБЕЗРАБОТИЦА\nДЕНЕЖНЫЙ РЫНОК\n")["questions"]) == 3


def test_a_course_name_before_the_first_ticket_is_skipped_with_a_warning():
    result = parse_question_list("Макроэкономика\nБилет № 1\n1. Предмет и метод.\n2. Инфляция.\nБилет № 2\n1. Безработица.\n2. Деньги.\n")
    assert [q["text"] for q in result["questions"]] == ["Предмет и метод.", "Инфляция.", "Безработица.", "Деньги."]
    assert codes(result) == ["unparsed_lines"] and "1 строку" in result["warnings"][0]["message"]
    # Ссылки на номера вопросов: строки до билетов — это сам список, его не трогаем.
    refs = parse_question_list("1. Предмет.\n2. Инфляция.\n3. Деньги.\nБилет № 1 — вопросы 1 и 3\n")
    assert len(refs["questions"]) == 3 and refs["tickets"] == [{"number": 1, "positions": [1, 3]}]


def test_pdf_wraps_with_a_capital_letter_are_joined_when_numbers_continue():
    result = parse_question_list(
        "1. Предмет и метод макроэкономики. Макроэкономические агрегаты.\n"
        "2. Система национальных счетов. Основные макроэкономические показатели и их\n"
        "ВВП, ВНД: способы измерения.\n"
        "3. Инфляция: сущность, виды, причины и социально-экономические последствия.\n"
        "4. Цели и инструменты денежно-кредитной политики. Денежно-кредитная Политика\n"
        "Центрального банка РФ.\n"
        "5. Безработица: виды, причины и последствия. Закон Оукена.\n"
    )
    assert compact(result) == [
        (1, "1", "Предмет и метод макроэкономики. Макроэкономические агрегаты.", None, []),
        (2, "2", "Система национальных счетов. Основные макроэкономические показатели и их ВВП, ВНД: способы измерения.",
         None, ["joined_lines"]),
        (3, "3", "Инфляция: сущность, виды, причины и социально-экономические последствия.", None, []),
        (4, "4", "Цели и инструменты денежно-кредитной политики. Денежно-кредитная Политика Центрального банка РФ.",
         None, ["joined_lines"]),
        (5, "5", "Безработица: виды, причины и последствия. Закон Оукена.", None, []),
    ]
    assert result["warnings"] == []
    # Перенос последнего вопроса: строка дошла до края, следующего номера нет.
    last = parse_question_list(
        "1. Предмет и метод макроэкономики. Макроэкономические агрегаты.\n"
        "2. Цели и инструменты денежно-кредитной политики. Денежно-кредитная Политика\n"
        "Центрального банка РФ.\n"
    )
    assert [q["text"] for q in last["questions"]][-1] == (
        "Цели и инструменты денежно-кредитной политики. Денежно-кредитная Политика Центрального банка РФ."
    )


def test_capitalised_sub_items_after_a_colon_stay_in_their_question():
    result = parse_question_list(
        "1. Предмет макроэкономики.\n2. Функции денег:\n- Мера стоимости;\n- Средство обращения;\n- Средство накопления.\n"
        "3. Денежные агрегаты.\n"
    )
    assert compact(result) == [
        (1, "1", "Предмет макроэкономики.", None, []),
        (2, "2", "Функции денег: Мера стоимости; Средство обращения; Средство накопления.", None, ["joined_lines"]),
        (3, "3", "Денежные агрегаты.", None, []),
    ]


def test_question_n_wraps_with_a_capital_letter_are_joined():
    result = parse_question_list(
        "Вопрос 1. Предмет макроэкономики и основные\nМакроэкономические показатели.\nВопрос 2. Инфляция.\nВопрос 3. Безработица.\n"
    )
    assert [(q["label"], q["text"]) for q in result["questions"]] == [
        ("1", "Предмет макроэкономики и основные Макроэкономические показатели."),
        ("2", "Инфляция."),
        ("3", "Безработица."),
    ]


def test_lines_that_may_be_questions_or_headings_are_not_glued():
    # Пропуск в нумерации: строка может быть вопросом 3 с потерянным номером.
    gap = parse_question_list("1. Предмет и метод\n2. ВВП и его измерение\nИнфляция и её виды\n4. Безработица\n")
    assert [q["label"] for q in gap["questions"]] == ["1", "2", "", "4"] and codes(gap) == ["mixed_numbering"]
    # Короткая строка без точки между номерами подряд — заголовок раздела (с 2026-10-07: не вопрос и не перенос).
    heading = parse_question_list("1. Предмет и метод\n2. ВВП\nДенежный рынок\n3. Спрос на деньги\n4. Предложение денег\n")
    assert [q["text"] for q in heading["questions"]] == ["Предмет и метод", "ВВП", "Спрос на деньги", "Предложение денег"]
    assert heading["sections"] == [{"index": 1, "title": "Денежный рынок"}] and codes(heading) == []
    # Короткая строка после последнего вопроса — не перенос.
    tail = parse_question_list("17. Инфляция\n18. Безработица и её виды\nИнфляция!\n")
    assert [q["text"] for q in tail["questions"]] == ["Инфляция", "Безработица и её виды"]


# ── Заголовки разделов внутри списка (живой прогон 2026-10-07) ─────────────────

LIVE_COURSE = Path(__file__).resolve().parent / "fixtures" / "live_course"


def test_the_live_list_has_two_parts_and_no_heading_becomes_a_question():
    # Список «Базы данных» живого прогона: «Практические задания» был вопросом № 19, «Теоретические
    # вопросы» молча пропадал как шапка, и стояло «У части строк нет номера».
    result = parse_question_list((LIVE_COURSE / "questions_databases.txt").read_text(encoding="utf-8"))
    questions = result["questions"]
    assert [q["label"] for q in questions] == [str(n) for n in range(1, 21)]
    assert result["sections"] == [{"index": 1, "title": "Теоретические вопросы"}, {"index": 2, "title": "Практические задания"}]
    assert [q["section"] for q in questions] == [1] * 18 + [2] * 2
    assert questions[18]["text"].startswith("Составьте SQL-запрос")
    assert not any(q["text"] in ("Практические задания", "Теоретические вопросы") for q in questions)
    # Остаётся одна пропущенная строка — гриф «Рассмотрено и одобрено…» под списком.
    assert codes(result) == ["unparsed_lines"] and "1 строку" in result["warnings"][0]["message"]


def test_the_live_list_from_a_docx_gives_the_same_questions_and_parts():
    # DOCX теперь приходит блоками с пустой строкой между ними (parser.py): разбору это не мешает.
    lines = [line for line in (LIVE_COURSE / "questions_databases.txt").read_text(encoding="utf-8").splitlines() if line.strip()]
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    text = extract_text(buffer.getvalue(), "Вопросы.docx")
    assert "\n\n" in text
    from_docx = parse_question_list(text)
    from_text = parse_question_list("\n".join(lines))
    assert compact(from_docx) == compact(from_text) and len(from_docx["questions"]) == 20
    assert from_docx["sections"] == from_text["sections"] and len(from_docx["sections"]) == 2


# ── Пропущенные строки списком (QA 2026-10-07) ──────────────────────────────────
# Сервер сказал «Пропустили 1 строку», а предпросмотр сам угадывал строки по исходнику: показывал восемь
# строк шапки и не доходил до грифа под списком. Теперь ответ перечисляет ровно то, что посчитано.


def skipped_count(result):
    """Число из «Пропустили N строк…»; 0 — такого предупреждения нет."""
    message = next((warning["message"] for warning in result["warnings"] if warning["code"] == "unparsed_lines"), "")
    match = re.match(r"Пропустили (\d+) ", message)
    return int(match.group(1)) if match else 0


def test_the_live_list_lists_its_one_skipped_line_the_signature_under_the_list():
    result = parse_question_list((LIVE_COURSE / "questions_databases.txt").read_text(encoding="utf-8"))
    assert result["skipped"] == [{"text": "Рассмотрено и одобрено на заседании кафедры, протокол № 2 от 01.09.2026.", "before": 20}]
    assert skipped_count(result) == 1


def test_skipped_lines_of_a_ticket_sheet_come_in_order_with_their_place():
    result = parse_question_list(official_sheet())
    assert result["skipped"][:4] == [
        {"text": "Зав. кафедрой ____________ И. И. Иванов", "before": 2},  # как вставили: «__» не съедено как разметка
        {"text": "«___» ____________ 2025 г.", "before": 2},
        {"text": "ФГБОУ ВО «Российский экономический университет имени Г. В. Плеханова»", "before": 2},
        {"text": "Кафедра экономической теории", "before": 2},
    ]
    assert [line["before"] for line in result["skipped"]] == [2, 2, 2, 2, 4, 4, 4, 4, 6, 6]
    assert skipped_count(result) == len(result["skipped"]) == 10


@pytest.mark.parametrize(
    ("text", "skipped"),
    [
        # название курса над билетами — перед первым вопросом
        ("Макроэкономика\nБилет № 1\n1. Предмет и метод.\n2. Инфляция.\nБилет № 2\n1. Безработица.\n2. Деньги.\n",
         [{"text": "Макроэкономика", "before": 0}]),
        # заглавными над списком без номеров
        ("МАКРОЭКОНОМИКА\nПредмет и метод макроэкономики.\nСистема национальных счетов.\nИнфляция и её виды.\n",
         [{"text": "МАКРОЭКОНОМИКА", "before": 0}]),
        # шапка в строке каждого билета — у каждого билета
        ("Билет № 1 Макроэкономика\n1. Предмет и метод.\n2. Инфляция.\nБилет № 2 Макроэкономика\n1. Безработица.\n2. Деньги.\n",
         [{"text": "Макроэкономика", "before": 0}, {"text": "Макроэкономика", "before": 2}]),
        # длинная строка шапки над вопросами с номерами («Перечень вопросов» уходит молча)
        ("Перечень вопросов\nКурс для второго года обучения, весенний семестр\n1. Предмет.\n2. Инфляция.\n3. Деньги.\n",
         [{"text": "Курс для второго года обучения, весенний семестр", "before": 0}]),
        # подпись посреди списка — между вторым и третьим вопросом, как её вставили
        ("1. Предмет и метод.\n2. Инфляция.\nЗав. кафедрой __Иванов__\n3. Безработица.\n",
         [{"text": "Зав. кафедрой __Иванов__", "before": 2}]),
        # чистый список
        ("1. Предмет.\n2. Инфляция.\n3. Деньги.\n", []),
    ],
    ids=["course-over-tickets", "caps-title", "header-on-ticket-lines", "long-header-line", "signature-mid-list", "clean-list"],
)
def test_every_counted_line_is_listed_at_its_place(text, skipped):
    result = parse_question_list(text)
    assert result["skipped"] == skipped
    assert skipped_count(result) == len(skipped)


def test_the_skipped_list_is_capped_and_the_warning_still_counts_every_line():
    result = parse_question_list("1. Предмет и метод.\n2. Инфляция.\n" + "Зав. кафедрой ____ И. И. Иванов\n" * 130)
    assert len(result["skipped"]) == MAX_SKIPPED_LINES and skipped_count(result) == 130
    assert all(line == {"text": "Зав. кафедрой ____ И. И. Иванов", "before": 2} for line in result["skipped"])
    long_line = parse_question_list("1. Предмет.\n2. Инфляция.\n3. Деньги.\nЗав. кафедрой " + "очень длинная подпись " * 20)
    assert len(long_line["skipped"]) == 1 and len(long_line["skipped"][0]["text"]) <= SKIPPED_LINE_CHARS
    assert long_line["skipped"][0]["text"].startswith("Зав. кафедрой очень длинная подпись")


@pytest.mark.parametrize(
    "heading",
    ["Практические задания", "ПРАКТИЧЕСКИЕ ЗАДАНИЯ:", "Задачи", "Раздел 2. SQL", "Запросы к нескольким таблицам",
     "Практические задания (решаются письменно)"],
)
def test_a_heading_between_numbered_questions_is_a_section(heading):
    result = parse_question_list(f"1. Модели данных.\n2. Ключи отношения.\n{heading}\n3. Запрос с соединением.\n4. Подзапросы.\n")
    assert [q["text"] for q in result["questions"]] == ["Модели данных.", "Ключи отношения.", "Запрос с соединением.", "Подзапросы."]
    assert [q["section"] for q in result["questions"]] == [None, None, 1, 1]
    assert len(result["sections"]) == 1 and not result["sections"][0]["title"].endswith(":")
    assert codes(result) == []


def test_a_heading_before_restarted_numbering_is_a_section():
    result = parse_question_list("1. Модели данных.\n2. Ключи отношения.\nОсновы SQL\n1. Оператор SELECT.\n2. Соединения.\n")
    assert compact(result) == [
        (1, "1", "Модели данных.", None, []),
        (2, "2", "Ключи отношения.", None, []),
        (3, "1", "Оператор SELECT.", 1, []),
        (4, "2", "Соединения.", 1, []),
    ]
    assert result["sections"] == [{"index": 1, "title": "Основы SQL"}] and codes(result) == []


@pytest.mark.parametrize(
    "line",
    [
        "Что такое ключ отношения?",  # вопрос со знаком вопроса
        "Нормализация отношений.",  # с точкой — предложение, а не заголовок
        "Функциональные зависимости и аномалии вставки, обновления и удаления в отношениях",  # длинная
        "нормализация отношений",  # со строчной — перенос предыдущего вопроса
    ],
)
def test_lines_that_are_not_headings_stay_as_before(line):
    result = parse_question_list(f"1. Модели данных.\n2. Ключи отношения.\n{line}\n3. Запрос с соединением.\n")
    assert result["sections"] == []
    texts = " | ".join(q["text"] for q in result["questions"])
    assert line.rstrip(".?") in texts


def test_a_part_heading_above_question_one_is_a_section_and_the_title_above_it_still_header():
    result = parse_question_list("Базы данных\nТеоретические вопросы\n1. Модели данных.\n2. Ключи отношения.\n3. Индексы.\n")
    assert [q["text"] for q in result["questions"]] == ["Модели данных.", "Ключи отношения.", "Индексы."]
    assert result["sections"] == [{"index": 1, "title": "Теоретические вопросы"}]
    assert [q["section"] for q in result["questions"]] == [1, 1, 1]


def test_a_list_without_numbers_keeps_every_line_and_takes_only_part_headings_as_sections():
    plain = parse_question_list("Понятие базы данных\nМодели данных\nРеляционная алгебра\nНормализация отношений\n")
    assert [q["text"] for q in plain["questions"]] == ["Понятие базы данных", "Модели данных", "Реляционная алгебра",
                                                      "Нормализация отношений"]
    assert plain["sections"] == [] and codes(plain) == []
    parts = parse_question_list("Теоретические вопросы\nПонятие базы данных\nМодели данных\nПрактические задания\n"
                                "Составьте запрос с группировкой\n")
    assert [(q["text"], q["section"]) for q in parts["questions"]] == [
        ("Понятие базы данных", 1), ("Модели данных", 1), ("Составьте запрос с группировкой", 2)]
    with pytest.raises(ParseError):
        parse_question_list("Теоретические вопросы\nПрактические задания\n")  # одни заголовки — вопросов нет


def test_a_part_heading_inside_a_ticket_does_not_close_it():
    result = parse_question_list(
        "Билет № 1\n1. Модели данных.\nПрактическое задание\n2. Запрос с соединением.\n"
        "Билет № 2\n1. Ключи отношения.\nПрактическое задание\n2. Запрос с группировкой.\n"
        "Билет № 3\n1. Индексы.\nПрактическое задание\n2. Подзапрос.\n"
    )
    assert result["sections"] == []
    assert [ticket["positions"] for ticket in result["tickets"]] == [[1, 2], [3, 4], [5, 6]]


# ── Пределы и пустой ввод ───────────────────────────────────────────────────────


@pytest.mark.parametrize("text", ["", "   \n\n ", "Вопросы к экзамену\nПеречень вопросов\n", "1.\n2.\n— 3 —\n"])
def test_nothing_found_is_400(text):
    with pytest.raises(ParseError) as error:
        parse_question_list(text)
    assert error.value.status == 400 and error.value.detail == NO_QUESTIONS_DETAIL


def test_more_than_150_questions_is_400():
    text = "\n".join(f"{n}. Вопрос номер {n} об экономике" for n in range(1, MAX_QUESTIONS + 2))
    with pytest.raises(ParseError) as error:
        parse_question_list(text)
    assert error.value.status == 400 and error.value.detail == TOO_MANY_DETAIL
    exactly = "\n".join(f"{n}. Вопрос номер {n} об экономике" for n in range(1, MAX_QUESTIONS + 1))
    assert len(parse_question_list(exactly)["questions"]) == MAX_QUESTIONS


def test_more_than_60000_chars_is_413():
    with pytest.raises(ParseError) as error:
        parse_question_list("1. Вопрос\n" + "а" * 60_000)
    assert error.value.status == 413 and error.value.detail == TOO_LONG_DETAIL


@pytest.mark.parametrize("line", ["а", ",а", "а-"])
def test_60000_chars_of_continuation_lines_parse_in_linear_time(line):
    """Склейка переносов линейна: раньше такой ввод занимал 4–9 с и держал цикл событий сервера."""
    text = "1. Вопрос\n" + (line + "\n") * ((60_000 - 10) // (len(line) + 1))
    assert len(text) <= 60_000
    elapsed = []
    for _ in range(2):
        started = time.perf_counter()
        result = parse_question_list(text)
        elapsed.append(time.perf_counter() - started)
    assert len(result["questions"]) == 1 and len(result["questions"][0]["text"]) <= 1000
    assert result["questions"][0]["flags"] == ["joined_lines", "long"]
    assert min(elapsed) < 0.3


def test_150_questions_parse_in_under_300_ms():
    text = "\n".join(
        f"{n}. Вопрос {n}: денежно-кредитная политика, инфляция и безработица в долгосрочном периоде" for n in range(1, 151)
    )
    started = time.perf_counter()
    result = parse_question_list(text)
    elapsed = time.perf_counter() - started
    assert len(result["questions"]) == 150
    assert elapsed < 0.3
