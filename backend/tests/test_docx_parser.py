import io
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from parser import extract_text  # noqa: E402
from utils import SOURCE_CHUNK_CHARS, build_source_chunks, pack_source_blocks  # noqa: E402

LIVE_COURSE = Path(__file__).resolve().parent / "fixtures" / "live_course"


def docx_bytes(doc) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def extract(doc) -> list[str]:
    """Lines of the text in order; the blank lines between blocks have tests of their own below."""
    return [line for line in extract_text(docx_bytes(doc), "Вопросы.docx").split("\n") if line]


_next_abstract = [900]


def add_list(
    doc,
    formats: list[str],
    *,
    start_override: int | None = None,
    texts: list[str | None] | None = None,
    legal: frozenset[int] = frozenset(),
    override_text: str | None = None,
) -> str:
    """Новое определение списка: уровни с форматами formats; возвращает numId.

    texts — шаблоны w:lvlText по уровням (None — без шаблона), legal — уровни с w:isLgl,
    override_text — шаблон уровня 0 в lvlOverride самого numId.
    """
    numbering = doc.part.numbering_part.element
    _next_abstract[0] += 1
    abstract_id = str(_next_abstract[0])
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), abstract_id)
    for level, fmt in enumerate(formats):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(start)
        lvl.append(num_fmt)
        if level in legal:
            lvl.append(OxmlElement("w:isLgl"))
        if texts is not None and level < len(texts) and texts[level] is not None:
            lvl_text = OxmlElement("w:lvlText")
            lvl_text.set(qn("w:val"), texts[level])
            lvl.append(lvl_text)
        abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)  # abstractNum должны идти до num
    else:
        numbering.append(abstract)
    num_id = str(max([int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))] + [0]) + 1)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), num_id)
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), abstract_id)
    num.append(ref)
    if start_override is not None or override_text is not None:
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        if start_override is not None:
            start = OxmlElement("w:startOverride")
            start.set(qn("w:val"), str(start_override))
            override.append(start)
        if override_text is not None:
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), "0")
            lvl_text = OxmlElement("w:lvlText")
            lvl_text.set(qn("w:val"), override_text)
            lvl.append(lvl_text)
            override.append(lvl)
        num.append(override)
    numbering.append(num)
    return num_id


def numbered(paragraph, num_id: str, level: int = 0):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), num_id)
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)
    return paragraph


def test_document_order_tables_and_list_number_style():
    doc = Document()
    doc.add_paragraph("Вопросы к экзамену по макроэкономике")
    doc.add_paragraph("Раздел 1. Введение")
    table = doc.add_table(rows=3, cols=2)
    for row, (number, question) in enumerate([("№", "Вопрос"), ("1", "Предмет и метод макроэкономики."), ("2", "Система национальных счетов.")]):
        table.cell(row, 0).text = number
        table.cell(row, 1).text = question
    doc.add_paragraph("Раздел 2. Деньги")
    doc.add_paragraph("Денежный рынок и его структура", style="List Number")
    doc.add_paragraph("Спрос на деньги", style="List Number")
    doc.add_paragraph("Примечание без номера", style="List Bullet")
    doc.add_paragraph("Конец списка")

    assert extract(doc) == [
        "Вопросы к экзамену по макроэкономике",
        "Раздел 1. Введение",
        "№ | Вопрос",
        "1 | Предмет и метод макроэкономики.",
        "2 | Система национальных счетов.",
        "Раздел 2. Деньги",
        "1. Денежный рынок и его структура",
        "2. Спрос на деньги",
        "Примечание без номера",
        "Конец списка",
    ]


def test_merged_cells_are_written_once_and_empty_cells_skipped():
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Билет № 1"
    table.cell(0, 0).merge(table.cell(0, 2))  # по горизонтали
    table.cell(1, 0).text = "Раздел 1"
    table.cell(1, 0).merge(table.cell(2, 0))  # по вертикали
    table.cell(1, 1).text = "1"
    table.cell(1, 2).text = "Предмет макроэкономики"
    table.cell(2, 1).text = "21"
    table.cell(2, 2).text = "Инфляция"
    assert extract(doc) == [
        "Билет № 1",
        "Раздел 1 | 1 | Предмет макроэкономики",
        "21 | Инфляция",
    ]


def test_cell_text_is_one_line():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "5"
    cell = table.cell(0, 1)
    cell.paragraphs[0].text = "Денежно-кредитная политика:"
    cell.add_paragraph("цели и инструменты")
    run = cell.paragraphs[0].add_run()
    run.add_break()
    assert extract(doc) == ["5 | Денежно-кредитная политика: цели и инструменты"]


def test_auto_numbered_cells_and_direct_numbering():
    doc = Document()
    num_id = add_list(doc, ["decimal"])
    table = doc.add_table(rows=2, cols=2)
    for row, question in enumerate(["Предмет и метод макроэкономики.", "ВВП и способы его измерения."]):
        numbered(table.cell(row, 0).paragraphs[0], num_id)  # номер ставит Word, текста в ячейке нет
        table.cell(row, 1).text = question
    assert extract(doc) == ["1. | Предмет и метод макроэкономики.", "2. | ВВП и способы его измерения."]


def test_levels_counters_and_restart():
    doc = Document()
    outline = add_list(doc, ["decimal", "decimal", "bullet"])
    numbered(doc.add_paragraph("Деньги"), outline, 0)
    numbered(doc.add_paragraph("Функции денег"), outline, 1)
    numbered(doc.add_paragraph("Денежные агрегаты"), outline, 1)
    numbered(doc.add_paragraph("маркер внутри"), outline, 2)
    numbered(doc.add_paragraph("Инфляция"), outline, 0)
    numbered(doc.add_paragraph("Виды инфляции"), outline, 1)
    other = add_list(doc, ["decimal"])
    numbered(doc.add_paragraph("Отдельный список"), other)
    continued = add_list(doc, ["decimalZero"], start_override=21)
    numbered(doc.add_paragraph("Двадцать первый вопрос"), continued)
    numbered(doc.add_paragraph("Двадцать второй вопрос"), continued)
    assert extract(doc) == [
        "1. Деньги",
        "1) Функции денег",
        "2) Денежные агрегаты",
        "маркер внутри",
        "2. Инфляция",
        "1) Виды инфляции",
        "1. Отдельный список",
        "21. Двадцать первый вопрос",
        "22. Двадцать второй вопрос",
    ]


def test_multilevel_labels_follow_the_level_text_like_word():
    # Список кафедры «1. / 1.1. / 1.2.»: метки вопросов должны совпасть с распечаткой.
    doc = Document()
    outline = add_list(doc, ["decimal", "decimal"], texts=["%1.", "%1.%2."])
    numbered(doc.add_paragraph("Введение"), outline, 0)
    numbered(doc.add_paragraph("Предмет и метод макроэкономики."), outline, 1)
    numbered(doc.add_paragraph("Система национальных счетов."), outline, 1)
    numbered(doc.add_paragraph("Деньги"), outline, 0)
    numbered(doc.add_paragraph("Функции денег."), outline, 1)
    numbered(doc.add_paragraph("Денежные агрегаты."), outline, 1)
    assert extract(doc) == [
        "1. Введение",
        "1.1. Предмет и метод макроэкономики.",
        "1.2. Система национальных счетов.",
        "2. Деньги",
        "2.1. Функции денег.",
        "2.2. Денежные агрегаты.",
    ]


def test_level_text_with_words_brackets_and_an_override():
    doc = Document()
    worded = add_list(doc, ["decimal"], texts=["Вопрос %1."])
    numbered(doc.add_paragraph("Рынок труда"), worded)
    numbered(doc.add_paragraph("Инфляция"), worded)
    bracket = add_list(doc, ["decimal", "decimal"], texts=["%1)", "(%2)"])
    numbered(doc.add_paragraph("Спрос на деньги"), bracket, 0)
    numbered(doc.add_paragraph("Мотивы спроса"), bracket, 1)
    overridden = add_list(doc, ["decimal"], texts=["%1."], start_override=7, override_text="№ %1.")
    numbered(doc.add_paragraph("Бюджет"), overridden)
    assert extract(doc) == [
        "Вопрос 1. Рынок труда",
        "Вопрос 2. Инфляция",
        "1) Спрос на деньги",
        "(1) Мотивы спроса",
        "№ 7. Бюджет",
    ]


def test_level_text_falls_back_when_it_needs_a_non_decimal_level():
    doc = Document()
    # Верхний уровень римскими цифрами: «I.1.» в метку вопроса не годится.
    roman = add_list(doc, ["upperRoman", "decimal"], texts=["%1.", "%1.%2."])
    numbered(doc.add_paragraph("Раздел без номера в тексте"), roman, 0)
    numbered(doc.add_paragraph("Первый вопрос раздела"), roman, 1)
    # isLgl: Word пишет все уровни арабскими цифрами, и шаблон снова годится.
    legal = add_list(doc, ["upperRoman", "decimal"], texts=["%1.", "%1.%2."], legal=frozenset({1}))
    numbered(doc.add_paragraph("Раздел"), legal, 0)
    numbered(doc.add_paragraph("Вопрос по закону"), legal, 1)
    # Список начат сразу со второго уровня: ещё не начатый первый уровень — его start.
    nested_first = add_list(doc, ["decimal", "decimal"], texts=["%1.", "%1.%2."])
    numbered(doc.add_paragraph("Сразу вложенный"), nested_first, 1)
    # Пустой шаблон Word показывает без номера.
    hidden = add_list(doc, ["decimal"], texts=[""])
    numbered(doc.add_paragraph("Номер скрыт"), hidden)
    assert extract(doc) == [
        "Раздел без номера в тексте",
        "1) Первый вопрос раздела",
        "Раздел",
        "1.1. Вопрос по закону",
        "1.1. Сразу вложенный",
        "Номер скрыт",
    ]


def test_bullets_letters_and_removed_numbering_get_no_number():
    doc = Document()
    letters = add_list(doc, ["lowerLetter"])
    numbered(doc.add_paragraph("буквенный пункт"), letters)
    doc.add_paragraph("маркированный пункт", style="List Bullet")
    plain = doc.add_paragraph("стиль списка, номер снят", style="List Number")
    numbered(plain, "0")
    doc.add_paragraph("Первый настоящий", style="List Number")
    assert extract(doc) == ["буквенный пункт", "маркированный пункт", "стиль списка, номер снят", "1. Первый настоящий"]


def test_list_number_style_name_is_a_fallback_without_numbering_definition():
    doc = Document()
    style = doc.styles["List Number"]
    p_pr = style.element.find(qn("w:pPr"))
    p_pr.remove(p_pr.find(qn("w:numPr")))  # стиль без numPr, как у некоторых генераторов
    doc.add_paragraph("Первый вопрос", style="List Number")
    doc.add_paragraph("Второй вопрос", style="List Number")
    assert extract(doc) == ["1. Первый вопрос", "2. Второй вопрос"]


def test_empty_numbered_paragraph_advances_the_counter_like_word():
    doc = Document()
    doc.add_paragraph("Первый", style="List Number")
    doc.add_paragraph("", style="List Number")
    doc.add_paragraph("Третий", style="List Number")
    assert extract(doc) == ["1. Первый", "3. Третий"]


def test_tracked_changes_hyperlinks_and_content_controls():
    doc = Document()
    paragraph = doc.add_paragraph("Рынок ")
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), "1")
    inserted.set(qn("w:author"), "Кафедра")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "труда"
    run.append(text)
    inserted.append(run)
    paragraph._p.append(inserted)
    deleted = OxmlElement("w:del")
    deleted.set(qn("w:id"), "2")
    deleted.set(qn("w:author"), "Кафедра")
    del_run = OxmlElement("w:r")
    del_text = OxmlElement("w:delText")
    del_text.text = " и капитала"
    del_run.append(del_text)
    deleted.append(del_run)
    paragraph._p.append(deleted)

    linked = doc.add_paragraph("См. ")
    hyperlink = OxmlElement("w:hyperlink")
    link_run = OxmlElement("w:r")
    link_text = OxmlElement("w:t")
    link_text.text = "программу курса"
    link_run.append(link_text)
    hyperlink.append(link_run)
    linked._p.append(hyperlink)

    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    inner = doc.add_paragraph("Вопрос из поля формы")._p
    content.append(inner)  # переносим абзац внутрь элемента управления содержимым
    sdt.append(content)
    doc.element.body.insert(len(doc.element.body) - 1, sdt)

    assert extract(doc) == ["Рынок труда", "См. программу курса", "Вопрос из поля формы"]


def test_russian_default_style_id_is_resolved():
    doc = Document()
    styles = doc.styles.element
    normal = next(s for s in styles.findall(qn("w:style")) if s.get(qn("w:default")) == "1" and s.get(qn("w:type")) == "paragraph")
    normal.set(qn("w:styleId"), "a")  # так называет «Обычный» русский Word
    doc.add_paragraph("Обычный абзац")
    doc.add_paragraph("Номер по стилю", style="List Number")
    assert extract(doc) == ["Обычный абзац", "1. Номер по стилю"]


def test_broken_docx_raises_a_readable_error():
    with pytest.raises(ValueError, match="Ошибка чтения DOCX"):
        extract_text(b"not a zip at all", "x.docx")


def test_other_formats_are_unchanged():
    assert extract_text("Строка 1\r\nСтрока 2".encode("utf-8"), "a.txt") == "Строка 1\nСтрока 2"
    assert extract_text("<html><body><p>Абзац</p><script>x()</script></body></html>".encode("utf-8"), "a.html") == "Абзац"


# ── Фрагменты [S01]… режутся между абзацами DOCX (живой прогон 2026-10-07) ──────


def sentence(n: int, words: int = 30) -> str:
    return " ".join(f"слово{n}_{k}" for k in range(words)) + "."


def whole_lines(text: str, units: list) -> None:
    """Every line of every fragment is a whole paragraph or table row of the document."""
    known = set()
    for unit in units:
        known.update([unit] if isinstance(unit, str) else unit)
    for chunk in build_source_chunks(text, max_chunks=None):
        assert len(chunk["text"]) <= SOURCE_CHUNK_CHARS
        for line in chunk["text"].split("\n"):
            assert line in known, (chunk["id"], line[:60])


def test_paragraphs_without_blank_lines_are_cut_into_fragments_only_between_paragraphs():
    # Word ставит интервал между абзацами стилем, пустых строк нет: раньше весь текст был одним
    # «абзацем» и резался окнами по 850 знаков посреди фраз.
    doc = Document()
    paragraphs = [sentence(n) for n in range(12)]  # ~300 знаков каждый
    for text in paragraphs:
        doc.add_paragraph(text)
    text = extract_text(docx_bytes(doc), "Лекция.docx")
    assert "\n\n" in text
    whole_lines(text, paragraphs)
    chunks = build_source_chunks(text)
    assert [chunk["id"] for chunk in chunks][:2] == ["S01", "S02"] and len(chunks) == 6  # по два абзаца в фрагменте


def test_a_table_stays_in_one_fragment_or_is_cut_between_rows():
    doc = Document()
    doc.add_paragraph(sentence(1, 70))  # 690 знаков: с вводной строкой и таблицей — больше 850
    doc.add_paragraph("Стандарт SQL определяет четыре уровня изоляции:")
    small = doc.add_table(rows=5, cols=4)
    rows = [("Уровень", "Грязное", "Неповторяющееся", "Фантом"), ("READ UNCOMMITTED", "да", "да", "да"),
            ("READ COMMITTED", "нет", "да", "да"), ("REPEATABLE READ", "нет", "нет", "да"), ("SERIALIZABLE", "нет", "нет", "нет")]
    for r, cells in enumerate(rows):
        for c, value in enumerate(cells):
            small.cell(r, c).text = value
    doc.add_paragraph(sentence(2))
    big = doc.add_table(rows=40, cols=2)
    for r in range(40):
        big.cell(r, 0).text = f"Строка {r}"
        big.cell(r, 1).text = f"значение номер {r} с пояснением в несколько слов"
    text = extract_text(docx_bytes(doc), "Лекция.docx")
    small_rows = [" | ".join(cells) for cells in rows]
    big_rows = [f"Строка {r} | значение номер {r} с пояснением в несколько слов" for r in range(40)]
    whole_lines(text, [sentence(1, 70), "Стандарт SQL определяет четыре уровня изоляции:", sentence(2)] + small_rows + big_rows)
    chunks = build_source_chunks(text, max_chunks=None)
    holders = [chunk["id"] for chunk in chunks if all(row in chunk["text"] for row in small_rows)]
    assert len(holders) == 1  # маленькая таблица целиком
    # Вводная строка с двоеточием уходит в фрагмент таблицы, а не остаётся в конце прошлого.
    assert next(c for c in chunks if c["id"] == holders[0])["text"].startswith("Стандарт SQL определяет")
    assert sum(1 for chunk in chunks if any(row in chunk["text"] for row in big_rows)) >= 2  # большая — по строкам


def test_a_heading_opens_a_fragment_and_a_lead_in_moves_on_with_its_list():
    blocks = pack_source_blocks([
        "Лекция 3. Транзакции",
        sentence(1, 20),
        "1. Понятие транзакции",  # заголовок после законченной фразы — новый фрагмент
        sentence(2, 50),
        "Транзакциями управляют операторы:",
        "START TRANSACTION — начать транзакцию;",
        "COMMIT — зафиксировать транзакцию;" + " очень" * 60,  # не помещается: вводная строка с пунктом — к нему
        "Конец.",
    ]).split("\n\n")
    assert blocks[0].split("\n")[0] == "Лекция 3. Транзакции"
    assert blocks[1].startswith("1. Понятие транзакции\n")
    assert blocks[2].startswith("Транзакциями управляют операторы:\nSTART TRANSACTION")
    assert all(len(block) <= SOURCE_CHUNK_CHARS for block in blocks)


def test_one_paragraph_longer_than_a_fragment_stands_alone_and_short_documents_stay_one_block():
    long_paragraph = sentence(9, 200)
    blocks = pack_source_blocks(["Введение.", long_paragraph, "Итог."]).split("\n\n")
    assert blocks == ["Введение.", long_paragraph, "Итог."]
    assert pack_source_blocks(["Первый абзац.", "Второй абзац.", "третий;"]) == "Первый абзац.\nВторой абзац.\nтретий;"
    # Строка без точки после законченной фразы — заголовок: с неё начинается новый фрагмент.
    assert pack_source_blocks(["Первый абзац.", "Второй раздел", "текст;"]) == "Первый абзац.\n\nВторой раздел\nтекст;"
    assert pack_source_blocks(["Мягкий перенос\n\n\nвнутри абзаца", ["a | b", "  ", "c | d"]]) == "Мягкий перенос\nвнутри абзаца\na | b\nc | d"


def test_the_live_lecture_keeps_the_isolation_levels_table_in_one_fragment():
    # Q18 «Уровни изоляции» получил «частично»: строки таблицы попали в S14 и S15.
    data = (LIVE_COURSE / "lecture3_normalization_transactions.docx").read_bytes()
    text = extract_text(data, "Лекция 3. Нормализация и транзакции.docx")
    chunks = build_source_chunks(text, max_chunks=None)
    table = ["Уровень изоляции | «Грязное» чтение | Неповторяющееся чтение | Фантомное чтение",
             "READ UNCOMMITTED | возможно | возможно | возможно", "READ COMMITTED | нет | возможно | возможно",
             "REPEATABLE READ | нет | нет | возможно", "SERIALIZABLE | нет | нет | нет"]
    holders = [chunk for chunk in chunks if any(row in chunk["text"] for row in table)]
    assert len(holders) == 1 and all(row in holders[0]["text"] for row in table)
    assert "Стандарт SQL определяет четыре уровня изоляции" in holders[0]["text"]
    assert max(len(chunk["text"]) for chunk in chunks) <= SOURCE_CHUNK_CHARS and len(chunks) < 80
    # Ни одна строка фрагмента не обрезана: каждая — целый абзац или строка таблицы документа.
    lines = [line for line in text.split("\n") if line]
    assert all(line in lines for chunk in chunks for line in chunk["text"].split("\n"))
