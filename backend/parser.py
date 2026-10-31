import io
import math
import re
from dataclasses import dataclass, replace
from typing import Any, Optional, Union

from utils import pack_source_blocks

PPT_LEGACY_DETAIL = "Старый формат PowerPoint (.ppt) не читается. Сохраните презентацию как PPTX или PDF и загрузите снова."
DOC_LEGACY_DETAIL = "Старый формат Word (.doc) не читается. Сохраните файл как DOCX или PDF и загрузите снова."
PPTX_NO_TEXT_DETAIL = (
    "В презентации нет текста: похоже, слайды сделаны картинками. "
    "Сохраните её как PDF — такой файл Meldy распознает как скан."
)

# Отказ по формату файла (до запуска анализа, без списания): что подойдёт вместо него.
# Записи лекций Meldy открывает в MP4, MP3, WAV, MPEG и M4A, фото страниц — в JPG и PNG
# (main.RECORDING_EXTENSIONS, main.PHOTO_EXTENSIONS); те же строки показывает выбор файла.
# {subject} — «Формат OGG» (расширение заглавными) или «Этот формат», если расширения нет.
RECORDING_FORMAT_DETAIL = (
    "{subject} не подходит: для записей лекций подходят MP4, MP3, WAV, MPEG и M4A. "
    "Конвертируйте запись, например в MP3, и загрузите снова."
)
PHOTO_FORMAT_DETAIL = (
    "{subject} не подходит: для фото страниц подходят JPG и PNG. Сохраните фото в одном из них и загрузите снова."
)
DOCUMENT_FORMAT_DETAIL = "Этот формат не подходит. Загрузите PDF, DOCX, PPTX, TXT, MD или HTML."
MEDIA_FORMAT_DETAIL = (
    "Этот формат не подходит: для записей лекций подходят MP4, MP3, WAV, MPEG и M4A, для фото страниц — JPG и PNG."
)
# Голосовые сообщения, видео с телефона и прочие записи, которых Meldy не открывает.
_RECORDING_EXTENSIONS = frozenset({
    "ogg", "oga", "opus", "webm", "weba", "mov", "qt", "avi", "mkv", "mka", "m4v", "3gp", "3g2", "aac", "flac",
    "wma", "wmv", "amr", "aif", "aiff", "caf", "mpg", "mpga", "mp2", "ts", "mts", "m2ts", "flv", "f4v", "vob",
})
_PHOTO_EXTENSIONS = frozenset({"heic", "heif", "webp", "gif", "bmp", "tif", "tiff", "avif", "jfif", "dng"})


class UnsupportedFormat(ValueError):
    """A format the app knows but cannot read. The message is shown to the user as is."""


class FileTooLarge(ValueError):
    """A DOCX or PPTX that would unpack to more than the reader takes (a zip bomb). The message is shown as is."""


# A Word or PowerPoint file is checked by its archive's declared sizes before anything is unpacked (review of
# release 49: a 200 KB DOCX unpacked to 200 MB took 427 MB of memory): one XML part up to EXTRACT_XML_PART_BYTES (a
# 300-page lecture's document.xml is a few MB), the whole archive up to bank_docx.MAX_ARCHIVE_BYTES. A PDF gives the
# text of its first PDF_TEXT_MAX_PAGES pages.
EXTRACT_XML_PART_BYTES = 24 * 1024 * 1024
PDF_TEXT_MAX_PAGES = 500
FILE_TOO_LARGE_DETAIL = "Файл слишком большой после распаковки. Разделите его на части и загрузите по очереди."


def _check_office_archive(contents: bytes) -> None:
    import zipfile

    from bank_docx import ArchiveTooLarge, check_archive

    try:
        check_archive(contents, max_part=EXTRACT_XML_PART_BYTES)
    except ArchiveTooLarge:
        raise FileTooLarge(FILE_TOO_LARGE_DETAIL) from None
    except (zipfile.BadZipFile, ValueError, OSError) as exc:
        raise ValueError(f"not an archive: {type(exc).__name__}") from None


def unsupported_format_detail(filename: str, content_type: Optional[str] = None, *, media: bool = False) -> str:
    """Why a file of this format is not taken, in Russian: the old Word and PowerPoint formats by name,
    a recording or a photo by its kind with the formats that fit it, anything else with the list of
    the route's formats (``media`` — /api/analyze/media, else the document analysis)."""
    name = (filename or "").rsplit("/", 1)[-1]
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    mime = (content_type or "").lower()
    if ext == "doc":
        return DOC_LEGACY_DETAIL
    if ext == "ppt":
        return PPT_LEGACY_DETAIL
    subject = f"Формат {ext.upper()}" if ext.isascii() and ext.isalnum() and len(ext) <= 5 else "Этот формат"
    if ext in _RECORDING_EXTENSIONS or (ext not in _PHOTO_EXTENSIONS and mime.startswith(("audio/", "video/"))):
        return RECORDING_FORMAT_DETAIL.format(subject=subject)
    if ext in _PHOTO_EXTENSIONS or mime.startswith("image/"):
        return PHOTO_FORMAT_DETAIL.format(subject=subject)
    return MEDIA_FORMAT_DETAIL if media else DOCUMENT_FORMAT_DETAIL


def extract_text(contents: bytes, filename: str) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else "txt"

    if ext == "pdf":
        return _extract_pdf(contents)
    elif ext == "docx":
        _check_office_archive(contents)
        return _extract_docx(contents)
    elif ext == "pptx":
        _check_office_archive(contents)
        return _extract_pptx(contents)
    elif ext == "ppt":
        raise UnsupportedFormat(PPT_LEGACY_DETAIL)
    elif ext == "doc":
        raise UnsupportedFormat(DOC_LEGACY_DETAIL)
    elif ext in ("html", "htm"):
        return _extract_html(contents)
    elif ext in ("txt", "md"):
        return _extract_text(contents)
    else:
        return _extract_text(contents)


def pdf_page_count(contents: bytes) -> Optional[int]:
    """Pages of a PDF, or None when PyPDF2 cannot read it."""
    try:
        from PyPDF2 import PdfReader

        return len(PdfReader(io.BytesIO(contents)).pages)
    except Exception:
        return None


def _extract_pdf(contents: bytes) -> str:
    """Текст PDF: страницы по порядку, абзацы — блоками для фрагментов лекции (как у DOCX).

    Каждую страницу читает _PdfPage (пробелы по промежуткам между словами, абзацы по
    интервалам); страницу, которую он не умеет измерить, — PyPDF2 (page.extract_text), как раньше.
    """
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(contents))
        units: list[str] = []
        for page in reader.pages[:PDF_TEXT_MAX_PAGES]:
            paragraphs = _pdf_page_paragraphs(page, reader)
            if paragraphs is None:
                page_text = page.extract_text()
                paragraphs = [page_text] if page_text and page_text.strip() else []
            units.extend(paragraphs)
        return clean_extracted_text(pack_source_blocks(units))
    except Exception as e:
        raise ValueError(f"Ошибка чтения PDF: {e}")


# ── Текст страницы PDF по положению букв ─────────────────────────────────────────
#
# PyPDF2 3.0 (page.extract_text) ставит пробел, когда текст сдвигается дальше ~1,9 em от того
# места, где НАЧАЛСЯ предыдущий кусок, а не где он кончился. Chromium и Skia (браузеры, Google
# Docs) пишут жирный заголовок кусками по кернинговым парам, и длинный кусок перед сдвигом
# получал пробел посреди слова: «Строение эук арио тиче ской клетки», «М ембрана и цит оплаз ма»
# (7 из 8 заголовков PDF живого прогона 2026-10-11). Здесь каждый кусок измеряется ширинами
# глифов его шрифта, и пробел встаёт только там, где между концом одного куска и началом
# следующего есть промежуток: на странице, где пробелы — это символы (Chromium, Word), — от
# 0,3 em (табуляция, ячейка таблицы), на странице без символов пробела (TeX) — от 0,15 em.
# Строки делятся по базовой линии; заметно больший интервал между строками или другой размер
# шрифта начинает новый абзац — фрагменты лекции ([S01]…) идут по её абзацам и разделам
# (pack_source_blocks в _extract_pdf), а не режутся по 850 знаков посреди раздела; номер страницы
# в поле под текстом или над ним не попадает в текст. Страница, где есть то, чего этот разбор не измеряет
# (шрифт без ширин — 14 стандартных, Type3, кодировка CMap кроме Identity-H, повёрнутый или
# вертикальный текст, form XObject, письмо справа налево), остаётся за PyPDF2.

_PDF_SPACE_GAP_EM = 0.3  # пробел между кусками на странице, где пробелы — символы
_PDF_TIGHT_SPACE_GAP_EM = 0.15  # пробел на странице без символов пробела
_PDF_EXPLICIT_SPACE_SHARE = 0.05  # столько пробелов среди знаков — пробелы на странице символами
_PDF_SAME_LINE_EM = 0.6  # сдвиг базовой линии меньше — та же строка (индексы, сноски)
_PDF_PARAGRAPH_GAP = 1.25  # интервал больше обычного во столько раз — новый абзац
_PDF_SIZE_CHANGE = 0.05  # размер шрифта другой на столько — новый абзац (заголовок)
_RTL_RE = re.compile("[\u0590-\u08ff\ufb1d-\ufdff\ufe70-\ufeff]")


class _PdfUnsupported(Exception):
    """На странице есть то, чего _PdfPage не измеряет: её текст берётся у PyPDF2."""


@dataclass
class _PdfFont:
    two_byte: bool  # Identity-H: код глифа — два байта
    encoding: Any  # однобайтовый шрифт: {код: знак} или None (знак = код)
    to_unicode: dict
    widths: dict  # код → ширина в долях em
    default_width: float

    def glyphs(self, data: bytes) -> list[tuple[str, float, bool]]:
        """(текст, ширина в em, одиночный байт 32 — к нему прибавляется Tw) для каждого кода."""
        out = []
        if self.two_byte:
            for index in range(0, len(data) - 1, 2):
                code = data[index] << 8 | data[index + 1]
                char = chr(code)
                out.append((self.to_unicode.get(char, char), self.widths.get(code, self.default_width), False))
            return out
        for code in data:
            char = self.encoding.get(code, chr(code)) if isinstance(self.encoding, dict) else chr(code)
            text = self.to_unicode.get(char, char)
            out.append((text if isinstance(text, str) else char, self.widths.get(code, self.default_width), code == 32))
        return out


def _pdf_value(value: Any) -> Any:
    """Значение без косвенной ссылки (PyPDF2 IndirectObject → сам объект)."""
    return value.get_object() if hasattr(value, "get_object") else value


def _pdf_number(value: Any) -> float:
    number = float(_pdf_value(value))
    if not math.isfinite(number):
        raise _PdfUnsupported("number")
    return number


def _pdf_cid_widths(descendant: Any) -> dict:
    """Ширины /W шрифта CID: [c [w1 w2 …]] и [c_first c_last w], в долях em."""
    widths: dict = {}
    raw = descendant.get("/W")
    items = [_pdf_value(item) for item in _pdf_value(raw)] if raw is not None else []
    index = 0
    while index < len(items):
        first = int(items[index])
        following = items[index + 1]
        if isinstance(following, list):
            for offset, width in enumerate(following):
                widths[first + offset] = _pdf_number(width) / 1000.0
            index += 2
        else:
            last = int(following)
            if last - first > 65535:
                raise _PdfUnsupported("widths")
            width = _pdf_number(items[index + 2]) / 1000.0
            for code in range(first, last + 1):
                widths[code] = width
            index += 3
    return widths


def _pdf_font(page: Any, name: str) -> _PdfFont:
    from PyPDF2._cmap import build_char_map

    _, _, encoding, to_unicode, font = build_char_map(name, 200.0, page)
    font = _pdf_value(font)
    subtype = _pdf_value(font.get("/Subtype"))
    if "/DescendantFonts" in font:
        if _pdf_value(font.get("/Encoding")) != "/Identity-H" or encoding != "utf-16-be":
            raise _PdfUnsupported("cmap")
        descendant = _pdf_value(_pdf_value(font["/DescendantFonts"])[0])
        default = _pdf_number(descendant.get("/DW", 1000)) / 1000.0
        return _PdfFont(True, None, to_unicode, _pdf_cid_widths(descendant), default)
    if subtype not in ("/Type1", "/TrueType", "/MMType1") or "/Widths" not in font:
        raise _PdfUnsupported("font")  # Type3 или один из 14 стандартных шрифтов без /Widths
    if not (encoding == "charmap" or isinstance(encoding, dict)):
        raise _PdfUnsupported("encoding")
    first = int(_pdf_value(font.get("/FirstChar", 0)))
    widths = {first + offset: _pdf_number(width) / 1000.0 for offset, width in enumerate(_pdf_value(font["/Widths"]))}
    descriptor = _pdf_value(font.get("/FontDescriptor")) or {}
    default = _pdf_number(descriptor.get("/MissingWidth", 0)) / 1000.0
    return _PdfFont(False, encoding if isinstance(encoding, dict) else None, to_unicode, widths, default)


@dataclass
class _PdfRun:
    text: str
    x0: float
    x1: float
    y: float
    size: float


def _mul(m: list, n: list) -> list:
    return [
        m[0] * n[0] + m[1] * n[2], m[0] * n[1] + m[1] * n[3],
        m[2] * n[0] + m[3] * n[2], m[2] * n[1] + m[3] * n[3],
        m[4] * n[0] + m[5] * n[2] + n[4], m[4] * n[1] + m[5] * n[3] + n[5],
    ]


class _PdfPage:
    """Куски текста страницы с их положением (PDF 32000, 9.4): строки, пробелы, абзацы."""

    def __init__(self, page: Any, reader: Any) -> None:
        from PyPDF2.generic import ContentStream

        self.page = page
        resources = page.get("/Resources")
        resources = resources.get_object() if resources is not None else {}
        fonts = resources.get("/Font")
        self.font_names = set(fonts.get_object().keys()) if fonts is not None else set()
        xobjects = resources.get("/XObject")
        self.forms = set()
        if xobjects is not None:
            for key, value in xobjects.get_object().items():
                if value.get_object().get("/Subtype") == "/Form":
                    self.forms.add(key)
        contents = page.get("/Contents")
        self.operations = ContentStream(contents.get_object(), reader, "bytes").operations if contents is not None else []
        self.fonts: dict = {}
        self.runs: list = []

    def _font(self, name: str) -> _PdfFont:
        if name not in self.fonts:
            if name not in self.font_names:
                raise _PdfUnsupported("missing font")
            self.fonts[name] = _pdf_font(self.page, name)
        return self.fonts[name]

    def read(self) -> list:
        ctm = [1.0, 0.0, 0.0, 1.0, 0.0, 0.0]
        stack: list = []
        tm = [1.0, 0.0, 0.0, 1.0, 0.0, 0.0]
        tlm = list(tm)
        state = {"font": None, "size": 0.0, "tc": 0.0, "tw": 0.0, "th": 1.0, "tl": 0.0, "rise": 0.0}

        def show(data: Any) -> None:
            nonlocal tm
            font = state["font"]
            if font is None or not isinstance(data, bytes):
                raise _PdfUnsupported("text before Tf")
            m = _mul(tm, ctm)
            if m[0] <= 0 or m[3] <= 0 or abs(m[1]) > 1e-3 * abs(m[0]) or abs(m[2]) > 1e-3 * abs(m[3]):
                raise _PdfUnsupported("rotated text")
            size = abs(state["size"] * m[3])
            start = m[4] + state["rise"] * m[2]
            y = m[5] + state["rise"] * m[3]
            text = []
            advance = 0.0
            for chars, width, single_space in font.glyphs(bytes(data)):
                text.append(chars)
                advance += (width * state["size"] + state["tc"] + (state["tw"] if single_space else 0.0)) * state["th"]
            tm = [tm[0], tm[1], tm[2], tm[3], tm[4] + advance * tm[0], tm[5] + advance * tm[1]]
            joined = "".join(text)
            if joined:
                self.runs.append(_PdfRun(joined, start, start + advance * m[0], y, size))

        def move(tx: float, ty: float) -> None:
            nonlocal tm, tlm
            tlm = _mul([1.0, 0.0, 0.0, 1.0, tx, ty], tlm)
            tm = list(tlm)

        for operands, operator in self.operations:
            if operator == b"q":
                stack.append((list(ctm), dict(state)))
            elif operator == b"Q":
                if stack:
                    ctm, state = stack.pop()
            elif operator == b"cm":
                ctm = _mul([_pdf_number(value) for value in operands[:6]], ctm)
            elif operator == b"BT":
                tm = [1.0, 0.0, 0.0, 1.0, 0.0, 0.0]
                tlm = list(tm)
            elif operator == b"Tf":
                state["font"] = self._font(operands[0])
                state["size"] = _pdf_number(operands[1])
            elif operator == b"Tc":
                state["tc"] = _pdf_number(operands[0])
            elif operator == b"Tw":
                state["tw"] = _pdf_number(operands[0])
            elif operator == b"Tz":
                state["th"] = _pdf_number(operands[0]) / 100.0
            elif operator == b"TL":
                state["tl"] = _pdf_number(operands[0])
            elif operator == b"Ts":
                state["rise"] = _pdf_number(operands[0])
            elif operator == b"Tm":
                tm = [_pdf_number(value) for value in operands[:6]]
                tlm = list(tm)
            elif operator == b"Td":
                move(_pdf_number(operands[0]), _pdf_number(operands[1]))
            elif operator == b"TD":
                state["tl"] = -_pdf_number(operands[1])
                move(_pdf_number(operands[0]), _pdf_number(operands[1]))
            elif operator == b"T*":
                move(0.0, -state["tl"])
            elif operator == b"Tj":
                show(operands[0])
            elif operator == b"'":
                move(0.0, -state["tl"])
                show(operands[0])
            elif operator == b'"':
                state["tw"] = _pdf_number(operands[0])
                state["tc"] = _pdf_number(operands[1])
                move(0.0, -state["tl"])
                show(operands[2])
            elif operator == b"TJ":
                for item in operands[0]:
                    if isinstance(item, bytes):
                        show(item)
                    else:
                        shift = -_pdf_number(item) / 1000.0 * state["size"] * state["th"]
                        tm = [tm[0], tm[1], tm[2], tm[3], tm[4] + shift * tm[0], tm[5] + shift * tm[1]]
            elif operator == b"Do" and operands and operands[0] in self.forms:
                raise _PdfUnsupported("form xobject")
        if any(_RTL_RE.search(run.text) for run in self.runs):
            raise _PdfUnsupported("right-to-left text")
        return self.runs

    def lines(self) -> list:
        """[(текст строки, базовая линия, размер шрифта)] в порядке потока страницы.

        Линия и размер строки — у её самого крупного текста, чтобы индекс («c²», «H₂O») или
        знак сноски не уводили строку и не начинали новую."""
        runs = self.read()
        spaces = sum(run.text.count(" ") for run in runs)
        letters = sum(len(run.text) for run in runs) - spaces
        explicit = letters > 0 and spaces >= _PDF_EXPLICIT_SPACE_SHARE * letters
        gap_em = _PDF_SPACE_GAP_EM if explicit else _PDF_TIGHT_SPACE_GAP_EM
        out: list = []
        parts: list = []
        base: Optional[_PdfRun] = None  # основной текст строки
        last: Optional[_PdfRun] = None  # предыдущий кусок: промежуток меряется от его конца

        def flush() -> None:
            text = " ".join("".join(parts).split())
            if text and base is not None:
                out.append((text, base.y, base.size))

        for run in runs:
            if base is not None and abs(run.y - base.y) > _PDF_SAME_LINE_EM * max(run.size, base.size, 1e-6):
                flush()
                parts, base, last = [], None, None
            if last is not None:
                gap = run.x0 - last.x1
                size = max(run.size, last.size)
                # A jump back far over the previous piece (a table drawn column by column) parts words
                # too; a kerning overlap is a fraction of an em.
                if (gap > gap_em * size or gap < -size) and not parts[-1].endswith(" ") and not run.text.startswith(" "):
                    parts.append(" ")
            parts.append(run.text)
            if base is None or run.size > base.size:
                base = run
            last = run
        flush()
        return out

    def paragraphs(self) -> list:
        """Абзацы страницы: строки абзаца — через перевод строки, номер страницы в поле — без него."""
        lines = self.lines()
        steps = [
            y0 - y1
            for (_, y0, size0), (_, y1, size1) in zip(lines, lines[1:])
            if 0.8 * size1 <= y0 - y1 <= 3.0 * size1 and abs(size0 - size1) <= _PDF_SIZE_CHANGE * max(size0, size1)
        ]
        # Обычный шаг строк: нижняя треть шагов, чтобы страница из коротких абзацев не приняла
        # отступ между ними за шаг строки.
        usual = sorted(steps)[int(len(steps) * 0.3)] if steps else None
        if usual is not None:
            lines = _without_page_numbers(lines, usual)
        paragraphs: list = []
        current: list = []
        for index, (text, y, size) in enumerate(lines):
            if current:
                _, y0, size0 = lines[index - 1]
                new_size = abs(size0 - size) > _PDF_SIZE_CHANGE * max(size0, size)
                if y0 - y < 0 or new_size or (usual is not None and y0 - y > usual * _PDF_PARAGRAPH_GAP):
                    paragraphs.append("\n".join(current))
                    current = []
            current.append(text)
        if current:
            paragraphs.append("\n".join(current))
        return paragraphs


_PAGE_NUMBER_RE = re.compile(r"^[-–—\s]*(?:(?:стр|страница|с)\.?\s*)?\d{1,4}(?:\s*(?:из|/)\s*\d{1,4})?[-–—\s]*$", re.IGNORECASE)


def _without_page_numbers(lines: list, usual: float) -> list:
    """Строки без номера страницы: строки из одного номера («2», «— 2 —», «стр. 2 из 9») ниже
    всего остального текста страницы (или выше него) больше чем на полтора шага строки. Номер
    вопроса из таблицы стоит рядом со своей строкой и остаётся."""
    kept = []
    for index, line in enumerate(lines):
        text, y, _ = line
        if _PAGE_NUMBER_RE.match(text):
            others = [other for position, (_, other, _) in enumerate(lines) if position != index]
            if others and (y < min(others) - 1.5 * usual or y > max(others) + 1.5 * usual):
                continue
        kept.append(line)
    return kept


def _pdf_page_paragraphs(page: Any, reader: Any) -> Optional[list]:
    """Абзацы страницы, или None — тогда её текст берёт PyPDF2."""
    try:
        return _PdfPage(page, reader).paragraphs()
    except Exception:
        return None


def _extract_docx(contents: bytes) -> str:
    """Текст DOCX в порядке документа: абзацы и таблицы, с автонумерацией Word.

    - Строка таблицы — одна строка текста: непустые ячейки через « | ».
      Объединённая ячейка (gridSpan, vMerge/hMerge) выводится один раз.
    - Абзац с автонумерацией (numPr у абзаца или его стиля, формат числа decimal*;
      запасной вариант — стиль «List Number») получает номер, как его показывает Word:
      по шаблону уровня w:lvlText, где %k — текущий номер уровня k («%1.» → «3. »,
      «%1.%2.» → «1.2. », «Вопрос %1.» → «Вопрос 4. »). Если шаблона нет или он ссылается
      на уровень не из цифр (буквы, римские), верхний уровень получает «<n>. »,
      вложенные — «<n>) ». Счётчик свой у каждого numId. Маркированные списки номеров не получают.
    - Абзацы и таблицы складываются в блоки до 850 знаков, блоки разделены пустой строкой
      (utils.pack_source_blocks): фрагменты лекции [S01]… режутся только между блоками, то есть
      не посреди абзаца и не посреди строки таблицы; таблица — целиком или по строкам. В Word
      между абзацами пустых строк нет, и раньше весь текст был одним «абзацем», который резался
      окнами по 850 знаков (живой прогон 2026-10-07: таблица уровней изоляции попала в два фрагмента).
    """
    try:
        from docx import Document

        doc = Document(io.BytesIO(contents))
        return clean_extracted_text(pack_source_blocks(_DocxReader(doc).units()))
    except Exception as e:
        raise ValueError(f"Ошибка чтения DOCX: {e}")


_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _w(tag: str) -> str:
    return f"{{{_W_NS}}}{tag}"


def _w_val(element, attribute: str = "val") -> Optional[str]:
    if element is None:
        return None
    return element.get(_w(attribute))


_W_P, _W_TBL, _W_TR, _W_TC, _W_SDT, _W_SDT_CONTENT, _W_CUSTOM_XML = (
    _w("p"), _w("tbl"), _w("tr"), _w("tc"), _w("sdt"), _w("sdtContent"), _w("customXml")
)
_W_T, _W_TAB, _W_PTAB, _W_BR, _W_CR, _W_NO_BREAK_HYPHEN = (
    _w("t"), _w("tab"), _w("ptab"), _w("br"), _w("cr"), _w("noBreakHyphen")
)
# Поддеревья, текст которых в абзац не входит: удалённые правки, надписи (у Word их
# содержимое ещё и продублировано в mc:Fallback), свойства и сноски-ссылки.
_SKIPPED_IN_PARAGRAPH = {
    _w("del"), _w("moveFrom"), _w("txbxContent"), _w("pPr"), _w("rPr"), _w("instrText"),
    f"{{{_MC_NS}}}Fallback",
}
_LIST_NUMBER_STYLES = {"list number": 0, "list number 2": 1, "list number 3": 2, "list number 4": 3, "list number 5": 4}
_LEVEL_REF = re.compile(r"%([1-9])")


@dataclass(frozen=True)
class _Level:
    """Уровень списка Word. Поля, которых нет в XML, — None (у legal — False)."""

    fmt: Optional[str] = None  # w:numFmt: decimal, bullet, lowerLetter…
    start: Optional[int] = None  # w:start
    text: Optional[str] = None  # w:lvlText: «%1.», «%1.%2.», «Вопрос %1.»
    legal: bool = False  # w:isLgl: все уровни в номере пишутся арабскими цифрами


class _DocxNumbering:
    """Определения списков Word: numId → формат, начальное значение и шаблон каждого уровня."""

    def __init__(self, doc) -> None:
        self._num_to_abstract: dict[str, str] = {}
        self._overrides: dict[str, dict[int, _Level]] = {}
        self._abstract_levels: dict[str, dict[int, _Level]] = {}
        self._abstract_style_link: dict[str, str] = {}
        self._styles: dict[str, dict] = {}
        self._default_style: Optional[str] = None
        self._load_numbering(doc)
        self._load_styles(doc)

    def _load_numbering(self, doc) -> None:
        try:
            root = doc.part.numbering_part.element
        except Exception:
            return
        for abstract in root.findall(_w("abstractNum")):
            abstract_id = _w_val(abstract, "abstractNumId")
            if abstract_id is None:
                continue
            link = abstract.find(_w("numStyleLink"))
            if _w_val(link):
                self._abstract_style_link[abstract_id] = _w_val(link)
            self._abstract_levels[abstract_id] = dict(
                _level_definition(lvl) for lvl in abstract.findall(_w("lvl"))
            )
        for num in root.findall(_w("num")):
            num_id = _w_val(num, "numId")
            abstract_id = _w_val(num.find(_w("abstractNumId")))
            if num_id is None or abstract_id is None:
                continue
            self._num_to_abstract[num_id] = abstract_id
            overrides: dict[int, _Level] = {}
            for override in num.findall(_w("lvlOverride")):
                level = _int(_w_val(override, "ilvl"), 0)
                start = _int(_w_val(override.find(_w("startOverride"))), None)
                lvl = override.find(_w("lvl"))
                definition = _level_definition(lvl)[1] if lvl is not None else _Level()
                if start is not None:
                    definition = replace(definition, start=start)
                overrides[level] = definition
            if overrides:
                self._overrides[num_id] = overrides

    def _load_styles(self, doc) -> None:
        try:
            root = doc.styles.element
        except Exception:
            return
        for style in root.findall(_w("style")):
            style_id = style.get(_w("styleId"))
            if not style_id:
                continue
            if style.get(_w("type")) == "paragraph" and style.get(_w("default")) in ("1", "true", "on"):
                # В русском Word у «Обычного» стиля styleId обычно «a», а не «Normal».
                self._default_style = style_id
            num_pr = style.find(_w("pPr") + "/" + _w("numPr"))
            self._styles[style_id] = {
                "name": (_w_val(style.find(_w("name"))) or "").strip().lower(),
                "based_on": _w_val(style.find(_w("basedOn"))),
                "num_id": _w_val(num_pr.find(_w("numId"))) if num_pr is not None else None,
                "ilvl": _int(_w_val(num_pr.find(_w("ilvl"))), None) if num_pr is not None else None,
            }

    def _style_numbering(self, style_id: Optional[str]) -> tuple[Optional[str], Optional[int], Optional[str]]:
        """(numId, уровень, имя первого стиля) по цепочке basedOn."""
        seen: set[str] = set()
        first_name: Optional[str] = None
        current = style_id
        while current and current not in seen:
            seen.add(current)
            style = self._styles.get(current)
            if style is None:
                break
            if first_name is None:
                first_name = style["name"]
            if style["num_id"] is not None:
                return style["num_id"], style["ilvl"], first_name
            current = style["based_on"]
        return None, None, first_name

    def level(self, num_id: str, level: int) -> _Level:
        """Действующее определение уровня: abstractNum, поверх него lvlOverride; start не None."""
        base = _Level()
        override = self._overrides.get(num_id, {}).get(level)
        abstract_id = self._num_to_abstract.get(num_id)
        for _ in range(4):  # numStyleLink → стиль нумерации → его numId → abstractNum с уровнями
            if abstract_id is None or self._abstract_levels.get(abstract_id):
                break
            linked_num, _, _ = self._style_numbering(self._abstract_style_link.get(abstract_id))
            abstract_id = self._num_to_abstract.get(linked_num) if linked_num else None
        if abstract_id is not None:
            base = self._abstract_levels.get(abstract_id, {}).get(level, base)
        if override is not None:
            base = _Level(
                fmt=override.fmt or base.fmt,
                start=override.start if override.start is not None else base.start,
                text=override.text if override.text is not None else base.text,
                legal=override.legal or base.legal,
            )
        return base if base.start is not None else replace(base, start=1)

    def level_format(self, num_id: str, level: int) -> tuple[Optional[str], int]:
        definition = self.level(num_id, level)
        return definition.fmt, definition.start if definition.start is not None else 1

    def paragraph_numbering(self, p) -> Optional[tuple[str, int, bool]]:
        """(ключ счётчика, уровень, нумерованный ли формат) или None, если абзац не в списке."""
        p_pr = p.find(_w("pPr"))
        num_pr = p_pr.find(_w("numPr")) if p_pr is not None else None
        style_id = _w_val(p_pr.find(_w("pStyle"))) if p_pr is not None else None
        num_id = _w_val(num_pr.find(_w("numId"))) if num_pr is not None else None
        level = _int(_w_val(num_pr.find(_w("ilvl"))), None) if num_pr is not None else None
        style_num, style_level, style_name = self._style_numbering(style_id or self._default_style)
        if num_id is None:
            num_id = style_num
        if level is None:
            level = style_level if style_level is not None else 0
        if num_id is not None:
            if num_id == "0":
                return None  # numId 0 явно снимает нумерацию
            fmt, _ = self.level_format(num_id, level)
            if fmt is None and style_name in _LIST_NUMBER_STYLES:
                return f"style:{style_name}", _LIST_NUMBER_STYLES[style_name], True
            return num_id, level, _decimal_like(fmt)
        if style_name in _LIST_NUMBER_STYLES:
            return f"style:{style_name}", _LIST_NUMBER_STYLES[style_name], True
        return None

    def start_value(self, key: str, level: int) -> int:
        if key.startswith("style:"):
            return 1
        return self.level_format(key, level)[1]


def _level_definition(lvl) -> tuple[int, _Level]:
    level = _int(_w_val(lvl, "ilvl"), 0)
    legal = lvl.find(_w("isLgl"))
    return level, _Level(
        fmt=_w_val(lvl.find(_w("numFmt"))),
        start=_int(_w_val(lvl.find(_w("start"))), None),
        text=_w_val(lvl.find(_w("lvlText"))),
        legal=legal is not None and (_w_val(legal) or "true") not in ("0", "false", "off"),
    )


def _decimal_like(fmt: Optional[str]) -> bool:
    return bool(fmt) and fmt.startswith("decimal")


def _int(value: Optional[str], default):
    try:
        return int(value) if value is not None else default
    except (TypeError, ValueError):
        return default


class _DocxReader:
    """Обходит тело документа по порядку: абзацы (строка) и таблицы (список строк таблицы)."""

    def __init__(self, doc) -> None:
        self._doc = doc
        self._numbering = _DocxNumbering(doc)
        self._counters: dict[str, dict[int, int]] = {}

    def units(self) -> list[Union[str, list[str]]]:
        out: list[Union[str, list[str]]] = []
        self._walk_blocks(self._doc.element.body, out)
        return out

    def _walk_blocks(self, parent, out: list[Union[str, list[str]]]) -> None:
        for child in parent:
            tag = child.tag
            if tag == _W_P:
                # Пустой нумерованный абзац сдвигает счётчик (как в Word), но строки не даёт.
                text = self._paragraph(child, keep_bare_number=False)
                if text.strip():
                    out.append(text)
            elif tag == _W_TBL:
                rows = self._table_rows(child)
                if rows:
                    out.append(rows)
            elif tag in (_W_SDT, _W_CUSTOM_XML):
                content = child.find(_W_SDT_CONTENT) if tag == _W_SDT else child
                if content is not None:
                    self._walk_blocks(content, out)

    def _paragraph(self, p, *, keep_bare_number: bool) -> str:
        text = _paragraph_text(p)
        numbering = self._numbering.paragraph_numbering(p)
        if numbering is None:
            return text
        key, level, numbered = numbering
        counters = self._counters.setdefault(key, {})
        for deeper in [lvl for lvl in counters if lvl > level]:
            del counters[deeper]  # верхний уровень начинает вложенные заново
        if level in counters:
            counters[level] += 1
        else:
            counters[level] = self._numbering.start_value(key, level)
        if not numbered:
            return text
        prefix = self._number_prefix(key, level, counters)
        if not text.strip():
            # В ячейке «№» с автонумерацией номер и есть содержимое: «1. | Вопрос».
            return prefix.rstrip() if keep_bare_number else ""
        return f"{prefix}{text.lstrip()}".rstrip()

    def _number_prefix(self, key: str, level: int, counters: dict[int, int]) -> str:
        """Номер абзаца с пробелом после него, как в Word: по шаблону w:lvlText уровня.

        %k заменяется текущим номером уровня k (ещё не начатый уровень — его start).
        Без шаблона или со ссылкой на уровень не из цифр — «<n>. » наверху, «<n>) » глубже.
        Пустой шаблон Word показывает без номера, и здесь номера тоже нет.
        """
        fallback = f"{counters[level]}. " if level == 0 else f"{counters[level]}) "
        if key.startswith("style:"):
            return fallback
        definition = self._numbering.level(key, level)
        if definition.text is None:
            return fallback
        unusable = False

        def number_of(match: re.Match) -> str:
            nonlocal unusable
            referenced = int(match.group(1)) - 1
            if referenced != level and not definition.legal:
                if not _decimal_like(self._numbering.level(key, referenced).fmt):
                    unusable = True
                    return ""
            value = counters.get(referenced)
            return str(value if value is not None else self._numbering.start_value(key, referenced))

        rendered = " ".join(_LEVEL_REF.sub(number_of, definition.text).split())
        if unusable:
            return fallback
        return f"{rendered} " if rendered else ""

    def _table_rows(self, tbl) -> list[str]:
        rows: list[str] = []
        for tr in tbl.iter(_W_TR):
            if _nearest_table(tr) is not tbl:
                continue  # строки вложенных таблиц попадают в текст своей ячейки
            cells: list[str] = []
            for tc in _row_cells(tr):
                if _is_merge_continuation(tc):
                    continue
                cell_text = self._cell_text(tc)
                if cell_text:
                    cells.append(cell_text)
            if cells:
                rows.append(" | ".join(cells))
        return rows

    def _cell_text(self, tc) -> str:
        parts: list[str] = []
        self._collect_cell(tc, parts)
        return " ".join(" ".join(part.split()) for part in parts if part.strip())

    def _collect_cell(self, parent, parts: list[str]) -> None:
        for child in parent:
            tag = child.tag
            if tag == _W_P:
                text = self._paragraph(child, keep_bare_number=True)
                if text.strip():
                    parts.append(text)
            elif tag == _W_TBL:
                parts.extend(self._table_rows(child))
            elif tag in (_W_SDT, _W_CUSTOM_XML):
                content = child.find(_W_SDT_CONTENT) if tag == _W_SDT else child
                if content is not None:
                    self._collect_cell(content, parts)


def _nearest_table(element):
    parent = element.getparent()
    while parent is not None and parent.tag != _W_TBL:
        parent = parent.getparent()
    return parent


def _row_cells(tr) -> list:
    cells = []
    for child in tr:
        if child.tag == _W_TC:
            cells.append(child)
        elif child.tag in (_W_SDT, _W_CUSTOM_XML):
            content = child.find(_W_SDT_CONTENT) if child.tag == _W_SDT else child
            if content is not None:
                cells.extend(c for c in content if c.tag == _W_TC)
    return cells


def _is_merge_continuation(tc) -> bool:
    tc_pr = tc.find(_w("tcPr"))
    if tc_pr is None:
        return False
    for tag in ("vMerge", "hMerge"):
        merge = tc_pr.find(_w(tag))
        if merge is not None and (_w_val(merge) or "continue") != "restart":
            return True
    return False


def _paragraph_text(p) -> str:
    parts: list[str] = []
    _collect_runs(p, parts)
    return "".join(parts)


def _collect_runs(element, parts: list[str]) -> None:
    for child in element:
        tag = child.tag
        if not isinstance(tag, str) or tag in _SKIPPED_IN_PARAGRAPH:
            continue
        if tag == _W_T:
            parts.append(child.text or "")
        elif tag in (_W_TAB, _W_PTAB):
            parts.append("\t")
        elif tag in (_W_BR, _W_CR):
            parts.append("\n")
        elif tag == _W_NO_BREAK_HYPHEN:
            parts.append("-")
        else:
            _collect_runs(child, parts)


# ── PPTX ─────────────────────────────────────────────────────────────────────────

# Фигуры, у которых верх отличается меньше чем на 1/8 дюйма, стоят в одном ряду и читаются слева
# направо: при ручной расстановке колонки редко выровнены точно.
_PPTX_ROW_TOLERANCE_EMU = 114300
_DGM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _extract_pptx(contents: bytes) -> str:
    """Текст презентации по слайдам.

    У каждого слайда с текстом — строка «Слайд N» (N — номер в презентации), затем заголовок,
    затем надписи в порядке чтения: сверху вниз, в одном ряду слева направо. Группы фигур
    раскрываются, таблица — строками «a | b | c», схемы SmartArt — их пункты, у диаграммы —
    её название. Заметки докладчика — «Заметки: …» в конце слайда. Слайды без текста и без
    заметок пропускаются; колонтитулы, дата и номер слайда в текст не входят.
    """
    try:
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(contents))
        blocks: list[str] = []
        for number, slide in enumerate(presentation.slides, start=1):
            lines = _pptx_slide_lines(slide)
            notes = _pptx_notes(slide)
            if notes:
                lines.append("Заметки: " + "\n".join(notes))
            if lines:
                blocks.append("\n".join([f"Слайд {number}", *lines]))
        return clean_extracted_text("\n\n".join(blocks))
    except Exception as e:
        raise ValueError(f"Ошибка чтения PPTX: {e}")


def _pptx_slide_lines(slide) -> list[str]:
    title = slide.shapes.title
    lines: list[str] = []
    if title is not None and title.has_text_frame:
        # Заголовок — одной строкой, даже если в нём перенос.
        heading = " ".join(_pptx_frame_lines(title.text_frame))
        if heading:
            lines.append(heading)
    others = [shape for shape in slide.shapes if title is None or shape.element is not title.element]
    for shape in _pptx_reading_order(others):
        lines.extend(_pptx_shape_lines(shape))
    return lines


def _pptx_position(shape) -> tuple[int, int]:
    try:
        top, left = shape.top, shape.left
    except Exception:
        return 0, 0
    return int(top or 0), int(left or 0)


def _pptx_reading_order(shapes) -> list:
    rows: list[tuple[int, list]] = []
    for shape in sorted(shapes, key=_pptx_position):
        top = _pptx_position(shape)[0]
        if rows and top - rows[-1][0] <= _PPTX_ROW_TOLERANCE_EMU:
            rows[-1][1].append(shape)
        else:
            rows.append((top, [shape]))
    return [shape for _, row in rows for shape in sorted(row, key=lambda item: _pptx_position(item)[1])]


def _pptx_is_service_placeholder(shape) -> bool:
    """Колонтитул, дата и номер слайда: служебные поля, не текст лекции."""
    try:
        if not shape.is_placeholder:
            return False
        from pptx.enum.shapes import PP_PLACEHOLDER

        return shape.placeholder_format.type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER)
    except Exception:
        return False


def _pptx_shape_lines(shape) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if _pptx_is_service_placeholder(shape):
        return []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        lines: list[str] = []
        for child in _pptx_reading_order(list(shape.shapes)):
            lines.extend(_pptx_shape_lines(child))
        return lines
    if getattr(shape, "has_table", False):
        return _pptx_table_rows(shape.table)
    if getattr(shape, "has_text_frame", False):
        return _pptx_frame_lines(shape.text_frame)
    if getattr(shape, "has_chart", False):
        return _pptx_chart_title(shape.chart)
    return _pptx_smartart_lines(shape)


def _pptx_frame_lines(frame) -> list[str]:
    lines: list[str] = []
    for paragraph in frame.paragraphs:
        # python-pptx передаёт перенос строки внутри абзаца (a:br) символом \v.
        for line in paragraph.text.replace("\v", "\n").split("\n"):
            text = " ".join(line.split())
            if text:
                lines.append(text)
    return lines


def _pptx_table_rows(table) -> list[str]:
    rows: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            if getattr(cell, "is_spanned", False):
                continue  # ячейка под объединением: её текст уже у первой ячейки
            text = " ".join(cell.text.replace("\v", " ").split())
            if text:
                cells.append(text)
        if cells:
            rows.append(" | ".join(cells))
    return rows


def _pptx_chart_title(chart) -> list[str]:
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            return _pptx_frame_lines(chart.chart_title.text_frame)
    except Exception:
        pass
    return []


def _pptx_smartart_lines(shape) -> list[str]:
    """Пункты схемы SmartArt: их текст лежит в отдельной части файла (данные схемы)."""
    try:
        rel_ids = shape.element.find(f".//{{{_DGM_NS}}}relIds")
        if rel_ids is None:
            return []
        data_part = shape.part.related_part(rel_ids.get(f"{{{_R_NS}}}dm"))
        from lxml import etree

        root = etree.fromstring(data_part.blob)
    except Exception:
        return []
    lines: list[str] = []
    for point in root.iter(f"{{{_DGM_NS}}}pt"):
        if point.get("type") not in (None, "node"):
            continue  # служебные точки схемы: связи, оформление
        body = point.find(f"{{{_DGM_NS}}}t")
        if body is None:
            continue
        for paragraph in body.iter(f"{{{_A_NS}}}p"):
            text = " ".join("".join(node.text or "" for node in paragraph.iter(f"{{{_A_NS}}}t")).split())
            if text:
                lines.append(text)
    return lines


def _pptx_notes(slide) -> list[str]:
    if not slide.has_notes_slide:  # обращение к notes_slide создало бы пустые заметки
        return []
    frame = slide.notes_slide.notes_text_frame
    return _pptx_frame_lines(frame) if frame is not None else []


def _extract_text(contents: bytes) -> str:
    # BOM must win over heuristic decoding. Without this, UTF-16 files are
    # decoded as CP1251 because CP1251 accepts every byte sequence.
    if contents.startswith((b"\xff\xfe", b"\xfe\xff", b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        return clean_extracted_text(contents.decode("utf-16" if not contents.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")) else "utf-32"))
    if contents.startswith(b"\xef\xbb\xbf"):
        return clean_extracted_text(contents.decode("utf-8-sig"))

    # UTF-16 without a BOM has a high NUL-byte ratio. Detect it before the
    # single-byte fallbacks so it cannot become readable-looking gibberish.
    sample = contents[:4096]
    if sample and sample.count(b"\x00") / len(sample) > 0.20:
        for encoding in ("utf-16-le", "utf-16-be"):
            try:
                return clean_extracted_text(contents.decode(encoding))
            except UnicodeDecodeError:
                pass

    try:
        return clean_extracted_text(contents.decode("utf-8"))
    except UnicodeDecodeError:
        pass

    candidates = []
    for encoding in ("cp1251", "cp866", "latin-1"):
        try:
            decoded = contents.decode(encoding)
            cleaned = clean_extracted_text(decoded)
            if _looks_like_text(cleaned):
                candidates.append((_legacy_decode_score(cleaned), cleaned))
        except UnicodeDecodeError:
            continue
    if candidates:
        return max(candidates, key=lambda item: item[0])[1]
    return clean_extracted_text(contents.decode("utf-8", errors="replace"))


def clean_extracted_text(text: str) -> str:
    """Remove transport/control noise while preserving normal Unicode text."""
    text = text.replace("\ufeff", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x00", "")
    text = re.sub(r"[\u0001-\u0008\u000b\u000c\u000e-\u001f\u007f]", "", text)
    # Trailing spaces and tabs go line by line: the regex r"[ \t]+\n" did the same in quadratic time on a long run
    # of spaces with no line break after it (80,000 spaces took 7.5 s). The last line's tail goes with strip() below.
    text = "\n".join(line.rstrip(" \t") for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _looks_like_text(text: str) -> bool:
    if not text:
        return False
    replacement_count = text.count("\ufffd")
    if replacement_count > max(2, len(text) // 200):
        return False
    printable = sum(char.isprintable() or char in "\n\t" for char in text)
    return printable / len(text) >= 0.85


def _legacy_decode_score(text: str) -> float:
    """Prefer ordinary Russian text over printable legacy-codec mojibake."""
    if not text:
        return float("-inf")

    normal_cyrillic = sum("А" <= char <= "я" or char in "Ёё" for char in text)
    extended_cyrillic = sum("\u0400" <= char <= "\u052f" for char in text) - normal_cyrillic
    latin_letters = sum(char.isascii() and char.isalpha() for char in text)
    suspicious = sum(
        char in "\u00ad\ufffd" or (char.isprintable() and not char.isspace() and ord(char) < 32)
        for char in text
    )
    common_pairs = sum(
        text.lower().count(pair)
        for pair in ("ст", "но", "ен", "ов", "ни", "ра", "во", "ко", "пр", "то")
    )
    return normal_cyrillic * 3 + common_pairs * 2 + latin_letters - extended_cyrillic * 4 - suspicious * 8


def _extract_html(contents: bytes) -> str:
    try:
        from bs4 import BeautifulSoup
        html_str = _extract_text(contents)
        soup = BeautifulSoup(html_str, "html.parser")
        
        # Remove script and style elements
        for script in soup(["script", "style", "header", "footer", "nav"]):
            script.decompose()
            
        # Get text
        text = soup.get_text(separator="\n")
        
        # Clean up spacing and empty lines
        lines = (line.strip() for line in text.splitlines())
        chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
        text = "\n".join(chunk for chunk in chunks if chunk)
        
        return clean_extracted_text(text)
    except Exception as e:
        raise ValueError(f"Ошибка чтения HTML: {e}")
