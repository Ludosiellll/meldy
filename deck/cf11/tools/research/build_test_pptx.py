"""Build a small feasibility PPTX by code (python-pptx + raw OOXML where python-pptx has no API).

Covers: 16:9 1920x1080-equivalent canvas, title as PNG (rendered from HTML), title as vector
outlines (custGeom, no font needed), live text in Onest vs Arial, QR picture with hyperlink,
text hyperlink, Morph transition (named !!shapes, Fade fallback), embedded H.264 MP4 with poster
frame set to autoplay + loop, animated GIF, an Appear build, SVG with PNG fallback, speaker
notes, PowerPoint sections, a hidden reserve slide, optional embedded fonts (EOT or raw TTF).

Usage: python build_test_pptx.py <assets_dir> <out.pptx> [--embed-fonts none|eot|ttf] [--video file.mp4] [--poster file.png]
"""
import argparse, copy, json, struct, uuid
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from fontTools.ttLib import TTFont

NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'p14': 'http://schemas.microsoft.com/office/powerpoint/2010/main',
    'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
}
EMU_PER_PX = 6350          # 13.333 in x 7.5 in slide shown at 1920x1080 -> 1 px = 6350 EMU
px = lambda v: Emu(int(round(v * EMU_PER_PX)))
INK, MUTED, BG, YELLOW, BLUE, LINE = '0B121C', '4B515B', 'F3F3F0', 'F4C45B', '2869DD', 'E3E3DD'

ap = argparse.ArgumentParser()
ap.add_argument('assets'); ap.add_argument('out')
ap.add_argument('--embed-fonts', default='none', choices=['none', 'eot', 'ttf'])
ap.add_argument('--video', default=None); ap.add_argument('--poster', default=None); ap.add_argument('--gif', default=None)
ap.add_argument('--fonts-dir', default=None)
args = ap.parse_args()
A = Path(args.assets)

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)     # 13.333 x 7.5 in = 16:9
LAYOUT_TITLE_ONLY = prs.slide_layouts[5]


def bg(slide, color=BG):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = RGBColor.from_string(color)


def title(slide, text, size=40):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = px(96), px(56), px(1728), px(110)
    tf = t.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = True; r.font.name = 'Onest'; r.font.color.rgb = RGBColor.from_string(INK)
    return t


def textbox(slide, x, y, w, h, text, size=20, font='Onest', color=INK, bold=False, name=None):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = 0
    first = True
    for line in text.split('\n'):
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.name = font; r.font.bold = bold; r.font.color.rgb = RGBColor.from_string(color)
    return tb


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_vector_text(slide, spec_json, x, y, height_px, name, color=INK):
    """Insert outlines from text_to_paths.py as one custGeom freeform (renders without the font)."""
    spec = json.loads(Path(spec_json).read_text())
    w_units, h_units = spec['w'], spec['h']
    height = px(height_px); width = Emu(int(height * w_units / h_units))
    path_xml = []
    for c in spec['cmds']:
        if c[0] == 'M': path_xml.append('<a:moveTo><a:pt x="%d" y="%d"/></a:moveTo>' % tuple(c[1]))
        elif c[0] == 'L': path_xml.append('<a:lnTo><a:pt x="%d" y="%d"/></a:lnTo>' % tuple(c[1]))
        elif c[0] == 'Q': path_xml.append('<a:quadBezTo><a:pt x="%d" y="%d"/><a:pt x="%d" y="%d"/></a:quadBezTo>' % (tuple(c[1]) + tuple(c[2])))
        elif c[0] == 'C': path_xml.append('<a:cubicBezTo>' + ''.join('<a:pt x="%d" y="%d"/>' % tuple(p) for p in c[1:]) + '</a:cubicBezTo>')
        elif c[0] == 'Z': path_xml.append('<a:close/>')
    sp_id = max(int(e.get('id')) for e in slide.shapes._spTree.iter('{%s}cNvPr' % NS['p'])) + 1
    xml = f'''<p:sp xmlns:p="{NS['p']}" xmlns:a="{NS['a']}">
  <p:nvSpPr><p:cNvPr id="{sp_id}" name="{name}" descr="{spec['text']}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
  <p:spPr><a:xfrm><a:off x="{px(x)}" y="{px(y)}"/><a:ext cx="{width}" cy="{height}"/></a:xfrm>
    <a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>
      <a:pathLst><a:path w="{w_units}" h="{h_units}">{''.join(path_xml)}</a:path></a:pathLst></a:custGeom>
    <a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr>
  <p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="ru-RU"/></a:p></p:txBody>
</p:sp>'''
    slide.shapes._spTree.append(etree.fromstring(xml))


def named_shape(slide, kind, x, y, w, h, name, fill, text=None, size=24, color='FFFFFF', radius=None):
    s = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    s.name = name
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill); s.line.fill.background()
    s.shadow.inherit = False
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = radius
    if text:
        tf = s.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = True; r.font.name = 'Onest'
        r.font.color.rgb = RGBColor.from_string(color)
    return s


MORPH_XML = f'''<mc:AlternateContent xmlns:mc="{NS['mc']}" xmlns:p="{NS['p']}">
  <mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" xmlns:p14="{NS['p14']}" Requires="p159">
    <p:transition spd="slow" p14:dur="1600"><p159:morph option="byObject"/></p:transition>
  </mc:Choice>
  <mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback>
</mc:AlternateContent>'''


def set_morph(slide):
    sld = slide._element
    clr = sld.find('{%s}clrMapOvr' % NS['p'])
    clr.addnext(etree.fromstring(MORPH_XML))


def timing_autoplay_loop_plus_appear(video_id, appear_id, video_ms):
    """Main sequence: (1) video starts automatically with the slide (afterEffect, mediacall playFrom),
    (2) one click shows the caption (Appear). Media node loops forever (repeatCount=indefinite)."""
    return f'''<p:timing xmlns:p="{NS['p']}">
 <p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
  <p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
   <p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
    <p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
     <p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
      <p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{video_ms}" fill="hold"/><p:tgtEl><p:spTgt spid="{video_id}"/></p:tgtEl></p:cBhvr></p:cmd>
     </p:childTnLst></p:cTn></p:par>
    </p:childTnLst></p:cTn></p:par>
   </p:childTnLst></p:cTn></p:par>
   <p:par><p:cTn id="7" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
    <p:par><p:cTn id="8" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
     <p:par><p:cTn id="9" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
      <p:set><p:cBhvr><p:cTn id="10" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{appear_id}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
     </p:childTnLst></p:cTn></p:par>
    </p:childTnLst></p:cTn></p:par>
   </p:childTnLst></p:cTn></p:par>
  </p:childTnLst></p:cTn>
  <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
  <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
  <p:video><p:cMediaNode vol="80000"><p:cTn id="11" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{video_id}"/></p:tgtEl></p:cMediaNode></p:video>
 </p:childTnLst></p:cTn></p:par></p:tnLst>
 <p:bldLst><p:bldP spid="{appear_id}" grpId="0"/></p:bldLst>
</p:timing>'''


def ttf_to_eot(ttf_path):
    """Uncompressed EOT (version 0x00020001, no MTX, no XOR) as PowerPoint/LibreOffice .fntdata."""
    data = Path(ttf_path).read_bytes(); f = TTFont(ttf_path); os2 = f['OS/2']; name = f['name']
    def nm(i):
        s = (name.getDebugName(i) or '').encode('utf-16-le'); return struct.pack('<H', len(s)) + s
    panose = os2.panose
    pan = bytes([panose.bFamilyType, panose.bSerifStyle, panose.bWeight, panose.bProportion, panose.bContrast,
                 panose.bStrokeVariation, panose.bArmStyle, panose.bLetterForm, panose.bMidline, panose.bXHeight])
    body = (struct.pack('<I', len(data)) + struct.pack('<I', 0x00020001) + struct.pack('<I', 0) + pan
            + bytes([1, 1 if os2.fsSelection & 1 else 0]) + struct.pack('<I', os2.usWeightClass)
            + struct.pack('<H', os2.fsType) + struct.pack('<H', 0x504C)
            + struct.pack('<IIII', os2.ulUnicodeRange1, os2.ulUnicodeRange2, os2.ulUnicodeRange3, os2.ulUnicodeRange4)
            + struct.pack('<II', os2.ulCodePageRange1, os2.ulCodePageRange2)
            + struct.pack('<I', f['head'].checkSumAdjustment) + b'\0' * 16
            + struct.pack('<H', 0) + nm(1) + struct.pack('<H', 0) + nm(2) + struct.pack('<H', 0) + nm(5)
            + struct.pack('<H', 0) + nm(4) + struct.pack('<H', 0) + struct.pack('<H', 0))  # empty RootString
    header = struct.pack('<I', 4 + len(body) + len(data)) + body
    return header + data


# ------------------------------------------------------------------ slide 1: title techniques
s1 = prs.slides.add_slide(LAYOUT_TITLE_ONLY); bg(s1)
title(s1, 'Титульный', 28)
s1.shapes.add_picture(str(A / 'wordmark.png'), px(96), px(170), height=px(150))
s1.shapes.add_picture(str(A / 'headline.png'), px(96), px(340), width=px(1000))
textbox(s1, 96, 470, 1700, 50, 'Onest (живой текст, 24 пт): Шпаргалка стирается ответами', 24, 'Onest', INK, True, 'live-onest')
textbox(s1, 96, 535, 1700, 50, 'Arial (живой текст, 24 пт): Шпаргалка стирается ответами', 24, 'Arial', INK, True, 'live-arial')
add_vector_text(s1, A / 'title_vec.json', 96, 630, 64, 'vector-title')
textbox(s1, 96, 715, 1100, 40, '↑ вектор (custGeom из контуров Onest), шрифт не нужен', 16, 'Arial', MUTED)
qr = s1.shapes.add_picture(str(A / 'qr_try.png'), px(1530), px(630), height=px(260))
qr.name = 'qr-try'; qr.click_action.hyperlink.address = 'https://ludosiel-studyforge.hf.space/try'
tb = textbox(s1, 1400, 900, 500, 40, '', 16, 'Arial', BLUE)
r = tb.text_frame.paragraphs[0].runs[0]; r.text = 'ludosiel-studyforge.hf.space/try'
r.hyperlink.address = 'https://ludosiel-studyforge.hf.space/try'
textbox(s1, 96, 980, 1150, 40, 'Номинация: «Образование, культура» · CodingFest-11 (тестовый слайд)', 18, 'Arial', MUTED)
notes(s1, 'ЯДРО 0:00–0:30. Текст для чтения со сцены: «Meldy — тренажёр устного экзамена…». '
          'Здесь же — расширенная версия для эксперта, который читает файл без докладчика.')

# ------------------------------------------------------------------ slides 2-3: Morph pair
s2 = prs.slides.add_slide(LAYOUT_TITLE_ONLY); bg(s2); title(s2, 'Содержание')
named_shape(s2, MSO_SHAPE.ROUNDED_RECTANGLE, 96, 260, 520, 300, '!!card', 'FFFFFF', 'Экзамен по билетам', 28, INK, 0.12)
named_shape(s2, MSO_SHAPE.OVAL, 650, 380, 60, 60, '!!dot', YELLOW)
named_shape(s2, MSO_SHAPE.ROUNDED_RECTANGLE, 1300, 220, 300, 640, '!!phone', INK, 'Meldy', 30, 'FFFFFF', 0.18)
textbox(s2, 96, 900, 1500, 60, 'Morph A: три объекта с именами !!card, !!dot, !!phone', 20, 'Arial', MUTED)
notes(s2, 'ЯДРО 0:30–1:00. Камера «наезжает» на телефон на следующем слайде (Morph).')

s3 = prs.slides.add_slide(LAYOUT_TITLE_ONLY); bg(s3); title(s3, 'Представление блоков проекта')
named_shape(s3, MSO_SHAPE.ROUNDED_RECTANGLE, 96, 200, 340, 200, '!!card', 'FFFFFF', 'Экзамен по билетам', 22, INK, 0.12)
named_shape(s3, MSO_SHAPE.OVAL, 980, 160, 120, 120, '!!dot', YELLOW)
named_shape(s3, MSO_SHAPE.ROUNDED_RECTANGLE, 560, 170, 420, 840, '!!phone', INK, 'Meldy', 44, 'FFFFFF', 0.18)
textbox(s3, 1120, 400, 700, 300, 'Morph B: те же имена, другие места и размеры. В старых версиях — Fade.', 24, 'Arial', INK)
set_morph(s3)
notes(s3, 'ЯДРО 1:00–1:40. Переход Morph: телефон вырастает, карточка уходит влево.')

# ------------------------------------------------------------------ slide 4: video + gif + build
s4 = prs.slides.add_slide(LAYOUT_TITLE_ONLY); bg(s4); title(s4, 'Видео (автозапуск, петля) и GIF')
vid = None
if args.video:
    vid = s4.shapes.add_movie(args.video, px(96), px(190), px(1120), px(630), poster_frame_image=args.poster, mime_type='video/mp4')
    vid.name = 'clip-sheet'
if args.gif:
    g = s4.shapes.add_picture(args.gif, px(1260), px(190), width=px(560)); g.name = 'gif-sheet'
cap = textbox(s4, 96, 850, 1700, 80, 'Подпись появляется по щелчку (Appear) — но в обычном режиме и в PDF она видна сразу.', 24, 'Arial', INK, False, 'caption-build')
if vid is not None:
    sld = s4._element
    old = sld.find('{%s}timing' % NS['p'])
    if old is not None: sld.remove(old)
    new_t = etree.fromstring(timing_autoplay_loop_plus_appear(vid.shape_id, cap.shape_id, 6000))
    ext = sld.find('{%s}extLst' % NS['p'])
    (ext.addprevious(new_t) if ext is not None else sld.append(new_t))
notes(s4, 'ЯДРО 1:40–2:40. Видео стартует само и крутится по кругу; подпись — по щелчку.')

# ------------------------------------------------------------------ slide 5: SVG with PNG fallback
s5 = prs.slides.add_slide(LAYOUT_TITLE_ONLY); bg(s5); title(s5, 'Используемые источники (тест SVG)')
pic = s5.shapes.add_picture(str(A / 'title_vec_fallback.png'), px(96), px(300), width=px(1400)); pic.name = 'svg-title'
svg_part = Part(PackURI('/ppt/media/title_vec.svg'), 'image/svg+xml', prs.part.package, (A / 'title_vec.svg').read_bytes())
rid_svg = s5.part.relate_to(svg_part, RT.IMAGE)
blip = pic._element.find('.//{%s}blip' % NS['a'])
blip.append(etree.fromstring(
    '<a:extLst xmlns:a="%s" xmlns:r="%s"><a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
    '<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" r:embed="%s"/></a:ext></a:extLst>'
    % (NS['a'], NS['r'], rid_svg)))
textbox(s5, 96, 520, 1600, 80, 'Сверху: SVG (PowerPoint 2019+/365) с PNG-заменой для 2016, LibreOffice, PDF.', 22, 'Arial', MUTED)
notes(s5, 'ЯДРО 2:40–3:10.')

# ------------------------------------------------------------------ slide 6: hidden reserve slide
s6 = prs.slides.add_slide(LAYOUT_TITLE_ONLY); bg(s6, 'FFFFFF'); title(s6, 'Резерв: доп. доказательство (скрытый слайд)')
textbox(s6, 96, 300, 1600, 200, 'Скрытый слайд: в показе пропускается, вызывается номером + Enter.', 28, 'Arial', INK)
s6._element.set('show', '0')
notes(s6, 'ДОП +1:00. Показывать, только если осталось время.')

# ------------------------------------------------------------------ sections (PowerPoint 2010+)
pres = prs.part._element
ids = [s.get('id') for s in pres.find('{%s}sldIdLst' % NS['p'])]
sec = f'''<p:ext xmlns:p="{NS['p']}" uri="{{521415D9-36F7-43E2-AB2F-B90AF26B5E84}}">
 <p14:sectionLst xmlns:p14="{NS['p14']}">
  <p14:section name="Обязательные слайды" id="{{{str(uuid.uuid4()).upper()}}}"><p14:sldIdLst>{''.join('<p14:sldId id="%s"/>' % i for i in ids[:5])}</p14:sldIdLst></p14:section>
  <p14:section name="Резерв для защиты" id="{{{str(uuid.uuid4()).upper()}}}"><p14:sldIdLst><p14:sldId id="{ids[5]}"/></p14:sldIdLst></p14:section>
 </p14:sectionLst></p:ext>'''
extlst = pres.find('{%s}extLst' % NS['p'])
if extlst is None:
    extlst = etree.SubElement(pres, '{%s}extLst' % NS['p'])
extlst.insert(0, etree.fromstring(sec))

# ------------------------------------------------------------------ embedded fonts (optional)
if args.embed_fonts != 'none':
    FD = Path(args.fonts_dir)
    fonts = [('Onest', {'regular': FD / 'Onest-400.ttf', 'bold': FD / 'Onest-700.ttf'}),
             ('Onest SemiBold', {'regular': FD / 'Onest-600.ttf'})]
    lst = etree.Element('{%s}embeddedFontLst' % NS['p'])
    n = 0
    for face, styles in fonts:
        ef = etree.SubElement(lst, '{%s}embeddedFont' % NS['p'])
        etree.SubElement(ef, '{%s}font' % NS['p'], typeface=face, pitchFamily='2', charset='0')
        for style in ('regular', 'bold', 'italic', 'boldItalic'):
            if style not in styles: continue
            n += 1
            blob = ttf_to_eot(styles[style]) if args.embed_fonts == 'eot' else Path(styles[style]).read_bytes()
            fp = Part(PackURI('/ppt/fonts/font%d.fntdata' % n), 'application/x-fontdata', prs.part.package, blob)
            rid = prs.part.relate_to(fp, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/font')
            etree.SubElement(ef, '{%s}%s' % (NS['p'], style)).set('{%s}id' % NS['r'], rid)
    anchor = pres.find('{%s}notesSz' % NS['p'])
    anchor.addnext(lst)
    pres.set('embedTrueTypeFonts', '1'); pres.set('saveSubsetFonts', '1')

prs.save(args.out)
print('saved', args.out, Path(args.out).stat().st_size, 'bytes')
